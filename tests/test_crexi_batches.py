"""Crexi bulk batches: credit periods, export, Crexi's reply (which finishes the batch), history, the merge rule
(extra values: Units 2, Owner 2...), the numbered export columns, the DRILL's Crexi box, the web routes.

Everything is made up (addresses, owners, phones); no real Crexi data is used.
"""
from __future__ import annotations

import csv
import itertools
from datetime import date
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from bluemoon.crexi import batches as B
from bluemoon.crexi import merge as M
from bluemoon.crexi import period as P
from bluemoon.crexi import reply as R
from bluemoon.db import market_by_slug, now_iso, one, rows, scalar
from bluemoon.drill import credits as C
from bluemoon.drill import engine as E

MARKET = "apartments-springfield-oh"
OTHER = "self-storage-springfield-oh"
TODAY = date(2026, 10, 29)
_ids = itertools.count(1)

HEADERS = ["Property Link", "Property Type", "Address", "City", "State", "Zip Code", "County", "APN", "Owner Name",
           "Mailing Address", "Mailing Address City", "Mailing Address State", "Mailing Address Zip Code",
           "Contact Name", "Phone 1", "Phone 2", "Phone 3", "Phone 4", "Phone 5",
           "Email 1", "Email 2", "Email 3", "Email 4", "Email 5"]


# --- builders -----------------------------------------------------------------------------------------------

def add_prop(conn, market=MARKET, **kw) -> int:
    n = next(_ids)
    m = market_by_slug(conn, market)
    data = {"market_id": m["id"], "property_key": f"cb:{n}", "address": f"{n} Batch Ave", "city": "SPRINGFIELD",
            "state": "OH", "zip": "45503", "county": "CLARK", "apn": f"B-{n:04d}", "owner_name": f"Owner{n} Example",
            "crexi_url": f"https://www.crexi.com/property-records/batch-{n}/h{n:05d}", "raw_json": "{}",
            "flags_json": "[]", "created_at": now_iso(), "updated_at": now_iso()}
    data.update(kw)
    cur = conn.execute(f"INSERT INTO properties({','.join(data)}) VALUES({','.join('?' for _ in data)})",
                       tuple(data.values()))
    conn.commit()
    return int(cur.lastrowid)


def score(conn, pid, value, in_buybox=1, reasons="[]"):
    conn.execute("INSERT OR REPLACE INTO scores(property_id, in_buybox, exclusion_reasons, score, rank, computed_at) "
                 "VALUES(?,?,?,?,?,?)", (pid, in_buybox, reasons, value, None, now_iso()))
    conn.commit()


def seed(conn, n, market=MARKET, top=0.9) -> list[int]:
    """n possible owners in the buy box; the first has the highest score."""
    out = []
    for i in range(n):
        pid = add_prop(conn, market)
        score(conn, pid, round(top - i * 0.001, 4))
        out.append(pid)
    return out


@pytest.fixture
def small(settings):
    settings.raw["crexi"]["credits_per_month"] = 20
    settings.raw["crexi"]["reset_day"] = 29
    return settings


def read_export(settings, rel) -> list[dict]:
    with (settings.path("output") / rel).open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def reply_file(path: Path, export_rows: list[dict], *, skip=(), not_matched=(), phones=True, contact=True,
               title_rows=2, extra=None) -> Path:
    """A made-up Crexi reply in Crexi's layout: blank row, note row, headers, then the rows. `skip`: row numbers
    Crexi never sent; `not_matched`: row numbers it answered 'Not Matched'."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ownership Information"
    ws.append([])
    for i in range(title_rows - 2):
        ws.append([f"Title line {i} (combined by hand)"])
    ws.append([f"{len(export_rows)} properties found, {len(list(not_matched))} not matched, 0 duplicates"])
    ws.append(HEADERS)
    for r in export_rows:
        n = int(r["Row"])
        if n in skip:
            continue
        zip4 = f"{r['Zip Code']}-1234"
        if n in not_matched:
            ws.append([None, "Not Matched", r["Address"].upper(), r["City"].upper(), r["State"], zip4, r["County"],
                       r["APN"]] + [None] * 16)
            continue
        row = [r["Property Link"], "Multifamily", r["Address"].upper(), r["City"].upper(), r["State"], zip4,
               r["County"], r["APN"], f"Test Owner {n}", f"{n} Mail St", "Dayton", "OH", "45402",
               f"Pat Example{n}" if contact else None,
               f"(937) 555-{1000 + n:04d}" if phones else None, None, None, None, None,
               f"pat{n}@example.test" if phones else None, None, None, None, None]
        if extra:
            row = extra(n, row)
        ws.append(row)
    wb.save(path)
    return path


# --- credit periods -----------------------------------------------------------------------------------------

def test_period_math_reset_day_29():
    d = date
    assert P.period_start(d(2026, 10, 29), 29) == d(2026, 10, 29)
    assert P.period_start(d(2026, 10, 28), 29) == d(2026, 9, 29)
    assert P.period_end(d(2026, 10, 29), 29) == d(2026, 11, 28)
    assert P.next_period_start(d(2026, 10, 29), 29) == d(2026, 11, 29)
    # the year boundary
    assert P.period_start(d(2027, 1, 15), 29) == d(2026, 12, 29)
    assert P.next_period_start(d(2026, 12, 29), 29) == d(2027, 1, 29)
    assert P.period_start(d(2027, 1, 29), 29) == d(2027, 1, 29)
    # February has no 29th (2027) so its period starts on its last day, 28th
    assert P.period_end(d(2027, 1, 29), 29) == d(2027, 2, 27)
    assert P.period_start(d(2027, 2, 27), 29) == d(2027, 1, 29)
    assert P.period_start(d(2027, 2, 28), 29) == d(2027, 2, 28)
    assert P.period_end(d(2027, 2, 28), 29) == d(2027, 3, 28)
    assert P.period_start(d(2027, 3, 1), 29) == d(2027, 2, 28)
    # leap year: Feb 29 exists
    assert P.period_start(d(2028, 2, 28), 29) == d(2028, 1, 29)
    assert P.period_start(d(2028, 2, 29), 29) == d(2028, 2, 29)
    # a 31st that April does not have
    assert P.period_start(d(2026, 4, 30), 31) == d(2026, 4, 30)
    assert P.period_start(d(2026, 5, 30), 31) == d(2026, 4, 30)
    assert P.period_start(d(2026, 5, 31), 31) == d(2026, 5, 31)
    # reset day 1 = calendar months
    assert P.period_start(d(2026, 10, 31), 1) == d(2026, 10, 1)
    assert P.period_end(d(2026, 10, 1), 1) == d(2026, 10, 31)


def test_every_day_is_in_exactly_one_period_and_periods_tile():
    day = date(2026, 12, 1)
    seen = {}
    for i in range(500):
        d = date.fromordinal(day.toordinal() + i)
        s = P.period_start(d, 29)
        e = P.period_end(s, 29)
        assert s <= d <= e
        seen[s] = e
    starts = sorted(seen)
    for a, b in zip(starts, starts[1:]):
        assert P.next_period_start(a, 29) == b and seen[a].toordinal() + 1 == b.toordinal()


def test_days_left_and_describe(small):
    assert P.days_left(date(2026, 10, 29), small) == 31          # Oct 29 .. Nov 28
    assert P.days_left(date(2026, 11, 28), small) == 1
    info = P.describe(date(2026, 11, 5), small)
    assert info["key"] == "2026-10-29" and info["end"] == "2026-11-28" and info["next_reset"] == "2026-11-29"
    assert P.label("2026-10-29", small) == "Oct 29 to Nov 28, 2026"
    assert C.days_left_in_month(date(2026, 10, 29), small) == 31
    small.raw["crexi"]["reset_day"] = "banana"
    assert P.reset_day(small) == 29


def test_credit_status_follows_the_period_and_counts_in_flight(conn, small):
    seed(conn, 12)
    s0 = C.credit_status(conn, today=TODAY, settings=small)
    assert (s0["period"], s0["per_month"], s0["left"], s0["returned"], s0["in_flight"]) == ("2026-10-29", 20, 20, 0, 0)
    res = B.export_batch(conn, 5, today=TODAY, settings=small)
    s1 = C.credit_status(conn, today=TODAY, settings=small)
    assert (s1["in_flight"], s1["returned"], s1["used"], s1["left"]) == (5, 0, 5, 15)
    # the day before the reset is the previous period: nothing is counted there
    old = C.credit_status(conn, today=date(2026, 10, 28), settings=small)
    assert (old["period"], old["left"]) == ("2026-09-29", 20)
    assert res["kind"] == "first"
    assert "waiting at Crexi" in s1["message"]


# --- export -------------------------------------------------------------------------------------------------

def test_export_writes_the_file_in_crexi_column_order(conn, small):
    ids = seed(conn, 6)
    res = B.export_batch(conn, 4, today=TODAY, settings=small, created_by="seth@example.test")
    assert res["count"] == 4 and res["file"] == f"crexi/2026-10-29/batch-{res['batch_id']}-4.csv"
    data = read_export(small, res["file"])
    assert list(data[0])[:8] == ["Property Link", "Property Type", "Address", "City", "State", "Zip Code", "County", "APN"]
    assert list(data[0])[-3:] == ["Property ID", "Batch ID", "Row"]
    assert [int(r["Property ID"]) for r in data] == ids[:4]                    # best first
    assert [r["Row"] for r in data] == ["1", "2", "3", "4"] and {r["Batch ID"] for r in data} == {str(res["batch_id"])}
    b = B.get_batch(conn, res["batch_id"])
    assert b["created_by"] == "seth@example.test" and b["requested"] == 4 and b["status"] == "open"
    assert b["markets"] == [MARKET] and b["market_id"] == market_by_slug(conn, MARKET)["id"]


def test_export_default_is_every_free_credit_and_n_is_respected(conn, small):
    seed(conn, 30)
    pv = B.preview_export(conn, None, today=TODAY, settings=small)
    assert pv["count"] == 20 and pv["requested"] == 20 and pv["free"] == 20
    res = B.export_batch(conn, 7, today=TODAY, settings=small)
    assert res["count"] == 7
    st = B.status(conn, today=TODAY, settings=small)
    assert st["default_count"] == 13


def test_export_never_overspends_the_budget(conn, small):
    seed(conn, 30)
    B.export_batch(conn, 15, today=TODAY, settings=small)
    with pytest.raises(B.CrexiError, match="Only 5 credits are free"):
        B.export_batch(conn, 6, today=TODAY, settings=small)
    with pytest.raises(B.CrexiError):
        B.export_batch(conn, 0, today=TODAY, settings=small)
    assert B.export_batch(conn, 5, today=TODAY, settings=small)["count"] == 5
    with pytest.raises(B.CrexiError):
        B.export_batch(conn, 1, today=TODAY, settings=small)


def test_export_leaves_out_sent_dnc_manual_and_confirmed(conn, small):
    ids = seed(conn, 10)
    dnc, confirmed, manual, sent = ids[0], ids[1], ids[2], ids[3]
    E.ensure_states(conn, MARKET)
    E.level_action(conn, confirmed, "confirm", person="Done Deal")
    conn.execute("UPDATE drill_state SET status='dnc' WHERE property_id=?", (dnc,))
    score(conn, manual, 0.95, in_buybox=0, reasons='["manual"]')
    conn.commit()
    first = B.export_batch(conn, 1, today=TODAY, settings=small)
    got_first = [int(r["Property ID"]) for r in read_export(small, first["file"])]
    assert got_first == [ids[3]] or got_first[0] not in (dnc, confirmed, manual)
    sent = got_first[0]
    second = B.export_batch(conn, 20 - 1, today=TODAY, settings=small)
    got = [int(r["Property ID"]) for r in read_export(small, second["file"])]
    assert sent not in got                                                       # never sent twice in a period
    assert not ({dnc, confirmed, manual} & set(got))
    assert len(got) == len(set(got))


def test_export_dedupes_by_link_then_apn_then_address(conn, small):
    a = add_prop(conn, apn="SH-1", address="1 Same St", crexi_url="https://www.crexi.com/property-records/x/dup1")
    b = add_prop(conn, apn="SH-2", address="2 Link St", crexi_url="https://www.crexi.com/property-records/x/dup1")   # same link
    c = add_prop(conn, apn="SH-1", address="3 Parcel St", crexi_url="https://www.crexi.com/property-records/x/c")      # same APN
    d = add_prop(conn, apn="D-9", address="1 SAME ST.", crexi_url="https://www.crexi.com/property-records/x/d")        # same address
    e = add_prop(conn, apn="E-1", address="5 Fine St", crexi_url="https://www.crexi.com/property-records/x/e")
    for i, pid in enumerate((a, b, c, d, e)):
        score(conn, pid, 0.9 - i * 0.01)
    pv = B.preview_export(conn, 5, today=TODAY, settings=small)
    assert pv["count"] == 2 and pv["duplicates_left_out"] == 3
    assert "3 duplicates left out" in pv["left_out_text"]
    res = B.export_batch(conn, 5, today=TODAY, settings=small)
    assert sorted(int(r["Property ID"]) for r in read_export(small, res["file"])) == sorted([a, e])
    # a property that duplicates one already sent this period is never sent again
    f = add_prop(conn, apn="SH-1", address="9 Late St", crexi_url="https://www.crexi.com/property-records/x/f")
    score(conn, f, 0.99)
    nxt = B.preview_export(conn, 5, today=TODAY, settings=small)
    assert f not in [s["property_id"] for s in nxt["sample"]] and nxt["duplicates_left_out"] >= 1


def test_export_dry_run_writes_nothing(conn, small):
    seed(conn, 5)
    pv = B.preview_export(conn, 3, today=TODAY, settings=small)
    assert pv["count"] == 3 and pv["reasons"][0]["count"] == 3 and len(pv["sample"]) == 3
    assert scalar(conn, "SELECT COUNT(*) FROM crexi_batches") == 0
    assert not (small.path("output") / "crexi").exists()


def test_export_fills_the_pool_across_markets_and_marks_the_batch(conn, small):
    seed(conn, 3, MARKET)
    seed(conn, 3, OTHER, top=0.5)
    res = B.export_batch(conn, 6, today=TODAY, settings=small)
    b = B.get_batch(conn, res["batch_id"])
    assert sorted(b["markets"]) == sorted([MARKET, OTHER]) and b["market_id"] is None


def test_cells_that_start_like_a_formula_are_quoted_but_numbers_are_not(conn, small, tmp_path):
    assert B.safe_cell("=SUM(A1)") == "'=SUM(A1)" and B.safe_cell("+1") == "'+1"
    assert B.safe_cell("-5 evil") == "'-5 evil" and B.safe_cell("@cmd") == "'@cmd" and B.safe_cell("\tx") == "'\tx"
    assert B.safe_cell(-83.8088) == -83.8088 and B.safe_cell(None) == "" and B.safe_cell("plain") == "plain"
    pid = add_prop(conn, address="=HYPERLINK(1)", name="@bad", lat=39.9, lng=-83.8)
    score(conn, pid, 0.9)
    res = B.export_batch(conn, 1, today=TODAY, settings=small)
    row = read_export(small, res["file"])[0]
    assert row["Address"] == "'=HYPERLINK(1)" and row["Property Name"] == "'@bad" and row["Longitude"] == "-83.8"


def test_export_file_lookup_is_by_id_inside_the_output_folder(conn, small, tmp_path):
    seed(conn, 2)
    res = B.export_batch(conn, 2, today=TODAY, settings=small)
    assert B.export_file(conn, res["batch_id"], small).name.startswith("batch-")
    conn.execute("UPDATE crexi_batches SET file_name = ? WHERE id = ?", ("../../../etc/passwd", res["batch_id"]))
    with pytest.raises(KeyError):
        B.export_file(conn, res["batch_id"], small)
    with pytest.raises(KeyError):
        B.export_file(conn, 9999, small)


# --- reply: matching ----------------------------------------------------------------------------------------

def _batch(conn, small, n, extra_ids=()):
    ids = seed(conn, n)
    res = B.export_batch(conn, n, today=TODAY, settings=small)
    return ids, res, read_export(small, res["file"])


def test_reply_for_a_batch_partial_return_and_next_batch_never_repeats(conn, small, tmp_path):
    ids, res, data = _batch(conn, small, 20)
    seed(conn, 10, top=0.5)                       # more owners to export later
    path = reply_file(tmp_path / "reply.xlsx", data, skip={19, 20}, not_matched={17, 18})
    pv = R.preview(conn, path, "reply.xlsx", settings=small)
    assert "reply to batch #%d (20 sent): Crexi returned 16 of 20, 4 missing" % res["batch_id"] in pv["message"]
    b = pv["summary"]["batches"][0]
    assert (b["returned"], b["missing"], b["not_matched"], b["not_in_file"]) == (16, 4, 2, 2)
    assert scalar(conn, "SELECT COUNT(*) FROM crexi_batch_items WHERE status='returned'") == 0   # preview wrote nothing
    assert C.credit_status(conn, today=TODAY, settings=small)["in_flight"] == 20
    out = R.apply_reply(conn, path, "reply.xlsx", actor="seth@example.test", settings=small)
    assert out["returned"] == 16 and out["not_matched"] == 2 and out["found"] == 16
    assert out["headline"] == "Crexi returned 16 of 20. 4 missing." and out["missing"] == 4
    assert "go to the DRILL's other sources" in out["missing_text"]
    assert out["credits_text"].startswith("You have 4 credits left until Nov 28.")
    # importing finished the batch at once: the 4 Crexi didn't have cost nothing and are free again (no close step)
    st = C.credit_status(conn, today=TODAY, settings=small)
    assert (st["returned"], st["in_flight"], st["left"]) == (16, 0, 4)
    assert B.get_batch(conn, res["batch_id"])["status"] == "closed"
    tip = B.status(conn, today=TODAY, settings=small)["topup"]
    assert tip["suggested"] == 4 and "Crexi returned 16 of 20. 4 missing." in tip["message"]
    assert "Export the next 4" in tip["message"]
    nxt = B.export_batch(conn, 4, today=TODAY, settings=small)
    assert nxt["kind"] == "topup"
    sent_before = {int(r["Property ID"]) for r in data}
    sent_now = {int(r["Property ID"]) for r in read_export(small, nxt["file"])}
    assert len(sent_now) == 4 and not (sent_before & sent_now)


def test_matching_order_link_then_apn_then_address(conn, small, tmp_path):
    ids, res, data = _batch(conn, small, 4)

    def extra(n, row):
        if n == 2:
            row[0] = None                                                # no link: the APN finds it
            row[2], row[3] = "999 CHANGED ST", "ELSEWHERE"
        if n == 3:
            row[0], row[7] = None, None                                  # no link, no APN: address + city + zip
            row[5] = "45503-9776"                                        # ZIP+4
        if n == 4:
            row[0] = "https://www.crexi.com/property-records/other/zzzz"  # a link that is not ours -> falls back
            row[7] = "B-0000"
            row[2], row[3] = "1 NOWHERE RD", "NOWHERE"
        return row

    path = reply_file(tmp_path / "r.xlsx", data, extra=extra)
    pv = R.preview(conn, path, "r.xlsx", settings=small)
    c = pv["summary"]["counts"]
    assert c["new"] == 3 and c["none"] == 1
    assert pv["matched_by"]["link"] == 1 and pv["matched_by"]["apn"] == 1 and pv["matched_by"]["address"] == 1
    out = R.apply_reply(conn, path, "r.xlsx", settings=small)
    assert out["returned"] == 3 and out["unmatched"] == 1
    assert {r["property_id"] for r in rows(conn, "SELECT property_id FROM crexi_batch_items WHERE status='returned'")} \
        == set(ids[:3])


def test_unmatched_rows_are_listed_in_a_file_never_dropped(conn, small, tmp_path):
    ids, res, data = _batch(conn, small, 3)

    def extra(n, row):
        return row

    path = reply_file(tmp_path / "r.xlsx", data[:2] + [], skip=set())
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    ws.append(["https://www.crexi.com/property-records/x/nobody", "Retail", "1 STRANGER AVE", "OTHERTOWN", "OH", "44444",
               "LUCAS", "X-1", "Stranger Co", None, None, None, None, None, None, None, None, None, None, None, None,
               None, None, None])
    wb.save(path)
    out = R.apply_reply(conn, path, "r.xlsx", settings=small)
    assert out["unmatched"] == 1 and out["unmatched_file"]
    text = (small.path("output") / out["unmatched_file"]).read_text(encoding="utf-8-sig")
    assert "1 STRANGER AVE" in text and "Does not match any property" in text
    assert one(conn, "SELECT unmatched FROM crexi_batch_imports")["unmatched"] == 1
    assert scalar(conn, "SELECT COUNT(*) FROM properties WHERE address LIKE '%STRANGER%'") == 0   # never added


def test_unmatched_file_header_from_crexi_cannot_run_as_a_formula(conn, small, tmp_path):
    ids, res, data = _batch(conn, small, 2)
    path = reply_file(tmp_path / "r.xlsx", data)
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    hdr = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "Property Link")
    ws.cell(hdr, len(HEADERS) + 1).value = '=HYPERLINK("http://evil.example","x")'
    ws.cell(hdr, len(HEADERS) + 2).value = ' +cmd'
    ws.append(["https://www.crexi.com/property-records/x/nobody", "Retail", "1 STRANGER AVE", "OTHERTOWN", "OH",
               "44444"] + [None] * (len(HEADERS) - 6) + ["v", "w"])
    wb.save(path)
    out = R.apply_reply(conn, path, "r.xlsx", settings=small)
    with (small.path("output") / out["unmatched_file"]).open(encoding="utf-8-sig", newline="") as f:
        head = next(csv.reader(f))
    assert "'=HYPERLINK(\"http://evil.example\",\"x\")" in head and "'+cmd" in head
    assert not any(h.lstrip(" ").startswith(("=", "+", "-", "@")) for h in head)


def test_reply_importing_twice_changes_nothing(conn, small, tmp_path):
    ids, res, data = _batch(conn, small, 6)
    path = reply_file(tmp_path / "r.xlsx", data, skip={6}, not_matched={5})

    def snapshot():
        return {t: scalar(conn, f"SELECT COUNT(*) FROM {t}") for t in
                ("crexi_batch_items", "credit_ledger", "contact_points", "people", "drill_attempts",
                 "crexi_batch_imports", "property_alternates")} | {
            "status": C.credit_status(conn, today=TODAY, settings=small)["left"]}

    first = R.apply_reply(conn, path, "r.xlsx", settings=small)
    assert first["changed"] and first["returned"] == 4
    before = snapshot()
    second = R.apply_reply(conn, path, "r.xlsx", settings=small)
    assert second["changed"] is False and "Nothing new" in second["message"]
    assert snapshot() == before
    pv = R.preview(conn, path, "r.xlsx", settings=small)
    assert pv["changes"] is False and pv["summary"]["counts"].get("already") == 5


def test_one_file_can_hold_rows_from_several_batches_newest_first(conn, small, tmp_path):
    ids = seed(conn, 8)
    one_ = B.export_batch(conn, 4, today=TODAY, settings=small)
    two = B.export_batch(conn, 4, today=TODAY, settings=small)
    rows1, rows2 = read_export(small, one_["file"]), read_export(small, two["file"])
    combined = [dict(r, Row=str(i)) for i, r in enumerate(rows1 + rows2, start=1)]
    path = reply_file(tmp_path / "both.xlsx", combined, skip={4})
    pv = R.preview(conn, path, "both.xlsx", settings=small)
    assert [b["batch_id"] for b in pv["summary"]["batches"]] == [two["batch_id"], one_["batch_id"]]
    out = R.apply_reply(conn, path, "both.xlsx", settings=small)
    assert out["returned"] == 7 and sorted(out["batch_ids"]) == sorted([one_["batch_id"], two["batch_id"]])
    assert B.get_batch(conn, two["batch_id"])["status"] == "closed"         # every batch it touches is finished
    assert B.get_batch(conn, one_["batch_id"])["status"] == "closed"
    assert B.batch_counts(conn, one_["batch_id"])["missing"] == 1           # the row left out is missing
    assert out["headline"] == "Crexi returned 7 of 8. 1 missing."


def test_a_batch_with_no_row_in_the_file_is_not_touched(conn, small, tmp_path):
    seed(conn, 6)
    first = B.export_batch(conn, 3, today=TODAY, settings=small)
    second = B.export_batch(conn, 3, today=TODAY, settings=small)
    reply = reply_file(tmp_path / "r.xlsx", read_export(small, second["file"]), skip={3})
    out = R.apply_reply(conn, reply, "r.xlsx", settings=small)
    assert out["batch_ids"] == [second["batch_id"]]
    assert B.get_batch(conn, first["batch_id"])["status"] == "open"         # that file was not its reply
    assert B.batch_counts(conn, first["batch_id"])["waiting"] == 3
    assert B.batch_counts(conn, second["batch_id"])["missing"] == 1
    st = C.credit_status(conn, today=TODAY, settings=small)
    assert (st["returned"], st["in_flight"], st["left"]) == (2, 3, 15)


def test_duplicate_rows_in_a_reply_are_not_new_records(conn, small, tmp_path):
    ids, res, data = _batch(conn, small, 3)
    path = reply_file(tmp_path / "r.xlsx", data)
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    dup = [c.value for c in ws[4]]
    dup[14] = "(937) 555-9999"                                   # the copy carries one more phone
    dup[16 - 1 + 1] = None
    ws.append(dup)
    wb.save(path)
    props_before = scalar(conn, "SELECT COUNT(*) FROM properties")
    pv = R.preview(conn, path, "r.xlsx", settings=small)
    assert pv["enrichment"]["duplicates"] == 1 and pv["summary"]["counts"]["duplicate"] == 1
    out = R.apply_reply(conn, path, "r.xlsx", settings=small)
    assert out["returned"] == 3 and out["duplicates"] == 1
    assert scalar(conn, "SELECT COUNT(*) FROM properties") == props_before
    pid = int(data[0]["Property ID"])
    phones = {r["value"] for r in rows(conn, "SELECT value FROM contact_points WHERE property_id=? AND channel='phone'",
                                       (pid,))}
    assert phones == {"(937) 555-1001", "(937) 555-9999"}                       # unioned, not lost, not duplicated


# --- reply: contacts and the DRILL ------------------------------------------------------------------------

def test_contacts_are_untested_crexi_contacts_of_the_contact_person(conn, small, tmp_path):
    ids, res, data = _batch(conn, small, 3)
    pid = int(data[0]["Property ID"])

    def extra(n, row):
        row[14] = "937-555-2000"
        row[15] = "(937) 555-2001"
        row[19] = "pat@example.test; second@example.test"
        return row

    path = reply_file(tmp_path / "r.xlsx", data, extra=extra)
    out = R.apply_reply(conn, path, "r.xlsx", settings=small)
    assert out["new_phones"] == 6 and out["new_emails"] == 6
    cps = rows(conn, "SELECT * FROM contact_points WHERE property_id=? ORDER BY id", (pid,))
    assert {c["channel"] for c in cps} == {"phone", "email"} and {c["status"] for c in cps} == {"untested"}
    assert {c["source"] for c in cps} == {"crexi"}
    assert all(c["notes"] == f"Crexi credit, batch {res['batch_id']}" for c in cps)
    assert {c["value"] for c in cps if c["channel"] == "phone"} == {"(937) 555-2000", "(937) 555-2001"}
    person = one(conn, "SELECT * FROM people WHERE property_id=? AND name='Pat Example1'", (pid,))
    assert person["source"] == "crexi" and person["role"] == "other"
    assert {c["person_id"] for c in cps} == {person["id"]}                  # the phones and emails are Pat's
    item = one(conn, "SELECT * FROM crexi_batch_items WHERE property_id=?", (pid,))
    assert (item["status"], item["result"], item["phones"], item["emails"]) == ("returned", "found", 2, 2)
    att = one(conn, "SELECT * FROM drill_attempts WHERE property_id=? AND stage_key=?", (pid, E.CREDIT_STAGE))
    assert att["outcome"] == "found" and f"batch #{res['batch_id']}" in att["detail"]
    ledger = one(conn, "SELECT * FROM credit_ledger WHERE property_id=?", (pid,))
    assert (ledger["result"], ledger["batch_id"], ledger["period"]) == ("found", res["batch_id"], "2026-10-29")
    # phones, emails and the contact name are contacts, not original property columns (the exports number them)
    raw = one(conn, "SELECT raw_json FROM properties WHERE id=?", (pid,))["raw_json"]
    assert "Phone 1" not in raw and "Email 1" not in raw and "Contact Name" not in raw


def test_do_not_contact_is_respected(conn, small, tmp_path):
    ids, res, data = _batch(conn, small, 2)
    p0, p1 = int(data[0]["Property ID"]), int(data[1]["Property ID"])
    E.add_contact(conn, p1, "phone", "(937) 555-1001", source="manual")           # same number as owner 1's, on another property
    cid = one(conn, "SELECT id FROM contact_points WHERE property_id=?", (p1,))["id"]
    E.set_contact_status(conn, cid, "dnc")
    path = reply_file(tmp_path / "r.xlsx", data)
    out = R.apply_reply(conn, path, "r.xlsx", settings=small)
    assert out["dnc_marked"] >= 1
    mine = one(conn, "SELECT status FROM contact_points WHERE property_id=? AND channel='phone'", (p0,))
    assert mine["status"] == "dnc"                                                # on the do-not-contact list
    other = one(conn, "SELECT status FROM contact_points WHERE property_id=? AND channel='phone' AND "
                      "value='(937) 555-1002'", (p0,)) or one(conn, "SELECT status FROM contact_points WHERE "
                                                              "property_id=? AND channel='email'", (p0,))
    assert other["status"] == "untested"


def test_matched_row_with_a_name_but_no_phone_or_email_is_owner_only(conn, small, tmp_path):
    ids, res, data = _batch(conn, small, 3)
    path = reply_file(tmp_path / "r.xlsx", data, phones=False)
    out = R.apply_reply(conn, path, "r.xlsx", settings=small)
    assert (out["found"], out["owner_only"], out["no_data"]) == (0, 3, 0) and out["returned"] == 3
    pid = int(data[0]["Property ID"])
    assert scalar(conn, "SELECT COUNT(*) FROM contact_points WHERE property_id=?", (pid,)) == 0   # no phones claimed
    att = one(conn, "SELECT * FROM drill_attempts WHERE property_id=? AND stage_key=?", (pid, E.CREDIT_STAGE))
    # the credit stage's goal is a phone or email: a name only is 'nothing' for the DRILL and the hit rate
    assert att["outcome"] == "nothing" and "name only" in att["detail"] and "no_contact" not in att["detail"]
    st = C.credit_status(conn, today=TODAY, settings=small)
    assert st["returned"] == 3                                                  # a matched row still costs a credit
    assert (st["found"], st["nothing"], st["hit_rate"]) == (0, 3, 0.0)
    assert scalar(conn, "SELECT COUNT(*) FROM people WHERE property_id=? AND source='crexi'", (pid,)) >= 1  # names kept
    hist = B.property_history(conn, pid)
    assert hist["state"] == "data" and "owner name but no phone or email" in hist["text"]


def test_preview_report_matches_the_headline_and_a_reimport_says_nothing_new(conn, small, tmp_path):
    ids, res, data = _batch(conn, small, 4)
    path = reply_file(tmp_path / "r.xlsx", data, not_matched={4})
    pv = R.preview(conn, path, "r.xlsx", settings=small)
    assert "Crexi returned 3 of 4" in pv["message"]
    assert pv["enrichment"]["matched"] == 3 and "3 matched" in pv["enrichment_text"]
    R.apply_reply(conn, path, "r.xlsx", settings=small)
    again = R.preview(conn, path, "r.xlsx", settings=small)
    assert again["enrichment"]["matched"] == 0
    assert "Nothing new" in again["enrichment_text"] and "gained nothing" not in again["enrichment_text"]


def test_matched_row_with_nothing_at_all_is_nothing(conn, small, tmp_path):
    ids, res, data = _batch(conn, small, 2)

    def extra(n, row):
        row[8] = None
        return row

    path = reply_file(tmp_path / "r.xlsx", data, phones=False, contact=False, extra=extra)
    out = R.apply_reply(conn, path, "r.xlsx", settings=small)
    assert out["no_data"] == 2 and out["found"] == 0
    pid = int(data[0]["Property ID"])
    att = one(conn, "SELECT * FROM drill_attempts WHERE property_id=? AND stage_key=?", (pid, E.CREDIT_STAGE))
    assert att["outcome"] == "nothing"
    assert B.property_history(conn, pid)["state"] == "none"


def test_override_which_columns_hold_phones_and_emails(conn, small, tmp_path):
    ids, res, data = _batch(conn, small, 2)
    path = reply_file(tmp_path / "r.xlsx", data)
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    ws.cell(row=3, column=5).value = "Landline"            # 'State' ... is mapped, so use the free columns instead
    wb.close()
    pv = R.preview(conn, path, "r.xlsx", settings=small)
    assert pv["columns"]["phone"] == ["Phone 1", "Phone 2", "Phone 3", "Phone 4", "Phone 5"]
    assert pv["columns"]["email"][0] == "Email 1" and pv["columns"]["name"] == "Contact Name"
    only = R.preview(conn, path, "r.xlsx", columns={"phone": ["Phone 2"], "email": []}, settings=small)
    assert only["columns"]["phone"] == ["Phone 2"] and only["columns"]["email"] == []
    assert only["enrichment"]["new_phones"] == 0                       # Phone 2 is empty in the made-up file
    out = R.apply_reply(conn, path, "r.xlsx", columns={"phone": ["Phone 1"], "email": []}, settings=small)
    assert out["new_phones"] == 2 and out["new_emails"] == 0
    assert B.get_batch(conn, res["batch_id"])["column_overrides"] == {"phone": ["Phone 1"], "email": []}


def test_header_row_is_found_below_title_text(conn, small, tmp_path):
    ids, res, data = _batch(conn, small, 2)
    path = reply_file(tmp_path / "combined.xlsx", data, title_rows=4)      # title text above the headers
    assert R.preview(conn, path, "combined.xlsx", settings=small)["header_row"] >= 3
    assert R.apply_reply(conn, path, "combined.xlsx", settings=small)["returned"] == 2


# --- finishing, late files, cancel, retry -------------------------------------------------------------------

def test_import_finishes_the_batch_and_the_missing_go_to_the_drill(conn, small, tmp_path):
    ids, res, data = _batch(conn, small, 6)
    out = R.apply_reply(conn, reply_file(tmp_path / "r.xlsx", data, skip={5, 6}, not_matched={4}), "r.xlsx",
                        settings=small)
    assert out["headline"] == "Crexi returned 3 of 6. 3 missing."
    missed = [int(data[i]["Property ID"]) for i in (3, 4, 5)]
    items = rows(conn, f"SELECT status FROM crexi_batch_items WHERE property_id IN ({','.join('?' * 3)})", missed)
    assert {i["status"] for i in items} == {"not_returned"}                    # missing: Crexi didn't have it
    att = one(conn, "SELECT * FROM drill_attempts WHERE property_id=? AND stage_key=?", (missed[2], E.CREDIT_STAGE))
    assert att["outcome"] == "nothing" and "didn't have it" in att["detail"] and "left it out" in att["detail"]
    att4 = one(conn, "SELECT * FROM drill_attempts WHERE property_id=? AND stage_key=?", (missed[0], E.CREDIT_STAGE))
    assert "Not Matched" in att4["detail"]
    assert scalar(conn, "SELECT COUNT(*) FROM credit_ledger WHERE property_id=?", (missed[2],)) == 0   # no credit
    b = B.get_batch(conn, res["batch_id"])
    assert b["status"] == "closed" and b["closed_by"] == "auto"
    # the DRILL moves those owners on: the Crexi credit stage is no longer their next step
    for pid in missed:
        assert E.property_drill(conn, pid)["next_step"].get("stage_key") != E.CREDIT_STAGE
    # the missing list: address, owner, why
    text = (small.path("output") / out["missing_file"]).read_text(encoding="utf-8-sig")
    lines = text.strip().splitlines()
    assert lines[0] == "Batch,Address,City,State,ZIP,County,APN,Owner,Why missing,Property ID" and len(lines) == 4
    assert "Not in Crexi's file" in text and "Crexi answered 'Not Matched'" in text
    assert data[5]["Address"] in text
    assert B.missing_file(conn, res["batch_id"], small).read_text(encoding="utf-8-sig") == text
    assert "missing: Crexi didn't have it" in B.property_history(conn, missed[2])["text"]


def test_a_late_reply_for_a_finished_batch_is_recorded_and_charged_then(conn, small, tmp_path):
    ids, res, data = _batch(conn, small, 4)
    R.apply_reply(conn, reply_file(tmp_path / "first.xlsx", data, skip={3, 4}), "first.xlsx", settings=small)
    assert C.credit_status(conn, today=TODAY, settings=small)["returned"] == 2
    path = reply_file(tmp_path / "late.xlsx", data, skip={1, 2, 4})          # row 3 comes back late
    pv = R.preview(conn, path, "late.xlsx", settings=small)
    assert any("already finished" in w for w in pv["warnings"])
    out = R.apply_reply(conn, path, "late.xlsx", settings=small)
    assert out["returned"] == 1 and out["headline"] == "Crexi returned 3 of 4. 1 missing."
    pid = int(data[2]["Property ID"])
    att = rows(conn, "SELECT outcome FROM drill_attempts WHERE property_id=? AND stage_key=? ORDER BY id",
               (pid, E.CREDIT_STAGE))
    assert [a["outcome"] for a in att] == ["nothing", "found"]                  # the DRILL result is upgraded
    assert one(conn, "SELECT status FROM crexi_batch_items WHERE property_id=?", (pid,))["status"] == "returned"
    assert C.credit_status(conn, today=TODAY, settings=small)["returned"] == 3  # the credit is charged now
    assert B.batch_counts(conn, res["batch_id"])["missing"] == 1


def test_withdraw_a_batch_never_uploaded_frees_the_credits(conn, small):
    seed(conn, 5)
    res = B.export_batch(conn, 5, today=TODAY, settings=small)
    assert C.credit_status(conn, today=TODAY, settings=small)["left"] == 15
    assert B.withdraw_batch(conn, res["batch_id"], settings=small)["needs_confirm"] is True
    B.withdraw_batch(conn, res["batch_id"], confirm=True, settings=small)
    assert C.credit_status(conn, today=TODAY, settings=small)["left"] == 20
    assert scalar(conn, "SELECT COUNT(*) FROM drill_attempts WHERE stage_key=?", (E.CREDIT_STAGE,)) == 0
    again = B.export_batch(conn, 5, today=TODAY, settings=small)                # the same five are available again
    assert again["kind"] == "first"


def test_cannot_cancel_a_batch_that_has_rows_back(conn, small, tmp_path):
    ids, res, data = _batch(conn, small, 3)
    R.apply_reply(conn, reply_file(tmp_path / "r.xlsx", data, skip={3}), "r.xlsx", settings=small)
    with pytest.raises(B.CrexiError, match="can't be cancelled"):
        B.withdraw_batch(conn, res["batch_id"], confirm=True, settings=small)


def test_a_batch_partly_back_from_before_is_finished_once(conn, small, tmp_path):
    ids, res, data = _batch(conn, small, 3)
    conn.execute("UPDATE crexi_batches SET status='imported' WHERE id=?", (res["batch_id"],))   # the old 'Partly back'
    conn.commit()
    B.status(conn, today=TODAY, settings=small)
    assert B.get_batch(conn, res["batch_id"])["status"] == "closed"
    assert B.batch_counts(conn, res["batch_id"])["missing"] == 3


def test_not_returned_are_never_resent_in_the_same_period_unless_asked(conn, small, tmp_path):
    ids, res, data = _batch(conn, small, 4)
    R.apply_reply(conn, reply_file(tmp_path / "r.xlsx", data, not_matched={3, 4}), "r.xlsx", settings=small)
    missed = {int(data[2]["Property ID"]), int(data[3]["Property ID"])}
    seed(conn, 3, top=0.2)
    plain = B.export_batch(conn, 3, today=TODAY, settings=small)
    assert not (missed & {int(r["Property ID"]) for r in read_export(small, plain["file"])})
    retry = B.export_batch(conn, 2, today=TODAY, settings=small, retry=True)
    got = read_export(small, retry["file"])
    assert {int(r["Property ID"]) for r in got} == missed
    items = rows(conn, "SELECT is_retry FROM crexi_batch_items WHERE batch_id=?", (retry["batch_id"],))
    assert [i["is_retry"] for i in items] == [1, 1]
    hist = B.history(conn, settings=small)["periods"][0]
    assert hist["second_passes"] == 2


def test_hand_batch_records_rows_that_were_in_no_batch(conn, small, tmp_path):
    ids, res, data = _batch(conn, small, 2)
    stray = add_prop(conn)
    score(conn, stray, 0.1)
    row = {"Property Link": conn.execute("SELECT crexi_url FROM properties WHERE id=?", (stray,)).fetchone()[0],
           "Address": "1 Stray St", "City": "Springfield", "State": "OH", "Zip Code": "45503", "County": "Clark",
           "APN": conn.execute("SELECT apn FROM properties WHERE id=?", (stray,)).fetchone()[0], "Row": "3"}
    path = reply_file(tmp_path / "r.xlsx", data + [row])
    pv = R.preview(conn, path, "r.xlsx", settings=small)
    assert pv["summary"]["counts"]["no_batch"] == 1 and "hand batch" in pv["message"]
    no = R.apply_reply(conn, path, "r.xlsx", settings=small)
    assert no["unmatched"] == 1 and scalar(conn, "SELECT COUNT(*) FROM crexi_batches WHERE kind='manual'") == 0
    # now with the hand-batch choice (nothing in the first file was applied twice)
    path2 = reply_file(tmp_path / "r2.xlsx", [row])
    done = R.apply_reply(conn, path2, "r2.xlsx", adopt_hand=True, settings=small)
    assert done["returned"] == 1 and scalar(conn, "SELECT COUNT(*) FROM crexi_batches WHERE kind='manual'") == 1


# --- history ------------------------------------------------------------------------------------------------

def test_history_totals_and_property_views(conn, small, tmp_path):
    ids, res, data = _batch(conn, small, 10)
    R.apply_reply(conn, reply_file(tmp_path / "r.xlsx", data, skip={9, 10}, not_matched={8}, phones=True),
                  "r.xlsx", actor="seth@example.test", settings=small)
    h = B.history(conn, settings=small)
    cur = h["periods"][0]
    assert cur["period"] == P.period_key(None, small) or cur["period"] == "2026-10-29"
    p = next(x for x in h["periods"] if x["period"] == "2026-10-29")
    assert (p["sent"], p["returned"], p["missing"], p["waiting"], p["found"]) == (10, 7, 3, 0, 7)
    assert p["matched_pct"] == 0.7 and p["credits_used"] == 7 and len(p["batches"]) == 1
    b = p["batches"][0]
    assert b["created_by"] == "cli" and b["kind"] == "first" and b["has_file"] and b["matched_pct"] == 0.7
    assert b["missing"] == 3 and b["imports"][0]["summary"]["headline"] == "Crexi returned 7 of 10. 3 missing."
    assert b["imports"][0]["created_by"] == "seth@example.test"
    assert p["enrichment"]["new_phones"] == 7
    assert (b["imports"][0]["summary"]["enrichment"]["rows_read"]) == 8
    # per-property words and the consistency of 'credits left' with the status the rest of the app uses
    pid_ok, pid_missed = int(data[0]["Property ID"]), int(data[9]["Property ID"])
    assert B.property_history(conn, pid_ok)["state"] == "data"
    assert B.property_history(conn, pid_missed)["state"] == "none"
    assert "Crexi didn't have it" in B.property_history(conn, pid_missed)["text"]
    st = C.credit_status(conn, today=TODAY, settings=small)
    assert st["used"] == st["returned"] + st["in_flight"] == 7 and st["left"] == 13
    assert st["found"] == 7 and st["nothing"] == 0
    plan_ids = {x["property_id"] for x in C.account_plan(conn, today=TODAY, settings=small)}
    assert not (plan_ids & {int(r["Property ID"]) for r in data})              # DRILL never plans a sent property


def test_hand_used_credits_count_in_the_same_period(conn, small):
    ids = seed(conn, 3)
    C.use_credit(conn, ids[0], "nothing", today=TODAY, settings=small)
    st = C.credit_status(conn, today=TODAY, settings=small)
    assert (st["returned"], st["left"]) == (1, 19)
    h = B.history(conn, settings=small)
    assert next(x for x in h["periods"] if x["period"] == "2026-10-29")["by_hand"] == 1
    assert B.property_history(conn, ids[0])["events"][0]["kind"] == "hand"
    pick = B.preview_export(conn, 5, today=TODAY, settings=small)
    assert ids[0] not in [s["property_id"] for s in pick["sample"]]


# --- the merge rule ---------------------------------------------------------------------------------------

def test_merge_rule_fills_blanks_unions_contacts_and_keeps_differences_as_extra_values(conn, small, tmp_path):
    ids, res, data = _batch(conn, small, 2)
    pid = int(data[0]["Property ID"])
    conn.execute("UPDATE properties SET units=12, year_built=1975, owner_name='Original Owner LLC', "
                 "mailing_address='1 Old Mail St', mailing_city='Dayton', mailing_state='OH', mailing_zip='45402', "
                 "building_sqft=NULL WHERE id=?", (pid,))
    conn.commit()

    def extra(n, row):
        row[8] = "New Owner Holdings"
        row[9], row[10], row[11], row[12] = "77 New Mail Rd", "Columbus", "OH", "43004"
        return row

    path = reply_file(tmp_path / "r.xlsx", data, extra=extra)
    pv = R.preview(conn, path, "r.xlsx", settings=small)
    e = pv["enrichment"]
    assert e["new_phones"] == 2 and e["new_emails"] == 2 and e["alternates"] == 3            # owner + mailing on the first, owner on the second
    assert e["new_mailing"] == 1 and e["blanks_filled"] >= 1 and e["rows_read"] == 2 and e["matched"] == 2
    assert "second values kept next to the first" in pv["enrichment_text"]
    out = R.apply_reply(conn, path, "r.xlsx", settings=small)
    p = one(conn, "SELECT * FROM properties WHERE id=?", (pid,))
    assert p["owner_name"] == "Original Owner LLC" and p["mailing_address"] == "1 Old Mail St"   # the original is kept
    extras = M.extra_values(conn, pid)
    assert extras["owner_name"][0]["header"] == "Owner 2" and extras["owner_name"][0]["value"] == "New Owner Holdings"
    assert extras["mailing_address"][0]["header"] == "Mailing address 2"
    assert extras["owner_name"][0]["batch_id"] == res["batch_id"] and extras["owner_name"][0]["source"] == "Crexi"
    assert {a["status"] for a in rows(conn, "SELECT status FROM property_alternates")} == {"kept"}
    assert "crexi_differs" not in p["flags_json"]                       # a second value is normal, not a warning
    names = {x["name"] for x in rows(conn, "SELECT name FROM people WHERE property_id=?", (pid,))}
    assert {"New Owner Holdings", "Pat Example1"} <= names                                        # owner also a person
    mails = rows(conn, "SELECT value, status, source FROM contact_points WHERE property_id=? AND channel='mail'", (pid,))
    assert len(mails) == 1 and mails[0]["status"] == "untested" and mails[0]["source"] == "crexi"
    assert out["enrichment"]["alternates"] == e["alternates"] and out["enrichment"]["new_phones"] == 2
    # idempotent: nothing is stored a second time
    count = scalar(conn, "SELECT COUNT(*) FROM property_alternates")
    again = R.apply_reply(conn, path, "r.xlsx", settings=small)
    assert again["changed"] is False and scalar(conn, "SELECT COUNT(*) FROM property_alternates") == count
    # the report is stored with the import for the History screen
    assert one(conn, "SELECT summary_json FROM crexi_batch_imports")["summary_json"].count("alternates") >= 1


def test_extra_values_are_numbered_kept_once_and_never_overwrite(conn, small):
    pid = add_prop(conn, units=24, year_built=1975, owner_name="Smith Farms LLC")
    prop = one(conn, "SELECT * FROM properties WHERE id=?", (pid,))
    for owner, units in (("Jones Holdings", 26), ("Jones Holdings", 26), ("Third Owner Co", 24), ("smith farms llc", 28)):
        plan = M.plan_facts(prop, {"owner_name": owner, "units": units, "year_built": 1975},
                            ["owner_name", "units", "year_built"])
        M.apply_alternates(conn, pid, plan, import_id=None, batch_id=7)
    p = one(conn, "SELECT * FROM properties WHERE id=?", (pid,))
    assert (p["owner_name"], p["units"], p["year_built"]) == ("Smith Farms LLC", 24, 1975)   # the first stays
    ex = M.extra_values(conn, pid)
    assert [(x["header"], x["value"]) for x in ex["owner_name"]] == [("Owner 2", "Jones Holdings"),
                                                                     ("Owner 3", "Third Owner Co")]
    assert [(x["header"], x["value"]) for x in ex["units"]] == [("Units 2", 26), ("Units 3", 28)]
    assert "year_built" not in ex                                       # the same value is not repeated
    assert scalar(conn, "SELECT COUNT(*) FROM property_alternates WHERE property_id=?", (pid,)) == 4
    assert all(x["batch_id"] == 7 and x["day"] for v in ex.values() for x in v)
    assert "crexi_differs" not in p["flags_json"]
    # the next Crexi import keeps a correction made by hand (property_edits still wins)
    again = M.plan_facts({"units": 20}, {"units": 12}, ["units"], skip=["units"])
    assert not again.alternates


def test_old_open_alternates_become_kept_and_the_warning_flag_goes(tmp_path, settings):
    """A database from before: property_alternates with 'open' rows waiting for a click and the crexi_differs flag."""
    import sqlite3

    from bluemoon import db as D
    path = tmp_path / "old.sqlite"
    old = sqlite3.connect(path)
    old.executescript(D.SCHEMA.replace("'kept', 'used', 'dismissed'", "'open', 'used', 'dismissed'")
                      .replace("DEFAULT 'kept'", "DEFAULT 'open'").replace("  batch_id INTEGER,                -- crexi_batches.id when it came in Crexi's reply to a batch\n", ""))
    m = old.execute("INSERT INTO markets(slug, name, asset_class, geography, created_at) VALUES('x','x','apartments','x','t')").lastrowid
    pid = old.execute("INSERT INTO properties(market_id, property_key, units, flags_json, created_at, updated_at) "
                      "VALUES(?, 'k', 24, '[\"crexi_differs\", \"year_unknown\"]', 't', 't')", (m,)).lastrowid
    for status, value in (("open", "26"), ("dismissed", "30"), ("used", "24")):
        old.execute("INSERT INTO property_alternates(property_id, field, value, original, seen_at, status) "
                    "VALUES(?, 'units', ?, '20', '2026-10-01T00:00:00+00:00', ?)", (pid, value, status))
    old.commit()
    old.close()
    conn = sqlite3.connect(path)
    conn.row_factory = sqlite3.Row
    D.init_db(conn, settings)
    D.init_db(conn, settings)                                             # safe to run again
    st = [r[0] for r in conn.execute("SELECT status FROM property_alternates ORDER BY id")]
    assert st == ["kept", "dismissed", "used"]
    assert conn.execute("SELECT flags_json FROM properties WHERE id=?", (pid,)).fetchone()[0] == '["year_unknown"]'
    ex = M.extra_values(conn, pid)
    # 26 kept; 30 was ignored by hand back then; 24 was swapped in by hand, so the 20 it replaced is kept too
    assert [x["value"] for x in ex["units"]] == [26, 20]
    with pytest.raises(sqlite3.IntegrityError):
        conn.execute("INSERT INTO property_alternates(property_id, field, value, seen_at, status) "
                     "VALUES(?, 'units', '1', 't', 'open')", (pid,))
    conn.close()


def test_blank_original_is_filled_and_same_value_does_nothing():
    plan = M.plan_facts({"units": None, "year_built": 1975, "county": "CLARK"},
                        {"units": 9, "year_built": 1975, "county": "Clark"}, ["units", "year_built", "county"])
    assert plan.fills == {"units": 9} and not plan.alternates and plan.merged["units"] == 9
    assert M.filled_count(plan) == 1


def test_plain_crexi_reimport_keeps_the_original_and_saves_alternates(conn, settings, sample_csv, tmp_path):
    from bluemoon.ingest import import_file
    from tests.test_round_trip import by_address, read_csv, write_csv   # the same helpers the round-trip tests use
    import_file(conn, MARKET, sample_csv)
    header, data = read_csv(sample_csv)
    row = by_address(data, "455 EXAMPLE AVE")
    old_year = one(conn, "SELECT id, year_built, owner_name FROM properties WHERE address = '455 EXAMPLE AVE'")
    row["Year Built"] = "1986"
    row["Owner Name"] = "A Brand New Buyer LLC"
    newer = write_csv(tmp_path / "newer.csv", header, data)
    res = import_file(conn, MARKET, newer)
    after = one(conn, "SELECT year_built, owner_name FROM properties WHERE id=?", (old_year["id"],))
    assert after["year_built"] == old_year["year_built"] and after["owner_name"] == old_year["owner_name"]
    fields = set(M.extra_values(conn, old_year["id"]))
    assert {"owner_name"} <= fields and res["enrichment"]["alternates"] >= 1
    assert "A Brand New Buyer LLC" in {r["name"] for r in rows(conn, "SELECT name FROM people WHERE property_id=?",
                                                               (old_year["id"],))}
    assert "second value" in res["enrichment_text"]
    count = scalar(conn, "SELECT COUNT(*) FROM property_alternates")
    import_file(conn, MARKET, newer)
    assert scalar(conn, "SELECT COUNT(*) FROM property_alternates") == count         # idempotent


def test_duplicate_rows_inside_an_import_file_keep_the_first_and_alternates(conn, settings, sample_csv, tmp_path):
    from bluemoon.ingest import import_file
    from tests.test_round_trip import by_address, read_csv, write_csv
    header, data = read_csv(sample_csv)
    first = by_address(data, "455 EXAMPLE AVE")
    copy = dict(first)
    copy["Year Built"] = str(int(first["Year Built"]) + 5)
    res = import_file(conn, MARKET, write_csv(tmp_path / "dups.csv", header, data + [copy]))
    pid = one(conn, "SELECT id, year_built FROM properties WHERE address = '455 EXAMPLE AVE'")
    assert pid["year_built"] == int(first["Year Built"])
    assert M.extra_values(conn, pid["id"])["year_built"][0]["header"] == "Year built 2"
    assert res["enrichment"]["duplicates"] >= 1


# --- numbered export columns -------------------------------------------------------------------------------

def _csv(path: Path) -> tuple[list[str], list[dict]]:
    with path.open(encoding="utf-8-sig", newline="") as f:
        r = csv.reader(f)
        header = next(r)
        return header, [dict(zip(header, line)) for line in r]


def test_exports_have_numbered_columns_and_the_round_trip_ignores_them(conn, small, tmp_path):
    from bluemoon.exports import export_market
    from bluemoon.ingest import import_file
    small.raw["exports"] = {"max_numbered": 2}
    pid = add_prop(conn, units=24, owner_name="Smith Farms LLC", mailing_address="1 Old Mail St",
                   mailing_city="Dayton", mailing_state="OH", mailing_zip="45402")
    other = add_prop(conn, units=8)
    score(conn, pid, 0.9)
    score(conn, other, 0.8)
    prop = one(conn, "SELECT * FROM properties WHERE id=?", (pid,))
    for owner, units in (("Jones Holdings", 26), ("Third Owner Co", 28), ("Fourth Owner Co", 30)):
        M.apply_alternates(conn, pid, M.plan_facts(prop, {"owner_name": owner, "units": units,
                                                          "mailing_address": "9 New Rd", "mailing_city": "Xenia",
                                                          "mailing_state": "OH", "mailing_zip": "45385"},
                                                   ["owner_name", "units", *M.MAILING_FIELDS]), import_id=None)
    conn.commit()
    E.add_contact(conn, pid, "phone", "(937) 555-0101", source="crexi")
    E.add_contact(conn, pid, "phone", "(937) 555-0102", source="crexi")
    E.add_contact(conn, pid, "phone", "(937) 555-0103", source="crexi")       # past the cap of 2
    E.add_contact(conn, pid, "email", "a@example.test", source="crexi")
    E.add_contact(conn, pid, "email", "b@example.test", source="crexi")
    ids = {r["value"]: r["id"] for r in rows(conn, "SELECT id, value FROM contact_points WHERE property_id=?", (pid,))}
    E.set_contact_status(conn, ids["(937) 555-0102"], "verified")              # working first
    E.set_contact_status(conn, ids["(937) 555-0101"], "wrong")                  # never in a callable column
    E.set_contact_status(conn, ids["a@example.test"], "dnc")                    # never in a callable column
    out = tmp_path / "out"
    files = export_market(conn, MARKET, out, stamp_handoffs=False)
    header, data = _csv(files["ranked_list.csv"])
    i = header.index("Owner")
    assert header[i:i + 4] == ["Owner", "Owner 2", "Mailing address", "Mailing address 2"]   # capped at 2
    assert header[header.index("Units") + 1] == "Units 2" and "Units 3" not in header and "Owner 3" not in header
    j = header.index("Contact info")
    assert header[j + 1:j + 4] == ["Phone 1", "Phone 2", "Email 1"] and "Email 2" not in header
    row = next(r for r in data if r["Property ID"] == str(pid))
    assert (row["Owner"], row["Owner 2"], row["Units"], row["Units 2"]) == ("Smith Farms LLC", "Jones Holdings", "24", "26")
    assert row["Mailing address"].startswith("1 Old Mail St") and row["Mailing address 2"].startswith("9 New Rd")
    assert (row["Phone 1"], row["Phone 2"], row["Email 1"]) == ("(937) 555-0102", "(937) 555-0103", "b@example.test")
    blank = next(r for r in data if r["Property ID"] == str(other))
    assert blank["Owner 2"] == "" and blank["Phone 1"] == ""
    wl_header, _ = _csv(files["drill_worklist.csv"])
    assert "Owner 2" in wl_header and wl_header[wl_header.index("Voicemail") + 1] == "Phone 1"
    crm_header, _ = _csv(files["crm_handoff.csv"])
    assert {"phone_1", "phone_2", "email_1", "owner_name_2", "mailing_address_2", "units_2"} <= set(crm_header)
    assert {"confirmed_person", "other_contacts_untested", "new_in_this_export"} <= set(crm_header)
    wb = openpyxl.load_workbook(files["ranked_list.xlsx"], read_only=True)
    first = [c.value for c in next(wb["Ranked"].iter_rows(max_row=1))]
    assert "Owner 2" in first and "Phone 1" in first and "Units 2" in first
    wb.close()
    # bring the file back with the numbered columns changed: they are ignored, never read as edits
    edited = tmp_path / "edited.csv"
    for r in data:
        if r["Property ID"] == str(pid):
            r["Units 2"], r["Owner 2"], r["Phone 1"] = "99", "Somebody Else", "(937) 555-9999"
    with edited.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=header)
        w.writeheader()
        w.writerows(data)
    res = import_file(conn, MARKET, edited)
    assert res["provider"] == "bluemoon" and res["properties_changed"] == 0
    assert one(conn, "SELECT units, owner_name FROM properties WHERE id=?", (pid,)) == {"units": 24,
                                                                                         "owner_name": "Smith Farms LLC"}


def test_drill_box_counts_sent_returned_missing_and_where_the_missing_are(conn, small, tmp_path):
    assert B.drill_box(conn, MARKET)["text"] == "Nothing sent to Crexi yet."
    ids, res, data = _batch(conn, small, 5)
    R.apply_reply(conn, reply_file(tmp_path / "r.xlsx", data, skip={5}, not_matched={4}), "r.xlsx", settings=small)
    box = B.drill_box(conn, MARKET)
    assert (box["sent"], box["returned"], box["missing"], box["waiting"], box["penetration"]) == (5, 3, 2, 0, 0.6)
    assert box["missing_working"] == 2 and "2 of the 2 missing are being worked" in box["text"]
    assert B.drill_box(conn, OTHER)["sent"] == 0                                 # per market


def test_screens_have_no_close_step_and_no_use_this_one():
    static = Path(__file__).resolve().parent.parent / "bluemoon" / "web" / "static"
    credits = (static / "credits.js").read_text(encoding="utf-8")
    prop = (static / "property.js").read_text(encoding="utf-8")
    drill = (static / "drill.js").read_text(encoding="utf-8")
    assert "Close the batch" not in credits and "ask-close" not in credits and "/close" not in credits
    assert "files/missing" in credits and "Export the next" in credits and "r.headline" in credits
    # after an import or a cancel frees credits, the export box starts again at every free credit (no stale N)
    assert credits.count("S.count = '';") >= 3
    assert "Use this one" not in prop and "/alternates/" not in prop and "crexi_differs" not in prop
    assert "withExtras(d, 'units'" in prop and "withExtras(d, 'owner_name'" in prop
    assert "crexiBox(" in drill and "missing_working" in drill


# --- the web routes ---------------------------------------------------------------------------------------

@pytest.fixture
def web(conn, small):
    from bluemoon.web.app import app
    with TestClient(app) as c:
        yield c


def test_routes_export_import_and_history(web, conn, small, tmp_path):
    seed(conn, 8)
    st = web.get("/api/crexi/status").json()
    assert st["status"]["left"] == 20 and st["default_count"] == 20
    pv = web.post("/api/crexi/export/preview", json={"count": 5}).json()
    assert pv["count"] == 5 and scalar(conn, "SELECT COUNT(*) FROM crexi_batches") == 0
    made = web.post("/api/crexi/export", json={"count": 5}).json()
    bid = made["batch_id"]
    assert made["status"]["in_flight"] == 5 and made["batch"]["created_by"]
    dl = web.get(f"/api/crexi/batches/{bid}/file")
    assert dl.status_code == 200 and "Property Link" in dl.text and dl.headers["content-type"].startswith("text/csv")
    data = read_export(small, made["file"])
    reply = reply_file(tmp_path / "reply.xlsx", data, skip={5})
    with reply.open("rb") as f:
        pre = web.post("/api/crexi/import/preview", files={"file": ("reply.xlsx", f)}).json()
    assert pre["summary"]["counts"]["new"] == 4 and "can_apply" in pre
    with reply.open("rb") as f:
        done = web.post("/api/crexi/import", files={"file": ("reply.xlsx", f)}).json()
    assert done["returned"] == 4 and done["status"]["in_flight"] == 0 and done["status"]["left"] == 16
    assert done["headline"] == "Crexi returned 4 of 5. 1 missing." and done["missing_file"]
    with reply.open("rb") as f:
        again = web.post("/api/crexi/import", files={"file": ("reply.xlsx", f)}).json()
    assert again["changed"] is False
    assert web.post(f"/api/crexi/batches/{bid}/close", json={"confirm": True}).status_code in (404, 405)  # no close step
    miss = web.get(f"/api/crexi/imports/{done['import_id']}/files/missing")
    assert miss.status_code == 200 and data[4]["Address"] in miss.text and "Not in Crexi's file" in miss.text
    assert web.get(f"/api/crexi/batches/{bid}/missing").status_code == 200
    hist = web.get("/api/crexi/history").json()
    assert hist["periods"][0]["batches"][0]["returned"] == 4 and hist["periods"][0]["missing"] == 1
    box = web.get(f"/api/m/{MARKET}/drill").json()["crexi"]
    assert (box["sent"], box["returned"], box["missing"], box["penetration"]) == (5, 4, 1, 0.8)
    assert box["missing_working"] == 1 and "being worked by the DRILL's other sources" in box["text"]
    prop = int(data[0]["Property ID"])
    one_ = web.get(f"/api/m/{MARKET}/properties/{prop}/crexi").json()
    assert one_["state"] == "data" and one_["events"]
    detail = web.get(f"/api/m/{MARKET}/properties/{prop}").json()
    assert detail["crexi"]["state"] == "data"
    listed = web.get(f"/api/m/{MARKET}/properties", params={"view": "crexi_data"}).json()
    assert listed["total"] == 4 and listed["counts"]["crexi_none"] == 1 and listed["counts"]["crexi_waiting"] == 0
    if (unm := done.get("unmatched_file")):
        assert web.get(f"/api/crexi/imports/{done['import_id']}/files/unmatched").status_code == 200
    assert web.get(f"/api/crexi/imports/{done['import_id']}/files/reply").status_code == 200
    assert web.get(f"/api/crexi/imports/{done['import_id']}/files/..%2f..%2fetc").status_code == 404
    assert web.get("/api/crexi/batches/9999/file").status_code == 404
    assert web.post(f"/api/m/{MARKET}/properties/{prop}/alternates/9999", json={"action": "ignore"}).status_code \
        in (404, 405)                                                       # no more Use this one / Ignore


def test_upload_names_cannot_escape_and_bad_files_are_refused(web, conn, small, tmp_path):
    seed(conn, 3)
    r = web.post("/api/crexi/import", files={"file": ("../../evil.exe", b"x")})
    assert r.status_code == 400
    r = web.post("/api/crexi/import/preview", files={"file": ("../../../tmp/evil.csv", b"Address\n1 A St\n")})
    assert r.status_code in (200, 400)                                  # read as a file called evil.csv, never a path
    assert not Path("/tmp/evil.csv").exists()
    r = web.post("/api/crexi/export", json={"count": 9999})
    assert r.status_code == 400 and "credits are free" in r.json()["detail"]


def test_every_crexi_route_needs_the_sign_in_when_a_password_is_set(web, conn, small, monkeypatch):
    monkeypatch.setenv("APP_PASSWORD", "correct horse battery")
    seed(conn, 2)
    checks = [("get", "/api/crexi/status"), ("get", "/api/crexi/history"), ("get", "/api/crexi/batches/1"),
              ("post", "/api/crexi/export/preview"), ("post", "/api/crexi/export"), ("get", "/api/crexi/batches/1/file"),
              ("post", "/api/crexi/import/preview"), ("post", "/api/crexi/import"),
              ("post", "/api/crexi/batches/1/withdraw"), ("get", "/api/crexi/batches/1/missing"),
              ("get", "/api/crexi/imports/1/files/unmatched"), ("get", "/api/crexi/imports/1/files/missing"),
              ("get", f"/api/m/{MARKET}/properties/1/crexi")]
    for method, path in checks:
        r = getattr(web, method)(path, **({"json": {}} if method == "post" else {}))
        assert r.status_code == 401, path
    assert scalar(conn, "SELECT COUNT(*) FROM crexi_batches") == 0


# --- the command line -------------------------------------------------------------------------------------

def test_cli_status_export_dry_run_and_export(conn, small, capsys, tmp_path):
    import run
    seed(conn, 6)
    conn.close()
    assert run.main(["crexi", "status"]) == 0
    assert "free" in capsys.readouterr().out
    assert run.main(["crexi", "export", "--count", "3", "--dry-run"]) == 0
    out = capsys.readouterr().out
    assert "3 properties would go to Crexi" in out and "Nothing was written" in out
    assert run.main(["crexi", "export", "--count", "3"]) == 0
    assert "batch-1-3.csv" in capsys.readouterr().out
    assert run.main(["crexi", "export", "--count", "99"]) == 1
    assert "credits are free" in capsys.readouterr().out
    assert run.main(["crexi", "history"]) == 0
    assert run.main(["crexi", "cancel", "1"]) == 3                        # asks for --yes
    assert run.main(["crexi", "cancel", "1", "--yes"]) == 0
    with pytest.raises(SystemExit):
        run.main(["crexi", "close", "1"])                                 # there is no close step any more


# --- review fixes (credits) ---------------------------------------------------------------------------------

def test_two_exports_at_once_never_share_the_free_credits(conn, small):
    import threading
    from bluemoon.db import connect
    seed(conn, 40)
    barrier = threading.Barrier(2)
    out: list = []

    def run():
        c = connect(settings=small)
        try:
            barrier.wait()
            out.append(B.export_batch(c, 15, today=TODAY, settings=small))
        except B.CrexiError as exc:
            out.append(exc)
        finally:
            c.close()

    ts = [threading.Thread(target=run) for _ in range(2)]
    for t in ts:
        t.start()
    for t in ts:
        t.join()
    st = C.credit_status(conn, today=TODAY, settings=small)
    assert st["in_flight"] <= 20 and st["used"] <= 20
    assert sum(isinstance(o, dict) for o in out) == 1 and sum(isinstance(o, B.CrexiError) for o in out) == 1
    assert scalar(conn, "SELECT COUNT(*) - COUNT(DISTINCT property_id) FROM crexi_batch_items") == 0


def test_export_rechecks_inside_the_write_lock(conn, small, monkeypatch):
    """A second program exports between this one's pick and its save: nothing is saved twice."""
    from bluemoon.db import connect
    seed(conn, 10)
    real_pick = B.pick
    other = connect(settings=small)

    done: list = []

    def pick_then_other_exports(*a, **k):
        pk = real_pick(*a, **k)
        if not done:
            done.append(1)
            B._export_batch(other, 10, created_by="other", today=TODAY, settings=small, retry=False, note=None)
        return pk

    monkeypatch.setattr(B, "pick", pick_then_other_exports)
    with pytest.raises(B.CrexiError, match="Nothing was saved"):
        B.export_batch(conn, 10, today=TODAY, settings=small)
    other.close()
    assert scalar(conn, "SELECT COUNT(*) FROM crexi_batches") == 1
    assert C.credit_status(conn, today=TODAY, settings=small)["in_flight"] == 10


def test_two_imports_of_the_same_reply_charge_each_row_once(conn, small, tmp_path, monkeypatch):
    ids, res, data = _batch(conn, small, 5)
    path = reply_file(tmp_path / "r.xlsx", data)
    table = R.read_reply(path)
    column_map, _ = R.map_columns(table.headers, R.ing._merged_overrides(small, None))
    cols = R.detect_columns(table.headers, table.rows, {}, skip=set(column_map.values()))
    stale = R.analyze(conn, table, column_map, cols)          # what a second, parallel import saw before writing
    assert R.apply_reply(conn, path, "r.xlsx", settings=small)["returned"] == 5
    monkeypatch.setattr(R, "analyze", lambda *a, **k: stale)
    R.apply_reply(conn, path, "r.xlsx", settings=small)
    assert scalar(conn, "SELECT COUNT(*) FROM credit_ledger") == 5
    st = C.credit_status(conn, today=TODAY, settings=small)
    assert (st["returned"], st["left"]) == (5, 15)


def test_hand_credit_on_a_property_waiting_at_crexi_is_refused(conn, small, tmp_path):
    ids, res, data = _batch(conn, small, 3)
    pid = int(data[0]["Property ID"])
    with pytest.raises(ValueError, match="waiting at Crexi in batch #%d" % res["batch_id"]):
        C.use_credit(conn, pid, None, today=TODAY, settings=small)
    with pytest.raises(ValueError, match="waiting at Crexi"):
        E.record_attempt(conn, pid, E.CREDIT_STAGE, "found", settings=small)
    assert scalar(conn, "SELECT COUNT(*) FROM credit_ledger") == 0
    # a pending hand credit made before this rule is adopted by the batch instead of counted twice
    conn.execute("INSERT INTO credit_ledger(property_id, month, result, created_at, period) VALUES(?,?,?,?,?)",
                 (pid, "2026-10", None, now_iso(), "2026-10-29"))
    conn.commit()
    R.apply_reply(conn, reply_file(tmp_path / "r.xlsx", data), "r.xlsx", settings=small)
    led = rows(conn, "SELECT batch_id, result FROM credit_ledger WHERE property_id=?", (pid,))
    assert led == [{"batch_id": res["batch_id"], "result": "found"}]
    st = C.credit_status(conn, today=TODAY, settings=small)
    assert (st["returned"], st["pending"]) == (3, 0)


def test_a_copy_of_a_property_sent_in_an_earlier_period_is_not_sent_again(conn, small):
    a = add_prop(conn, apn="OLD-1", address="1 Old St")
    score(conn, a, 0.9)
    B.export_batch(conn, 1, today=date(2026, 11, 27), settings=small)       # the Oct 29 period
    b = add_prop(conn, apn="OLD-1", address="2 Other St")                    # same parcel, a second record
    c = add_prop(conn)
    score(conn, b, 0.95)
    score(conn, c, 0.5)
    pk = B.pick(conn, 5, today=date(2026, 11, 30), settings=small)          # the Nov 29 period
    got = [it["id"] for it in pk.items]
    assert b not in got and c in got and pk.duplicates_left_out == 1


def test_a_batch_waiting_from_an_earlier_period_does_not_hold_this_periods_credits(conn, small):
    """credits-6, adapted to the import-finishes design (no close step): an older waiting batch gets a note and
    every credit of the new period is suggested."""
    seed(conn, 8)
    old = B.export_batch(conn, 5, today=date(2026, 11, 27), settings=small)
    tip = B.status(conn, today=date(2026, 11, 30), settings=small)["topup"]
    assert tip["suggested"] == 20 and "close" not in tip["message"].lower()
    assert "Batch #%d" % old["batch_id"] in tip["message"]
    assert "does not use this period's credits" in tip["message"]
    same = B.status(conn, today=date(2026, 11, 27), settings=small)["topup"]
    assert same["suggested"] == 15 and "earlier credit period" not in same["message"]


def test_reimporting_an_old_reply_does_not_flag_a_second_pass(conn, small, tmp_path):
    ids, res, data = _batch(conn, small, 4)
    path = reply_file(tmp_path / "r.xlsx", data, not_matched={3, 4})
    R.apply_reply(conn, path, "r.xlsx", settings=small)                 # finishes the batch: 3 and 4 missing
    retry = B.export_batch(conn, 2, today=TODAY, settings=small, retry=True)
    again = R.apply_reply(conn, path, "r.xlsx", settings=small)
    assert again["changed"] is False
    items = rows(conn, "SELECT status, not_matched_at FROM crexi_batch_items WHERE batch_id=?", (retry["batch_id"],))
    assert [(i["status"], i["not_matched_at"]) for i in items] == [("sent", None), ("sent", None)]
    assert scalar(conn, "SELECT COUNT(*) FROM crexi_batch_imports") == 1


def test_reimporting_an_old_reply_never_charges_or_finishes_a_newer_batch(conn, small, tmp_path):
    """No close step: a re-imported old file must not claim a property sent after it was first imported."""
    seed(conn, 3)
    first = B.export_batch(conn, 3, today=TODAY, settings=small)
    data = read_export(small, first["file"])
    B.withdraw_batch(conn, first["batch_id"], confirm=True, settings=small)
    B.export_batch(conn, 2, today=TODAY, settings=small)                  # the first two go to Crexi
    path = reply_file(tmp_path / "r.xlsx", data)                          # the file also has a row for the third
    out = R.apply_reply(conn, path, "r.xlsx", settings=small)
    assert out["returned"] == 2
    third = B.export_batch(conn, 1, today=TODAY, settings=small)
    again = R.apply_reply(conn, path, "r.xlsx", settings=small)
    assert again["changed"] is False
    items = rows(conn, "SELECT status FROM crexi_batch_items WHERE batch_id=?", (third["batch_id"],))
    assert [i["status"] for i in items] == ["sent"]
    assert one(conn, "SELECT status FROM crexi_batches WHERE id=?", (third["batch_id"],))["status"] != "closed"
    assert scalar(conn, "SELECT COUNT(*) FROM credit_ledger") == 2


# --- matching review fixes ------------------------------------------------------------------------------------

def test_a_second_row_for_a_retried_property_is_a_duplicate_not_a_second_credit(conn, small, tmp_path):
    ids, res, data = _batch(conn, small, 2)
    R.apply_reply(conn, reply_file(tmp_path / "nm.xlsx", data, not_matched={1, 2}), "nm.xlsx",
                  settings=small)                                               # both 'Not Matched': missing
    retry = B.export_batch(conn, 1, today=TODAY, settings=small, retry=True)
    sent = read_export(small, retry["file"])
    pid = int(sent[0]["Property ID"])
    path = reply_file(tmp_path / "r.xlsx", sent + [dict(sent[0], Row="2")])  # Crexi lists the row twice
    out = R.apply_reply(conn, path, "r.xlsx", settings=small)
    assert out["returned"] == 1 and out["duplicates"] == 1
    assert [(r["batch_id"]) for r in rows(conn, "SELECT batch_id FROM credit_ledger WHERE property_id=?", (pid,))] \
        == [retry["batch_id"]]                                                  # one credit, on the retry batch
    old = one(conn, "SELECT status FROM crexi_batch_items WHERE batch_id=? AND property_id=?", (res["batch_id"], pid))
    assert old["status"] == "not_returned"
    assert not any("already closed" in w for w in out["warnings"])


def test_two_properties_on_one_parcel_each_get_their_own_row(conn, small, tmp_path):
    a = add_prop(conn, crexi_url=None, apn="P-77")
    b = add_prop(conn, crexi_url=None, apn="TEMP-1")                # a different parcel while it is exported
    score(conn, a, 0.9)
    one_ = B.export_batch(conn, 1, today=TODAY, settings=small)
    score(conn, b, 0.8)
    two = B.export_batch(conn, 1, today=TODAY, settings=small)
    conn.execute("UPDATE properties SET apn='P-77' WHERE id=?", (b,))
    conn.commit()
    ra, rb = read_export(small, one_["file"])[0], read_export(small, two["file"])[0]
    rows_ = [dict(rb, Row="2", APN="P-77"), dict(ra, Row="1", APN="P-77")]     # B's row comes first
    out = R.apply_reply(conn, reply_file(tmp_path / "r.xlsx", rows_), "r.xlsx", settings=small)
    assert out["returned"] == 2

    def phones(pid):
        return {r["value"] for r in rows(conn, "SELECT value FROM contact_points WHERE property_id=? AND channel='phone'",
                                         (pid,))}
    assert phones(a) == {"(937) 555-1001"} and phones(b) == {"(937) 555-1002"}


def test_reply_facts_merge_into_the_matched_property_and_never_add_one(conn, small, tmp_path):
    pid = add_prop(conn, address="300 Oak St", unit="A", apn=None, crexi_url=None, source="list", owner_name=None)
    score(conn, pid, 0.9)
    res = B.export_batch(conn, 1, today=TODAY, settings=small)
    data = read_export(small, res["file"])

    def extra(n, row):
        row[0] = "https://www.crexi.com/property-records/oak/abcd1234"          # a link we have not stored
        row[7] = "P-300"
        return row

    before = scalar(conn, "SELECT COUNT(*) FROM properties")
    out = R.apply_reply(conn, reply_file(tmp_path / "r.xlsx", data, extra=extra), "r.xlsx", settings=small)
    assert out["returned"] == 1 and out["merge"]["new_rows"] == 0
    assert scalar(conn, "SELECT COUNT(*) FROM properties") == before                  # never a second '300 OAK ST'
    p = one(conn, "SELECT * FROM properties WHERE id=?", (pid,))
    assert p["owner_name"] == "Test Owner 1" and p["mailing_address"] == "1 Mail St"  # blanks filled
    assert p["unit"] == "A" and p["source"] == "list" and p["crexi_url"].endswith("abcd1234")


def test_owner_and_mailing_address_are_kept_for_a_matched_row_with_no_link(conn, small, tmp_path):
    ids, res, data = _batch(conn, small, 2)
    pid = int(data[0]["Property ID"])
    conn.execute("UPDATE properties SET owner_name=NULL WHERE id=?", (pid,))
    conn.commit()

    def extra(n, row):
        if n == 1:
            row[0] = None                                                            # matched by APN
            row[13:24] = [None] * 11                                                 # no contact, phones, emails
        return row

    path = reply_file(tmp_path / "r.xlsx", data, extra=extra)
    pv = R.preview(conn, path, "r.xlsx", settings=small)
    out = R.apply_reply(conn, path, "r.xlsx", settings=small)
    assert out["owner_only"] == 1
    p = one(conn, "SELECT owner_name, mailing_address FROM properties WHERE id=?", (pid,))
    assert p["owner_name"] == "Test Owner 1" and p["mailing_address"] == "1 Mail St"
    assert out["enrichment"]["blanks_filled"] == pv["enrichment"]["blanks_filled"]


def test_a_newer_value_of_a_column_kept_only_with_the_originals_is_kept_as_a_second_value(conn, settings, sample_csv, tmp_path):
    from bluemoon.ingest import import_file
    from tests.test_round_trip import by_address, read_csv, write_csv
    header, data = read_csv(sample_csv)
    row = by_address(data, "455 EXAMPLE AVE")
    row["Days on Market"], row["Occupancy"] = "10", "80"
    import_file(conn, MARKET, write_csv(tmp_path / "first.csv", header, data))
    pid = one(conn, "SELECT id FROM properties WHERE address = '455 EXAMPLE AVE'")["id"]
    row["Days on Market"], row["Occupancy"] = "95", "60"
    newer = write_csv(tmp_path / "newer.csv", header, data)
    res = import_file(conn, MARKET, newer)
    p = one(conn, "SELECT raw_json, flags_json FROM properties WHERE id=?", (pid,))
    raw = __import__("json").loads(p["raw_json"])
    assert (raw["Days on Market"], raw["Occupancy"]) == ("10", "80")                 # the original is kept
    ex = M.extra_values(conn, pid)
    assert [(x["header"], x["value"]) for x in ex["raw:Days on Market"]] == [("Days on Market 2", "95")]
    assert [(x["header"], x["value"]) for x in ex["raw:Occupancy"]] == [("Occupancy 2", "60")]
    assert "crexi_differs" not in p["flags_json"]                                     # a second value is normal
    assert res["enrichment"]["alternates"] >= 2
    count = scalar(conn, "SELECT COUNT(*) FROM property_alternates")
    import_file(conn, MARKET, newer)
    assert scalar(conn, "SELECT COUNT(*) FROM property_alternates") == count         # idempotent


def test_a_value_swapped_in_by_hand_before_oct_2026_is_not_kept_twice(conn, settings, sample_csv, tmp_path):
    """A legacy 'Use this one' row (status 'used'): the value it replaced already shows as Owner 2, so the old
    file imported again adds no new row and Owner 2 is not repeated."""
    from bluemoon.ingest import import_file
    from tests.test_round_trip import by_address, read_csv, write_csv
    import_file(conn, MARKET, sample_csv)
    pid = one(conn, "SELECT id, owner_name FROM properties WHERE address = '455 EXAMPLE AVE'")
    old_owner, pid = pid["owner_name"], pid["id"]
    header, data = read_csv(sample_csv)
    by_address(data, "455 EXAMPLE AVE")["Owner Name"] = "A Brand New Buyer LLC"
    import_file(conn, MARKET, write_csv(tmp_path / "newer.csv", header, data))
    alt = one(conn, "SELECT id FROM property_alternates WHERE property_id=? AND field='owner_name'", (pid,))
    conn.execute("UPDATE properties SET owner_name='A Brand New Buyer LLC' WHERE id=?", (pid,))
    conn.execute("UPDATE property_alternates SET status='used', decided_by='seth', decided_at=? WHERE id=?",
                 (now_iso(), alt["id"]))
    conn.commit()
    count = scalar(conn, "SELECT COUNT(*) FROM property_alternates")
    import_file(conn, MARKET, sample_csv)                                             # the old listing file again
    assert scalar(conn, "SELECT COUNT(*) FROM property_alternates") == count
    owners = M.extra_values(conn, pid)["owner_name"]
    assert [(x["header"], x["value"]) for x in owners] == [("Owner 2", old_owner)]
    assert "crexi_differs" not in one(conn, "SELECT flags_json FROM properties WHERE id=?", (pid,))["flags_json"]


# --- final review fixes (crexiflow) ---------------------------------------------------------------------------

def _resave(src: Path, dst: Path) -> Path:
    """The same reply opened and saved again (Excel, or a second download): same rows, different bytes."""
    wb = openpyxl.load_workbook(src)
    wb.active["A1"] = "saved again"
    wb.save(dst)
    return dst


def test_a_resaved_copy_of_an_old_reply_never_finishes_the_second_pass(conn, small, tmp_path):
    ids, res, data = _batch(conn, small, 6)
    path = reply_file(tmp_path / "r.xlsx", data, not_matched={3, 4})
    R.apply_reply(conn, path, "r.xlsx", settings=small)                    # 4 returned, 2 missing
    seed(conn, 2, top=0.1)
    retry = B.export_batch(conn, 4, today=TODAY, settings=small, retry=True)
    assert scalar(conn, "SELECT COUNT(*) FROM crexi_batch_items WHERE batch_id=?", (retry["batch_id"],)) == 4
    before = C.credit_status(conn, today=TODAY, settings=small)
    copy = _resave(path, tmp_path / "r copy.xlsx")
    assert R.first_imported(conn, copy) is None                              # the bytes differ
    again = R.apply_reply(conn, copy, "r copy.xlsx", settings=small)
    assert again["changed"] is False
    items = rows(conn, "SELECT status, not_matched_at FROM crexi_batch_items WHERE batch_id=?", (retry["batch_id"],))
    assert [(i["status"], i["not_matched_at"]) for i in items] == [("sent", None)] * 4
    assert B.get_batch(conn, retry["batch_id"])["status"] != "closed"
    after = C.credit_status(conn, today=TODAY, settings=small)
    assert (after["in_flight"], after["left"]) == (before["in_flight"], before["left"]) == (4, 12)


def test_a_late_row_from_the_old_reply_is_charged_to_the_old_batch_not_the_second_pass(conn, small, tmp_path):
    ids, res, data = _batch(conn, small, 6)
    R.apply_reply(conn, reply_file(tmp_path / "r.xlsx", data, skip={3}), "r.xlsx", settings=small)
    retry = B.export_batch(conn, 1, today=TODAY, settings=small, retry=True)
    pid = int(data[2]["Property ID"])
    late = reply_file(tmp_path / "late.xlsx", data)                          # Crexi's full file, row 3 matched now
    out = R.apply_reply(conn, late, "late.xlsx", settings=small)
    assert out["returned"] == 1 and out["batch_ids"] == [res["batch_id"]]
    assert one(conn, "SELECT status FROM crexi_batch_items WHERE batch_id=? AND property_id=?",
               (res["batch_id"], pid))["status"] == "returned"
    assert one(conn, "SELECT status FROM crexi_batch_items WHERE batch_id=? AND property_id=?",
               (retry["batch_id"], pid))["status"] == "sent"
    assert B.get_batch(conn, retry["batch_id"])["status"] != "closed"


def test_the_real_reply_to_a_second_pass_still_finishes_it(conn, small, tmp_path):
    ids, res, data = _batch(conn, small, 4)
    R.apply_reply(conn, reply_file(tmp_path / "r.xlsx", data, not_matched={3, 4}), "r.xlsx", settings=small)
    retry = B.export_batch(conn, 2, today=TODAY, settings=small, retry=True)
    sent = read_export(small, retry["file"])
    out = R.apply_reply(conn, reply_file(tmp_path / "r2.xlsx", sent, not_matched={2}), "r2.xlsx", settings=small)
    assert out["batch_ids"] == [retry["batch_id"]] and out["headline"] == "Crexi returned 1 of 2. 1 missing."
    assert B.get_batch(conn, retry["batch_id"])["status"] == "closed"


def test_streets_with_no_house_number_are_not_duplicates(conn, small, tmp_path):
    a = add_prop(conn, address="Channingway Dr", city="FAIRBORN", zip="45324", apn="CH-1")
    b = add_prop(conn, address="Channingway Dr", city="FAIRBORN", zip="45324", apn="CH-2")
    c = add_prop(conn, address="12 Main St", city="FAIRBORN", zip="45324", apn="M-1")
    d = add_prop(conn, address="12 Main St.", city="FAIRBORN", zip="45324", apn="M-2", crexi_url=None)
    for pid, v in ((a, 0.9), (b, 0.8), (c, 0.7), (d, 0.6)):
        score(conn, pid, v)
    assert B.address_key({"address": "Channingway Dr", "city": "Fairborn", "zip": "45324"}) is None
    picked = B.pick(conn, 4, today=TODAY, settings=small)
    assert {p["id"] for p in picked.items} == {a, b, c}                      # 12 Main St twice is still one
    assert picked.duplicates_left_out == 1
    res = B.export_batch(conn, 3, today=TODAY, settings=small)
    data = read_export(small, res["file"])
    for r in data:                                                            # Crexi answers by address only
        r["Property Link"], r["APN"] = "", ""
    out = R.apply_reply(conn, reply_file(tmp_path / "r.xlsx", data), "r.xlsx", settings=small)
    assert out["returned"] == 1                                               # only 12 Main St matches one property
    assert {r["property_id"] for r in rows(conn, "SELECT property_id FROM crexi_batch_items WHERE status='returned'")} \
        == {c}


def test_file_row_is_the_row_in_crexis_spreadsheet(conn, small, tmp_path):
    ids, res, data = _batch(conn, small, 3)
    path = reply_file(tmp_path / "r.xlsx", data[:2])          # blank row 1, note row 2, headers row 3, data rows 4-5
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    ws.append([])                                                             # a blank row 6
    ws.append(["https://www.crexi.com/property-records/x/nobody", "Retail", "1 STRANGER AVE", "OTHERTOWN", "OH",
               "44444", "LUCAS", "X-1", "Stranger Co"] + [None] * 15)          # row 7
    wb.save(path)
    out = R.apply_reply(conn, path, "r.xlsx", settings=small)
    with (small.path("output") / out["unmatched_file"]).open(encoding="utf-8-sig", newline="") as f:
        got = list(csv.DictReader(f))
    assert [r["File row"] for r in got] == ["7"]


def test_cli_import_ranks_the_markets_the_reply_changed(conn, small, tmp_path, monkeypatch, capsys):
    import argparse

    from bluemoon import pipeline
    from bluemoon.crexi import cli
    ids, res, data = _batch(conn, small, 2)
    path = reply_file(tmp_path / "r.xlsx", data)
    calls = []
    monkeypatch.setattr(pipeline, "after_import", lambda c, slug: calls.append(slug) or {"message": "Ranked."})
    monkeypatch.setattr(cli, "get_settings", lambda: small)
    monkeypatch.setattr(R, "get_settings", lambda: small)
    args = argparse.Namespace(action="import", target=str(path), count=None, retry=False, hand=False, yes=True,
                              dry_run=False, phone_columns=None, email_columns=None, contact_column=None)
    assert cli.run(args, conn) == cli.EXIT_OK
    assert calls == [MARKET]
    assert f"{MARKET}: Ranked." in capsys.readouterr().out
