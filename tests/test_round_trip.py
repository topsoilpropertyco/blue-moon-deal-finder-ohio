"""The round trip: download a Blue Moon export, fix it in Excel, import it back (bluemoon/ingest/roundtrip.py).

Uses the invented sample export (made-up owners and addresses); nothing touches the network.
"""
from __future__ import annotations

import csv
import json
import re
import zipfile
from datetime import datetime
from pathlib import Path

import httpx
import pytest
from openpyxl import load_workbook

from bluemoon import exports, pipeline
from bluemoon.db import loads, one, rows, scalar
from bluemoon.ingest import IngestError, detect_provider, import_file, preview_file, read_table
from bluemoon.ingest.columns import is_blue_moon_export, looks_like_blue_moon, other_blue_moon_export
from bluemoon.ingest.roundtrip import current_decision_text, decision_intent, parse_as_of, parse_cell
from bluemoon.scoring import overrides

APARTMENTS = "apartments-springfield-oh"


@pytest.fixture(autouse=True)
def no_network(monkeypatch):
    def refuse(self, request):
        raise AssertionError(f"Tried to reach the network: {request.url.host}")
    monkeypatch.setattr(httpx.HTTPTransport, "handle_request", refuse)


@pytest.fixture
def ranked(conn, sample_csv):
    res = pipeline.run_step(conn, APARTMENTS, "import", path=sample_csv)
    assert res["ok"], res["message"]
    return conn


# --- helpers ------------------------------------------------------------------------------------------

def read_csv(path: Path) -> tuple[list[str], list[dict]]:
    with path.open(encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        return list(reader.fieldnames or []), list(reader)


def write_csv(path: Path, header: list[str], data: list[dict]) -> Path:
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=header)
        w.writeheader()
        w.writerows(data)
    return path


def export_csv(conn) -> tuple[list[str], list[dict]]:
    return read_csv(exports.export_market(conn, APARTMENTS, only=["ranked_list.csv"])["ranked_list.csv"])


def by_address(data: list[dict], address: str) -> dict:
    return next(r for r in data if r["Address"] == address)


def pid_of(conn, address: str) -> int:
    return int(scalar(conn, "SELECT id FROM properties WHERE address = ?", (address,)))


def snapshot(conn) -> dict[int, dict]:
    return {r["id"]: r for r in rows(conn, "SELECT * FROM properties ORDER BY id")}


def sheet_cols(ws) -> dict[str, int]:
    """{header: column number} of a worksheet's first row."""
    return {c.value: i for i, c in enumerate(ws[1], start=1) if c.value}


def drop_columns(wb, header: str, sheets: list[str] | None = None) -> None:
    for ws in wb.worksheets:
        if (sheets is None or ws.title in sheets) and header in sheet_cols(ws):
            ws.delete_cols(sheet_cols(ws)[header])


def edited_csv(conn, tmp_path, edits: dict[str, dict[str, str]], name: str = "edited.csv") -> Path:
    """Export ranked_list.csv, change cells {address: {header: value}}, save it as another file."""
    header, data = export_csv(conn)
    for address, cells in edits.items():
        by_address(data, address).update(cells)
    return write_csv(tmp_path / name, header, data)


# --- knowing our own files ----------------------------------------------------------------------------

def test_our_exports_are_known_and_crexi_files_are_not(ranked, sample_csv):
    files = exports.export_market(ranked, APARTMENTS, only=["ranked_list.csv", "ranked_list.xlsx", "excluded"])
    for name in ("ranked_list.csv", "excluded_drive_time.csv"):
        header, _ = read_csv(files[name])
        assert "Row as of" in header and detect_provider(header) == "bluemoon", name
    assert detect_provider(read_csv(sample_csv)[0]) == "crexi"
    # Files exported before 'Row as of' existed: known by our other column names.
    older = ["Rank", "Blue Moon score %", "Why set aside", "Address", "Units", "Property ID", "Property Link",
             "Shared APN Flag", "PFC Indicator", "REO Sale Flag"]
    assert detect_provider(older) == "bluemoon"
    assert not is_blue_moon_export(["Blue Moon score %", "Why set aside", "Address"])   # no Property ID
    wb = load_workbook(files["ranked_list.xlsx"])
    assert wb.properties.creator == exports.WORKBOOK_CREATOR


def test_our_workbook_is_known_by_its_metadata_even_with_columns_deleted(ranked, tmp_path):
    pid = pid_of(ranked, "455 EXAMPLE AVE")
    cols = [exports.Col(h, lambda r, k=k: r[k]) for h, k in (("Address", "a"), ("Units", "u"), ("Property ID", "p"))]
    path = exports.write_xlsx(tmp_path / "trimmed.xlsx", [exports.Sheet("Mine", cols, [
        {"a": "455 EXAMPLE AVE", "u": 39, "p": pid}])])
    before = one(ranked, "SELECT source_id, raw_json FROM properties WHERE id = ?", (pid,))
    res = import_file(ranked, APARTMENTS, path)
    assert res["provider"] == "bluemoon" and res["fields_changed"] == {"units": 1}
    assert any("went by the time it was downloaded" in w for w in res["warnings"])
    assert one(ranked, "SELECT source_id, raw_json FROM properties WHERE id = ?", (pid,)) == before


def test_xlsx_reads_all_properties_not_just_the_ranked_sheet(ranked):
    path = exports.export_market(ranked, APARTMENTS, only=["ranked_list.xlsx"])["ranked_list.xlsx"]
    data = read_table(path)
    assert len(data) == scalar(ranked, "SELECT COUNT(*) FROM properties")   # 11, not the 6 on 'Ranked'


# --- only the editable columns, nothing else ------------------------------------------------------------

def test_csv_round_trip_changes_only_the_editable_columns(ranked, tmp_path):
    before = snapshot(ranked)
    target = pid_of(ranked, "455 EXAMPLE AVE")
    path = edited_csv(ranked, tmp_path, {"455 EXAMPLE AVE": {
        "Units": "40", "Owner": "SOMEONE ELSE LLC", "Address": "1 NEW ST", "Rank": "99", "Number of Units": "500",
        "Owner Name": "ANOTHER NAME", "Why this score": "made up", "Year Built": "1900"}})
    res = import_file(ranked, APARTMENTS, path, actor="seth@example.org")
    assert res["provider"] == "bluemoon" and res["rows_new"] == 0
    assert res["fields_changed"] == {"units": 1} and res["properties_changed"] == 1
    after = snapshot(ranked)
    assert set(after) == set(before)                                   # nothing added or deleted
    for pid, row in after.items():
        assert row["source_id"] == before[pid]["source_id"], pid      # our Property ID never becomes the source id
        assert row["raw_json"] == before[pid]["raw_json"], pid        # computed columns never land in raw_json
        assert row["owner_name"] == before[pid]["owner_name"] and row["address"] == before[pid]["address"]
        if pid != target:
            assert row == before[pid], pid
    assert after[target]["units"] == 40 and after[target]["year_built"] == before[target]["year_built"]
    edits = rows(ranked, "SELECT * FROM property_edits")
    assert len(edits) == 1
    assert edits[0]["property_id"] == target and edits[0]["field"] == "units"
    assert loads(edits[0]["old_value"]) == 36 and loads(edits[0]["new_value"]) == 40
    assert edits[0]["via"] == f"import:{res['import_id']}" and edits[0]["actor"] == "seth@example.org"
    assert scalar(ranked, "SELECT filename FROM imports WHERE id = ?", (res["import_id"],)) == "edited.csv"


def test_xlsx_round_trip_reads_every_property_sheet(ranked, tmp_path):
    before = snapshot(ranked)
    path = exports.export_market(ranked, APARTMENTS, only=["ranked_list.xlsx"])["ranked_list.xlsx"]
    wb = load_workbook(path)
    first = wb["Ranked"]
    head = [c.value for c in first[1]]
    ranked_pid = first.cell(2, head.index("Property ID") + 1).value
    first.cell(2, head.index("Year built") + 1).value = 1999
    aside = wb["Excluded"]
    head_x = [c.value for c in aside[1]]
    aside_pid = aside.cell(2, head_x.index("Property ID") + 1).value
    aside.cell(2, head_x.index("County") + 1).value = "madison county"
    aside.cell(2, head_x.index("Units") + 1).value = 25.0           # Excel numbers come back as floats
    wb.save(tmp_path / "edited.xlsx")
    res = import_file(ranked, APARTMENTS, tmp_path / "edited.xlsx")
    assert res["fields_changed"] == {"units": 1, "year_built": 1, "county": 1}, res
    after = snapshot(ranked)
    assert after[ranked_pid]["year_built"] == 1999
    assert after[aside_pid]["county"] == "MADISON" and after[aside_pid]["units"] == 25
    assert all(after[p]["raw_json"] == before[p]["raw_json"] and after[p]["source_id"] == before[p]["source_id"]
               for p in after)
    assert {r["actor"] for r in rows(ranked, "SELECT actor FROM property_edits")} == {"cli"}
    assert not any("has no 'Row as of' column" in w for w in res["warnings"])   # not about the DRILL worklist


def test_two_sheets_that_disagree_leave_the_value_alone(ranked, tmp_path):
    path = exports.export_market(ranked, APARTMENTS, only=["ranked_list.xlsx"])["ranked_list.xlsx"]
    wb = load_workbook(path)
    pid = None
    for name, value in (("Ranked", 30), ("All properties", 31)):
        ws = wb[name]
        head = [c.value for c in ws[1]]
        if pid is None:
            pid = ws.cell(2, head.index("Property ID") + 1).value
        row = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, head.index("Property ID") + 1).value == pid)
        ws.cell(row, head.index("Units") + 1).value = value
    wb.save(tmp_path / "edited.xlsx")
    units = scalar(ranked, "SELECT units FROM properties WHERE id = ?", (pid,))
    res = import_file(ranked, APARTMENTS, tmp_path / "edited.xlsx")
    assert res["properties_changed"] == 0 and res["conflicts"] == 1
    assert any("between two sheets" in w for w in res["warnings"])
    assert scalar(ranked, "SELECT units FROM properties WHERE id = ?", (pid,)) == units


def test_headers_and_data_do_not_snowball_over_two_round_trips(ranked, tmp_path):
    header1, data1 = export_csv(ranked)
    raw_before = {r["id"]: r["raw_json"] for r in rows(ranked, "SELECT id, raw_json FROM properties")}
    unchanged = write_csv(tmp_path / "same.csv", header1, data1)
    res = import_file(ranked, APARTMENTS, unchanged)
    assert res["properties_changed"] == 0 and res["rows_unchanged"] == 11
    assert "changes nothing" in preview_file(unchanged, market_slug=APARTMENTS, conn=ranked)["message"]
    path = edited_csv(ranked, tmp_path, {"1200 SAMPLE RIDGE DR": {"Units": "30"}}, "round1.csv")
    assert import_file(ranked, APARTMENTS, path)["properties_changed"] == 1
    header2, _ = export_csv(ranked)
    path = edited_csv(ranked, tmp_path, {"1200 SAMPLE RIDGE DR": {"Units": "31"}}, "round2.csv")
    assert import_file(ranked, APARTMENTS, path)["properties_changed"] == 1
    header3, _ = export_csv(ranked)
    assert header1 == header2 == header3
    assert not any(h.endswith("(Crexi) (Crexi)") for h in header3)
    assert {r["id"]: r["raw_json"] for r in rows(ranked, "SELECT id, raw_json FROM properties")} == raw_before
    assert scalar(ranked, "SELECT units FROM properties WHERE address = '1200 SAMPLE RIDGE DR'") == 31


def test_cells_that_cant_be_read_are_left_alone_and_empty_cells_clear(ranked, tmp_path):
    path = edited_csv(ranked, tmp_path, {"455 EXAMPLE AVE": {"Year built": "nineteen", "ZIP": "45x"},
                                         "77 FAKE MEADOW CT": {"Year built": ""}})
    res = import_file(ranked, APARTMENTS, path)
    assert res["fields_changed"] == {"year_built": 1}
    assert any("Year built 'nineteen'" in w for w in res["warnings"])
    gone = one(ranked, "SELECT year_built, flags_json FROM properties WHERE address = '77 FAKE MEADOW CT'")
    assert gone["year_built"] is None and "year_unknown" in loads(gone["flags_json"])
    kept = one(ranked, "SELECT year_built, zip FROM properties WHERE address = '455 EXAMPLE AVE'")
    assert kept["year_built"] == 1985 and kept["zip"] == "45505"


def test_unknown_ids_and_another_markets_file_change_nothing(ranked, tmp_path):
    header, data = export_csv(ranked)
    data[0]["Units"] = "99"
    data[1]["Property ID"] = "999999"
    path = write_csv(tmp_path / "edited.csv", header, data)
    before = snapshot(ranked)
    res = import_file(ranked, "land-springfield-oh", path)
    assert res["properties_changed"] == 0 and snapshot(ranked) == before
    assert "Apartments" in res["message"] and "switch to it" in res["message"]
    res = import_file(ranked, APARTMENTS, path)
    assert res["unknown_ids"] == 1 and res["properties_changed"] == 1
    assert any("not in the app" in w for w in res["warnings"])


def test_limit_reads_only_the_first_properties(ranked, tmp_path):
    header, data = export_csv(ranked)
    for r in data:
        r["Units"] = "7"
    path = write_csv(tmp_path / "edited.csv", header, data)
    res = import_file(ranked, APARTMENTS, path, limit=1)
    assert res["properties_changed"] == 1 and res["rows_total"] == 1


# --- stale rows ------------------------------------------------------------------------------------------

def test_rows_changed_in_the_app_after_the_download_are_left_alone(ranked, sample_csv, tmp_path):
    ranked.execute("UPDATE properties SET updated_at = '2000-01-01T00:00:00+00:00'")   # imported long ago
    ranked.commit()
    header, data = export_csv(ranked)
    assert {r["Row as of"] for r in data} == {"2000-01-01T00:00:00+00:00"}
    # The same Crexi file again changes nothing, so nothing goes stale.
    import_file(ranked, APARTMENTS, sample_csv)
    assert {r["updated_at"] for r in rows(ranked, "SELECT updated_at FROM properties")} == {"2000-01-01T00:00:00+00:00"}
    # A newer Crexi file changes one property after the download.
    crexi_header, crexi = read_csv(sample_csv)
    by_address(crexi, "455 EXAMPLE AVE")["Year Built"] = "1986"
    newer = write_csv(tmp_path / "crexi_newer.csv", crexi_header, crexi)
    import_file(ranked, APARTMENTS, newer)
    by_address(data, "455 EXAMPLE AVE")["Units"] = "50"         # an edit made on the old download
    by_address(data, "1200 SAMPLE RIDGE DR")["Units"] = "29"
    old_file = write_csv(tmp_path / "old_download.csv", header, data)
    preview = preview_file(old_file, market_slug=APARTMENTS, conn=ranked)
    assert preview["edits"]["skipped_stale"] == 1 and "after this file was downloaded" in preview["message"]
    res = import_file(ranked, APARTMENTS, old_file)
    assert res["skipped_stale"] == 1 and res["fields_changed"] == {"units": 1}
    assert scalar(ranked, "SELECT units FROM properties WHERE address = '455 EXAMPLE AVE'") == 36
    assert scalar(ranked, "SELECT year_built FROM properties WHERE address = '455 EXAMPLE AVE'") == 1986
    assert scalar(ranked, "SELECT units FROM properties WHERE address = '1200 SAMPLE RIDGE DR'") == 29


def test_row_as_of_survives_what_excel_does_to_it():
    assert parse_as_of("2026-10-21T14:03:22+00:00").isoformat() == "2026-10-21T14:03:22+00:00"
    assert parse_as_of("2026-10-21T14:03:22Z").isoformat() == "2026-10-21T14:03:22+00:00"
    assert parse_as_of("10/21/2026 14:03").isoformat() == "2026-10-21T14:03:00+00:00"
    assert parse_as_of("not a date") is None and parse_as_of(None) is None


@pytest.mark.parametrize("field,value,expected", [
    ("units", "12", (True, 12)), ("units", 12.0, (True, 12)), ("units", "", (True, None)), ("units", "0", (True, None)),
    ("units", "lots", (False, None)), ("year_built", "1985", (True, 1985)), ("year_built", "9999", (False, None)),
    ("county", "Clark County", (True, "CLARK")), ("zip", "4550", (True, "04550")), ("zip", "ABCDE", (False, None)),
    ("property_type", "'-Multifamily", (True, "-Multifamily")), ("lot_acres", "0.5", (True, 0.5)),
    # A minus sign, a unit count that rounds to 0 or a formula with no saved result is a typo, never "erase".
    ("units", "-12", (False, None)), ("units", -12, (False, None)), ("units", 0.4, (False, None)),
    ("units", 0, (True, None)), ("building_sqft", "-5000", (False, None)), ("lot_acres", -0.3, (False, None)),
    ("units", "=36+2", (False, None)), ("property_type", "=A2", (False, None)), ("county", "=B2", (False, None)),
])
def test_parse_cell(field, value, expected):
    assert parse_cell(field, value, 2026) == expected


# --- buy-box decisions from the spreadsheet ---------------------------------------------------------------

def test_decision_words():
    assert decision_intent("Add to buy box") == "include" and decision_intent("added by hand") == "include"
    assert decision_intent(" set aside ") == "exclude" and decision_intent("Set aside by hand") == "exclude"
    assert decision_intent("Follow the rules") == "rules" and decision_intent("Rules") == "rules"
    assert decision_intent("maybe") is None and decision_intent(None) is None
    assert current_decision_text({}, ["drive_time"]) == "Rules"
    assert current_decision_text({"drive_time": {"decision": "include"}}, ["drive_time"]) == \
        "Rule waived: too far to drive"


def test_decisions_column_creates_decisions_by_hand(ranked, tmp_path):
    aside, inside, decided = "1 PLACEHOLDER TOWER", "1200 SAMPLE RIDGE DR", "455 EXAMPLE AVE"
    overrides.decide_property(ranked, APARTMENTS, pid_of(ranked, decided), "exclude", "seth@example.org")
    assert scalar(ranked, "SELECT in_buybox FROM scores WHERE property_id = ?", (pid_of(ranked, decided),)) == 0
    header, data = export_csv(ranked)
    assert by_address(data, decided)["Buy box decision"] == "Set aside by hand"
    by_address(data, aside)["Buy box decision"] = "Add to buy box"
    by_address(data, inside)["Buy box decision"] = "set aside"
    by_address(data, decided)["Buy box decision"] = "Follow the rules"
    by_address(data, "77 FAKE MEADOW CT")["Buy box decision"] = "whatever"
    path = write_csv(tmp_path / "decisions.csv", header, data)
    preview = preview_file(path, market_slug=APARTMENTS, conn=ranked)
    assert preview["edits"]["decisions_changed"] == {"include": 1, "exclude": 1, "rules": 1}
    res = import_file(ranked, APARTMENTS, path, actor="pat@example.org")
    assert res["decisions_changed"] == {"include": 1, "exclude": 1, "rules": 1} and res["batch_id"]
    assert any("Buy box decision cell says something I can't use" in w for w in res["warnings"])
    found = {r["property_id"]: r for r in rows(ranked, "SELECT * FROM buybox_overrides")}
    assert set(found) == {pid_of(ranked, aside), pid_of(ranked, inside)}
    for r in found.values():
        assert r["reason"] == "*" and r["bucket"] is None and r["note"] == "From spreadsheet"
        assert r["decided_by"] == "pat@example.org" and r["batch_id"] == res["batch_id"]
    assert found[pid_of(ranked, aside)]["decision"] == "include"
    assert found[pid_of(ranked, inside)]["decision"] == "exclude"
    score = {r["property_id"]: r for r in rows(ranked, "SELECT property_id, in_buybox, exclusion_reasons FROM scores")}
    assert score[pid_of(ranked, aside)]["in_buybox"] == 1
    out = score[pid_of(ranked, inside)]
    assert out["in_buybox"] == 0 and "manual" in out["exclusion_reasons"]
    assert score[pid_of(ranked, decided)]["in_buybox"] == 1                      # back to the rules
    # The same file again: the cells now match the app, so nothing changes.
    again = import_file(ranked, APARTMENTS, path)
    assert again["decisions_changed"] == {} and again["batch_id"] is None
    # One undo puts all three back.
    overrides.undo(ranked, APARTMENTS, res["batch_id"], "pat@example.org")
    left = {r["property_id"]: r["decision"] for r in rows(ranked, "SELECT * FROM buybox_overrides")}
    assert left == {pid_of(ranked, decided): "exclude"}


def test_a_decision_made_in_the_app_after_the_download_wins(ranked, tmp_path):
    header, data = export_csv(ranked)
    pid = pid_of(ranked, "1 PLACEHOLDER TOWER")
    ranked.execute("UPDATE properties SET updated_at = '2000-01-01T00:00:00+00:00'")
    ranked.commit()
    for r in data:
        r["Row as of"] = "2020-01-01T00:00:00+00:00"
    by_address(data, "1 PLACEHOLDER TOWER")["Buy box decision"] = "Set aside"
    path = write_csv(tmp_path / "old.csv", header, data)
    overrides.decide_property(ranked, APARTMENTS, pid, "include", "seth@example.org")    # after the download
    res = import_file(ranked, APARTMENTS, path)
    assert res["decisions_changed"] == {} and res["skipped_stale"] == 1
    assert one(ranked, "SELECT decision FROM buybox_overrides WHERE property_id = ?", (pid,))["decision"] == "include"


# --- a later Crexi import keeps the corrections -----------------------------------------------------------

def test_corrections_survive_a_later_crexi_import(ranked, sample_csv, tmp_path):
    pid = pid_of(ranked, "77 FAKE MEADOW CT")
    path = edited_csv(ranked, tmp_path, {"77 FAKE MEADOW CT": {"Year built": "1999", "Units": "14"}})
    import_file(ranked, APARTMENTS, path)
    res = import_file(ranked, APARTMENTS, sample_csv)
    assert res["edits_kept"] == 2
    assert any(w.startswith("Kept 2 values you corrected by hand (units 1, year built 1)") for w in res["warnings"])
    row = one(ranked, "SELECT units, year_built, raw_json FROM properties WHERE id = ?", (pid,))
    assert (row["units"], row["year_built"]) == (14, 1999)
    assert loads(row["raw_json"])["Year Built"] == "1974"                 # Crexi's value is still kept
    # Crexi now says something new: the correction still counts, and the import says so.
    header, crexi = read_csv(sample_csv)
    by_address(crexi, "77 FAKE MEADOW CT")["Year Built"] = "1976"
    res = import_file(ranked, APARTMENTS, write_csv(tmp_path / "crexi2.csv", header, crexi))
    assert any("Crexi now shows a different value" in w for w in res["warnings"])
    assert scalar(ranked, "SELECT year_built FROM properties WHERE id = ?", (pid,)) == 1999


# --- preview and the review file --------------------------------------------------------------------------

def test_preview_counts_changes_and_shows_no_owner_data(ranked, tmp_path):
    path = edited_csv(ranked, tmp_path, {"1200 SAMPLE RIDGE DR": {"Units": "30", "Year built": "1979"},
                                         "455 EXAMPLE AVE": {"Units": "37"},
                                         "1 PLACEHOLDER TOWER": {"Buy box decision": "Add to buy box"}})
    before = snapshot(ranked)
    pv = preview_file(path, market_slug=APARTMENTS, conn=ranked)
    assert pv["provider"] == "bluemoon" and pv["provider_label"].startswith("Blue Moon export")
    assert pv["message"] == ("This is a Blue Moon export. 3 properties have changes: units 2, year built 1, "
                             "buy box decisions 1.")
    assert snapshot(ranked) == before                     # a preview writes nothing
    text = json.dumps(pv)
    for r in rows(ranked, "SELECT owner_name, mailing_address FROM properties"):
        for v in (r["owner_name"], r["mailing_address"]):
            assert not v or v not in text, v
    assert {c["label"] for c in pv["edits"]["changes"]} == {"Units", "Year built", "Buy box decision"}
    res = pipeline.run_step(ranked, APARTMENTS, "import", path=path, dry_run=True)
    assert "3 properties have changes" in res["message"] and snapshot(ranked) == before


def test_review_file_holds_exactly_one_group_and_comes_back(ranked, tmp_path):
    bucket = "drive_time:76-90"
    members = set(overrides.group_property_ids(ranked, 1, bucket))
    assert members
    path = exports.export_review(ranked, APARTMENTS, bucket)
    assert path.parent.name == "review" and path.name == "review_drive_time_76-90.xlsx"
    wb = load_workbook(path)
    assert wb.sheetnames == ["Review", "How to edit"]
    ws = wb["Review"]
    head = [c.value for c in ws[1]]
    assert head[:10] == ["Address", "City", "Buy box decision", "Units", "Building sq ft", "Year built",
                         "Property type", "County", "ZIP", "Lot acres"]
    assert head[-2:] == ["Property ID", "Row as of"] and ws.freeze_panes == "C2"
    ids = {ws.cell(r, head.index("Property ID") + 1).value for r in range(2, ws.max_row + 1)}
    assert ids == members
    notes = {r[0]: r[1] for r in wb["How to edit"].iter_rows(min_row=2, values_only=True)}
    assert "Too far to drive" in notes["Group"] and "Add to buy box" in notes["Buy box decision"]
    ws.cell(2, head.index("Units") + 1).value = 21
    changed = ws.cell(2, head.index("Property ID") + 1).value
    wb.save(tmp_path / "review_back.xlsx")
    res = import_file(ranked, APARTMENTS, tmp_path / "review_back.xlsx")
    assert res["provider"] == "bluemoon" and res["fields_changed"] == {"units": 1}
    assert scalar(ranked, "SELECT units FROM properties WHERE id = ?", (changed,)) == 21
    with pytest.raises(ValueError):
        exports.export_review(ranked, APARTMENTS, "no such group")


def test_the_dashboard_round_trip(settings, sample_csv, tmp_path):
    from fastapi.testclient import TestClient

    from bluemoon.db import connect
    from bluemoon.web.app import create_app

    connect(settings=settings).close()
    with TestClient(create_app()) as c:
        with sample_csv.open("rb") as fh:
            res = c.post(f"/api/m/{APARTMENTS}/import", files={"file": ("crexi.csv", fh, "text/csv")})
        assert res.status_code == 200
        res = c.get(f"/api/m/{APARTMENTS}/export/review.xlsx", params={"bucket": "drive_time:76-90"})
        assert res.status_code == 200 and res.content[:2] == b"PK"
        assert "review_drive_time_76-90.xlsx" in res.headers["content-disposition"]
        assert c.get(f"/api/m/{APARTMENTS}/export/review.xlsx", params={"bucket": "x;drop"}).status_code == 400
        assert c.get(f"/api/m/{APARTMENTS}/export/review.xlsx").status_code == 422
        csv_bytes = c.get(f"/api/m/{APARTMENTS}/export/ranked_list.csv").content
        (tmp_path / "dl.csv").write_bytes(csv_bytes)
        header, data = read_csv(tmp_path / "dl.csv")
        by_address(data, "455 EXAMPLE AVE")["Units"] = "38"
        path = write_csv(tmp_path / "back.csv", header, data)
        with path.open("rb") as fh:
            pv = c.post(f"/api/m/{APARTMENTS}/import/preview", files={"file": ("back.csv", fh, "text/csv")}).json()
        assert pv["provider"] == "bluemoon" and pv["edits"]["fields_changed"] == {"units": 1}
        with path.open("rb") as fh:
            done = c.post(f"/api/m/{APARTMENTS}/import", files={"file": ("back.csv", fh, "text/csv")}).json()
        assert done["provider"] == "bluemoon" and done["properties_changed"] == 1
        assert done["after_import"]["ok"]
    conn = connect(settings=settings)
    try:
        edit = one(conn, "SELECT actor, via FROM property_edits")
        assert edit["actor"] == "open" and edit["via"].startswith("import:")
    finally:
        conn.close()


# --- cells that are not what they seem, old files, files that aren't property lists ---------------------

def cache_formula_result(path: Path, formula: str, result: object) -> None:
    """Save a formula's result in the workbook, the way Excel does (openpyxl saves none)."""
    tmp = path.with_name(path.name + ".tmp")
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as out:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename.startswith("xl/worksheets/"):
                data = re.sub(rf"<f>{re.escape(formula)}</f>(<v\s*/>|<v></v>)?", f"<f>{formula}</f><v>{result}</v>",
                              data.decode("utf-8")).encode("utf-8")
            out.writestr(item, data)
    tmp.replace(path)


def units_of(conn, pid: int):
    return scalar(conn, "SELECT units FROM properties WHERE id = ?", (pid,))


def test_a_formula_reads_as_the_result_excel_saved_and_never_as_empty(ranked, tmp_path):
    path = exports.export_market(ranked, APARTMENTS, only=["ranked_list.xlsx"])["ranked_list.xlsx"]
    wb = load_workbook(path)
    ws = wb["Ranked"]
    col = sheet_cols(ws)
    picks = [r for r in range(2, ws.max_row + 1) if ws.cell(r, col["Units"]).value]
    (saved_row, bare_row) = picks[:2]
    saved_pid, saved_units = ws.cell(saved_row, col["Property ID"]).value, ws.cell(saved_row, col["Units"]).value
    bare_pid, bare_units = ws.cell(bare_row, col["Property ID"]).value, ws.cell(bare_row, col["Units"]).value
    ws.cell(saved_row, col["Units"]).value = f"={saved_units}+2"
    ws.cell(bare_row, col["Units"]).value = f"={bare_units}+3"
    wb.save(tmp_path / "formulas.xlsx")
    cache_formula_result(tmp_path / "formulas.xlsx", f"{saved_units}+2", saved_units + 2)
    pv = preview_file(tmp_path / "formulas.xlsx", market_slug=APARTMENTS, conn=ranked)
    assert pv["edits"]["fields_changed"] == {"units": 1} and pv["edits"]["clears"] == {}
    assert any(f"Units '={bare_units}+3' isn't a value I can read" in w for w in pv["warnings"])
    import_file(ranked, APARTMENTS, tmp_path / "formulas.xlsx")
    assert units_of(ranked, saved_pid) == saved_units + 2           # Excel's saved result
    assert units_of(ranked, bare_pid) == bare_units                 # no result saved: left alone, not cleared


def test_a_formula_or_a_minus_sign_in_a_csv_changes_nothing(ranked, tmp_path):
    before = snapshot(ranked)
    path = edited_csv(ranked, tmp_path, {"455 EXAMPLE AVE": {"Units": "=36+2", "Lot acres": "-0.5"},
                                         "1200 SAMPLE RIDGE DR": {"Units": "-12"}})
    res = import_file(ranked, APARTMENTS, path)
    assert res["properties_changed"] == 0 and snapshot(ranked) == before
    for text in ("Units '=36+2'", "Lot acres '-0.5'", "Units '-12'"):
        assert any(text in w for w in res["warnings"]), text


def test_a_csv_without_row_as_of_changes_nothing(ranked, tmp_path):
    header, data = export_csv(ranked)
    by_address(data, "455 EXAMPLE AVE")["Units"] = "40"
    for r in data:
        del r["Row as of"]
    path = write_csv(tmp_path / "old.csv", [h for h in header if h != "Row as of"], data)
    before = snapshot(ranked)
    pv = preview_file(path, market_slug=APARTMENTS, conn=ranked)
    assert pv["provider"] == "bluemoon" and pv["edits"]["skipped_unchecked"] == 1
    assert "Importing it would change nothing" in pv["message"] and "no 'Row as of' column" in pv["message"]
    res = import_file(ranked, APARTMENTS, path)
    assert res["properties_changed"] == 0 and res["rows_skipped"] == 1 and snapshot(ranked) == before


def test_an_old_workbook_without_row_as_of_goes_by_when_it_was_exported(ranked, tmp_path):
    path = exports.export_market(ranked, APARTMENTS, only=["ranked_list.xlsx"])["ranked_list.xlsx"]
    wb = load_workbook(path)
    drop_columns(wb, "Row as of")
    wb.properties.created = datetime(2020, 1, 1)
    ws = wb["Ranked"]
    col = sheet_cols(ws)
    (row_a, row_b) = [r for r in range(2, ws.max_row + 1) if ws.cell(r, col["Units"]).value][:2]
    pid_a, units_a = ws.cell(row_a, col["Property ID"]).value, ws.cell(row_a, col["Units"]).value
    pid_b, units_b = ws.cell(row_b, col["Property ID"]).value, ws.cell(row_b, col["Units"]).value
    ws.cell(row_a, col["Units"]).value = units_a + 5
    ws.cell(row_b, col["Units"]).value = units_b + 5
    wb.save(tmp_path / "old.xlsx")
    ranked.execute("UPDATE properties SET updated_at = '2000-01-01T00:00:00+00:00'")
    ranked.execute("UPDATE properties SET updated_at = '2021-06-01T00:00:00+00:00' WHERE id = ?", (pid_b,))
    ranked.commit()                                                  # B changed in the app after the export
    res = import_file(ranked, APARTMENTS, tmp_path / "old.xlsx")
    assert res["fields_changed"] == {"units": 1} and res["skipped_stale"] == 1
    assert units_of(ranked, pid_a) == units_a + 5 and units_of(ranked, pid_b) == units_b
    assert any("went by the time it was downloaded" in w for w in res["warnings"])
    # The same file again: 'All properties' still holds A's old value, but A changed since, so it stays.
    again = import_file(ranked, APARTMENTS, tmp_path / "old.xlsx")
    assert again["properties_changed"] == 0 and units_of(ranked, pid_a) == units_a + 5


def test_a_property_sheet_that_lost_row_as_of_is_named(ranked, tmp_path):
    path = exports.export_market(ranked, APARTMENTS, only=["ranked_list.xlsx"])["ranked_list.xlsx"]
    wb = load_workbook(path)
    ws = wb["Ranked"]
    col = sheet_cols(ws)
    ws.cell(2, col["Units"]).value = (ws.cell(2, col["Units"]).value or 0) + 7
    drop_columns(wb, "Row as of", ["Ranked"])
    wb.save(tmp_path / "edited.xlsx")
    pv = preview_file(tmp_path / "edited.xlsx", market_slug=APARTMENTS, conn=ranked)
    assert pv["edits"]["properties_changed"] == 0
    assert any("The 'Ranked' sheet has no 'Row as of' column, so its changes were not read" in w
               for w in pv["warnings"])
    assert not any("DRILL worklist" in w for w in pv["warnings"])


def test_our_file_without_its_property_id_column_is_refused(ranked, tmp_path):
    files = exports.export_market(ranked, APARTMENTS, only=["ranked_list.xlsx", "ranked_list.csv"])
    wb = load_workbook(files["ranked_list.xlsx"])
    drop_columns(wb, "Property ID")
    wb.save(tmp_path / "trimmed.xlsx")
    header, data = read_csv(files["ranked_list.csv"])
    for r in data:
        del r["Property ID"]
    header = [h for h in header if h != "Property ID"]
    trimmed_csv = write_csv(tmp_path / "trimmed.csv", header, data)
    assert looks_like_blue_moon(header) and not is_blue_moon_export(header)
    before = snapshot(ranked)
    for path in (tmp_path / "trimmed.xlsx", trimmed_csv):
        with pytest.raises(IngestError, match="Property ID column is missing"):
            preview_file(path, market_slug=APARTMENTS, conn=ranked)
        with pytest.raises(IngestError, match="Property ID column is missing"):
            import_file(ranked, APARTMENTS, path)
    assert snapshot(ranked) == before                    # no Crexi merge: raw_json and source ids untouched


def test_our_other_exports_are_refused_not_merged(ranked):
    files = exports.export_market(ranked, APARTMENTS, only=["drill_worklist.csv", "mailing_list.csv",
                                                            "crm_handoff.csv"], stamp_handoffs=False)
    before = snapshot(ranked)
    for name, what in (("drill_worklist.csv", "DRILL worklist"), ("mailing_list.csv", "mailing list"),
                       ("crm_handoff.csv", "CRM handoff")):
        assert other_blue_moon_export(read_csv(files[name])[0]) == what, name
        with pytest.raises(IngestError, match=what):
            preview_file(files[name], market_slug=APARTMENTS, conn=ranked)
        with pytest.raises(IngestError, match=what):
            import_file(ranked, APARTMENTS, files[name])
    assert snapshot(ranked) == before
    assert other_blue_moon_export(export_csv(ranked)[0]) is None


def test_emptied_cells_are_counted_apart_and_ask_first_over_the_limit(ranked, settings, tmp_path):
    header, data = export_csv(ranked)
    had = sum(1 for r in data if r["Units"])
    for r in data:
        r["Units"] = ""
    path = write_csv(tmp_path / "blank_units.csv", header, data)
    pv = preview_file(path, market_slug=APARTMENTS, conn=ranked)
    assert pv["edits"]["clears"] == {"units": had} and pv["edits"]["confirm_clears"] is (had > 10)
    assert any(w.startswith(f"{had} cells are empty in this file") and "erases those values" in w
               for w in pv["warnings"])
    settings.raw["safety"]["confirm_clears_over"] = 1
    assert preview_file(path, market_slug=APARTMENTS, conn=ranked)["edits"]["confirm_clears"] is True
    one_blank = edited_csv(ranked, tmp_path, {"455 EXAMPLE AVE": {"Units": ""}}, "one_blank.csv")
    assert preview_file(one_blank, market_slug=APARTMENTS, conn=ranked)["edits"]["confirm_clears"] is False
    res = import_file(ranked, APARTMENTS, path)
    assert res["clears"] == {"units": had} and any(w.startswith(f"Erased {had} values") for w in res["warnings"])


# --- the Import and Exports screens (Node, a tiny DOM stub; skipped without Node) -------------------------

SCREENS_HARNESS = r"""
const fs = require('fs');
const path = require('path');
const [fxPath, staticDir, outPath] = process.argv.slice(2);
const fx = JSON.parse(fs.readFileSync(fxPath, 'utf8'));
const out = {};
const calls = [];
class Element { closest() { return this; } }
class HTMLInputElement extends Element {}
Object.assign(global, {Element, HTMLInputElement});
global.CustomEvent = class { constructor(type, o) { this.type = type; this.detail = o && o.detail; } };
global.window = {location: {hash: ''}, localStorage: {getItem: () => null, setItem: () => {}},
  addEventListener() {}, removeEventListener() {}, scrollTo() {}};
global.document = {readyState: 'loading', addEventListener() {}, removeEventListener() {}, dispatchEvent() {},
  querySelector: () => null, visibilityState: 'visible'};
global.history = {replaceState() {}};
require(path.join(staticDir, 'core.js'));
const BM = window.BM;
const views = {};
BM.registerView = (key, def) => { views[key] = def; };
BM.api = async (p, opts) => {
  const o = opts || {};
  const key = (o.method || (o.form ? 'POST' : 'GET')) + ' ' + p;
  calls.push({key, body: o.body || null});
  if (!(key in fx.api)) throw new Error('no fixture for ' + key);
  return JSON.parse(JSON.stringify(fx.api[key]));
};
BM.refreshSummary = async () => BM.state.summary;
BM.toast = () => {};
Object.assign(BM.state, {market: fx.market, markets: fx.markets, summary: fx.summary});
['funnel.js', 'importer.js', 'exports.js'].forEach((f) => require(path.join(staticDir, f)));
function makeRoot() {
  const listeners = {};
  const parts = {};
  return {innerHTML: '', isConnected: true, parts, listeners,
    addEventListener(t, f) { (listeners[t] = listeners[t] || []).push(f); }, removeEventListener() {},
    contains: () => true, querySelectorAll: () => [],
    querySelector(sel) {
      if (!parts[sel]) parts[sel] = {innerHTML: '', textContent: '', style: {}, querySelector: () => null,
        setAttribute() {}, classList: {toggle() {}, add() {}, remove() {}}};
      return parts[sel];
    }};
}
const text = (root) => root.innerHTML + Object.values(root.parts).map((p) => p.innerHTML).join('');
function target(cls, attrs) {
  const e = new cls();
  e.getAttribute = (k) => (k in attrs ? String(attrs[k]) : null);
  e.matches = () => true;
  Object.assign(e, attrs.props || {});
  return e;
}
async function settle() { for (let i = 0; i < 6; i++) await new Promise((r) => setTimeout(r, 0)); }
async function fire(root, type, t) {
  (root.listeners[type] || []).forEach((f) => f({type, target: t, preventDefault() {}, dataTransfer: null}));
  await settle();
}
(async () => {
  const im = makeRoot();
  await views.import.render(im, []);
  await settle();
  const file = new File([fs.readFileSync(fx.editedCsv)], 'ranked_list.csv', {type: 'text/csv'});
  await fire(im, 'change', target(HTMLInputElement, {props: {files: [file], value: ''}}));
  out.preview = text(im);
  if (fx.askClears) {
    await fire(im, 'click', target(Element, {'data-im': 'ask-clears'}));
    out.asked = text(im);
    out.importedBeforeOk = calls.some((c) => c.key.endsWith('/import'));
  }
  await fire(im, 'click', target(Element, {'data-im': 'import'}));
  out.result = text(im);
  await fire(im, 'click', target(Element, {'data-im': 'undo', 'data-batch': fx.batch}));
  out.undone = text(im);
  const ex = makeRoot();
  await views.exports.render(ex, []);
  await settle();
  out.exports = text(ex);
  out.calls = calls;
  fs.writeFileSync(outPath, JSON.stringify(out));
  process.exit(0);
})().catch((e) => { console.error(e && e.stack || e); process.exit(1); });
"""


@pytest.mark.parametrize("ask_clears", [False, True])
def test_import_and_exports_screens_show_the_round_trip(settings, sample_csv, tmp_path, ask_clears):
    import shutil
    import subprocess

    from fastapi.testclient import TestClient

    from bluemoon.db import connect
    from bluemoon.web.app import create_app

    node = shutil.which("node")
    if not node:
        pytest.skip("Node is not installed")
    connect(settings=settings).close()
    base = f"/api/m/{APARTMENTS}"
    api: dict[str, object] = {}
    with TestClient(create_app()) as c:
        with sample_csv.open("rb") as fh:
            assert c.post(f"{base}/import", files={"file": ("crexi.csv", fh, "text/csv")}).status_code == 200
        (tmp_path / "dl.csv").write_bytes(c.get(f"{base}/export/ranked_list.csv").content)
        header, data = read_csv(tmp_path / "dl.csv")
        by_address(data, "455 EXAMPLE AVE")["Units"] = "38"
        by_address(data, "1 PLACEHOLDER TOWER")["Buy box decision"] = "Add to buy box"
        edited = write_csv(tmp_path / "edited.csv", header, data)
        for route in ("import/preview", "import"):
            with edited.open("rb") as fh:
                res = c.post(f"{base}/{route}", files={"file": ("edited.csv", fh, "text/csv")})
            assert res.status_code == 200, res.text
            api[f"POST {base}/{route}"] = res.json()
        batch = api[f"POST {base}/import"]["batch_id"]
        if ask_clears:      # as if the file emptied more cells than safety.confirm_clears_over
            api[f"POST {base}/import/preview"]["edits"].update({"clears": {"units": 12}, "confirm_clears": True})
        api[f"POST {base}/buybox/undo"] = c.post(f"{base}/buybox/undo", json={"batch_id": batch}).json()
        api[f"GET {base}/analytics"] = c.get(f"{base}/analytics").json()
        summary = c.get(f"{base}/summary").json()
        markets = c.get("/api/markets").json()
    fx = tmp_path / "fx.json"
    fx.write_text(json.dumps({"api": api, "market": APARTMENTS, "markets": markets, "summary": summary,
                              "editedCsv": str(edited), "batch": batch, "askClears": ask_clears}),
                  encoding="utf-8")
    script = tmp_path / "harness.js"
    script.write_text(SCREENS_HARNESS, encoding="utf-8")
    out = tmp_path / "out.json"
    static = Path(__file__).resolve().parent.parent / "bluemoon" / "web" / "static"
    done = subprocess.run([node, str(script), str(fx), str(static), str(out)], capture_output=True, text=True,
                          timeout=120)
    assert done.returncode == 0, done.stderr
    r = json.loads(out.read_text(encoding="utf-8"))
    for text in ("Check your edits", "Blue Moon export (your edits)", "2 properties have changes: units 1, buy box "
                 "decisions 1", "What will change", "455 EXAMPLE AVE", "Save the changes to 2 properties"):
        assert text in r["preview"], text
    assert "Our field" not in r["preview"]                     # not the Crexi column-matching preview
    if ask_clears:              # the first Save only asks; nothing is imported until "Yes, erase them"
        assert 'data-im="ask-clears"' in r["preview"] and "Erase 12 values?" not in r["preview"]
        assert "Erase 12 values?" in r["asked"] and "Yes, erase them and save" in r["asked"]
        assert r["importedBeforeOk"] is False
    else:
        assert 'data-im="ask-clears"' not in r["preview"] and 'data-im="import"' in r["preview"]
    for text in ("Imported your edits", "Properties changed", "Undo the buy box decisions", 'data-batch="'):
        assert text in r["result"], text
    assert "Undo the buy box decisions" not in r["undone"] and "Undone" in r["undone"]
    undo = [c for c in r["calls"] if c["key"] == f"POST {base}/buybox/undo"]
    assert undo and undo[0]["body"] == {"batch_id": batch}
    for heading in ("Work lists", "Hand off", "Fix data in a spreadsheet", "Summary"):   # the four Exports groups
        assert f'<h2 class="vw-sub">{heading}</h2>' in r["exports"], heading
    assert "Files for Brad" not in r["exports"] and "for Brad" not in r["exports"].split("Hand off", 1)[0]
    assert 'data-ex="review"' in r["exports"] and 'value="drive_time:76-90"' in r["exports"]
