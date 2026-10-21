"""Export files (bluemoon/exports.py): the xlsx/csv for Brad, excluded files, CRM handoff, mailing list,
DRILL worklist and run summary. Uses the invented sample export; nothing touches the network."""
from __future__ import annotations

import csv
import zipfile
from pathlib import Path

import httpx
import pytest
from openpyxl import load_workbook

from bluemoon import exports, pipeline
from bluemoon.db import now_iso, one, rows, scalar
from bluemoon.drill import engine
from bluemoon.scoring import EXCLUSION_LABELS, EXCLUSION_REASONS

FAKE_KEY = "TESTKEY-exports-should-never-appear"


@pytest.fixture(autouse=True)
def no_network(monkeypatch):
    def refuse(self, request):
        raise AssertionError(f"Tried to reach the network: {request.url.host}")
    monkeypatch.setattr(httpx.HTTPTransport, "handle_request", refuse)


@pytest.fixture
def ranked(conn, market_slug, sample_csv):
    res = pipeline.run_step(conn, market_slug, "import", path=sample_csv)
    assert res["ok"], res["message"]
    return conn


@pytest.fixture
def out(settings, market_slug) -> Path:
    return settings.path("output") / market_slug


def read_csv(path: Path) -> list[dict]:
    with path.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def ranked_ids(conn) -> list[int]:
    return [r["property_id"] for r in rows(conn, "SELECT property_id FROM scores WHERE rank IS NOT NULL ORDER BY rank")]


# --- every file --------------------------------------------------------------------------------------

def test_export_writes_every_file_to_the_market_folder(ranked, market_slug, out):
    files = exports.export_market(ranked, market_slug)
    assert list(files) == list(exports.EXPORT_FILES)
    for name, path in files.items():
        assert path == out / name and path.is_file() and path.stat().st_size > 0
    assert not list(out.glob(".*.tmp"))


def test_only_some_files_and_unknown_names(ranked, market_slug, tmp_path):
    files = exports.export_market(ranked, market_slug, tmp_path, only=["ranked_list.csv", "excluded"])
    assert set(files) == {"ranked_list.csv", *exports.EXCLUDED_FILES}
    assert all(p.parent == tmp_path for p in files.values())
    with pytest.raises(ValueError, match="Unknown export file"):
        exports.export_market(ranked, market_slug, only=["secrets.txt"])


def test_no_keys_in_any_file(ranked, market_slug, monkeypatch):
    monkeypatch.setenv("GOOGLE_MAPS_API_KEY", FAKE_KEY)
    files = exports.export_market(ranked, market_slug)
    for name, path in files.items():
        if path.suffix != ".xlsx":
            assert FAKE_KEY not in path.read_text(encoding="utf-8-sig"), name
    wb = load_workbook(files["ranked_list.xlsx"])
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            assert not any(isinstance(v, str) and FAKE_KEY in v for v in row)


# --- ranked_list.xlsx / .csv ----------------------------------------------------------------------------

def test_xlsx_sheets_and_formatting(ranked, market_slug):
    path = exports.export_market(ranked, market_slug, only=["ranked_list.xlsx"])["ranked_list.xlsx"]
    wb = load_workbook(path)
    assert wb.sheetnames == ["Ranked", "All properties", "Excluded", "DRILL worklist", "Summary"]
    ws = wb["Ranked"]
    header = [c.value for c in ws[1]]
    assert ws.max_row - 1 == len(ranked_ids(ranked))
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref.startswith("A1:")
    assert ws.cell(1, 1).font.b
    assert [ws.cell(r, 1).value for r in range(2, ws.max_row + 1)] == list(range(1, ws.max_row))
    # links are clickable
    crexi = ws.cell(2, header.index("Crexi") + 1)
    assert crexi.hyperlink is not None and crexi.hyperlink.target.startswith("https://www.crexi.com/")
    assert ws.cell(2, header.index("Google Maps") + 1).hyperlink is not None
    # colored color columns
    size = ws.cell(2, header.index("Size color") + 1)
    assert size.value in ("green", "yellow", "red")
    assert size.fill.fgColor.rgb.endswith(exports.FILLS[size.value])
    # original Crexi columns after ours, a clashing name kept apart
    assert header.index("Property ID") < header.index("Owner Name")
    assert "Address (Crexi)" in header and header.index("Address") < header.index("Property ID")
    assert ws.column_dimensions["E"].width >= 8
    assert wb["All properties"].max_row - 1 == 11
    assert wb["Excluded"].max_row - 1 == scalar(ranked, "SELECT COUNT(*) FROM scores WHERE in_buybox = 0")
    summary = {(r[0], r[1]): r[2] for r in wb["Summary"].iter_rows(min_row=2, values_only=True)}
    assert summary[("List", "Properties")] == 11


def test_ranked_csv_puts_the_buy_box_first(ranked, market_slug):
    path = exports.export_market(ranked, market_slug, only=["ranked_list.csv"])["ranked_list.csv"]
    data = read_csv(path)
    assert len(data) == 11
    ranks = [r["Rank"] for r in data]
    n = len(ranked_ids(ranked))
    assert ranks[:n] == [str(i) for i in range(1, n + 1)] and not any(ranks[n:])
    assert all(r["In buy box"] == "Yes" for r in data[:n])
    assert all(r["In buy box"] == "No" and r["Why set aside"] for r in data[n:])
    assert [int(r["Property ID"]) for r in data[:n]] == ranked_ids(ranked)
    assert data[0]["Owner Name"] and data[0]["Crexi"].startswith("https://")
    assert "Photo label" in data[0] and "Rent fit" in data[0]


def test_text_that_looks_like_a_formula_stays_text(ranked, market_slug):
    pid = ranked_ids(ranked)[0]
    ranked.execute("UPDATE properties SET owner_name = '=HYPERLINK(\"http://x\")' WHERE id = ?", (pid,))
    ranked.commit()
    files = exports.export_market(ranked, market_slug, only=["ranked_list.csv", "ranked_list.xlsx"])
    assert read_csv(files["ranked_list.csv"])[0]["Owner"].startswith("'=")
    ws = load_workbook(files["ranked_list.xlsx"])["Ranked"]
    header = [c.value for c in ws[1]]
    cell = ws.cell(2, header.index("Owner") + 1)
    assert cell.data_type == "s" and cell.value.startswith("=HYPERLINK")


def test_asset_class_without_photos_or_rent_has_no_such_columns(conn, sample_csv):
    pipeline.run_step(conn, "land-springfield-oh", "import", path=sample_csv)
    path = exports.export_market(conn, "land-springfield-oh", only=["ranked_list.csv"])["ranked_list.csv"]
    with path.open(encoding="utf-8-sig") as f:
        header = next(csv.reader(f))
    assert "Photo label" not in header and "Rent fit" not in header


def test_empty_market_still_gets_header_only_files(conn):
    files = exports.export_market(conn, "self-storage-springfield-oh")
    assert read_csv(files["ranked_list.csv"]) == []
    assert "Rank" in files["ranked_list.csv"].read_text(encoding="utf-8-sig")
    assert "Properties: 0" in files["run_summary.md"].read_text()


# --- excluded_<reason>.csv ------------------------------------------------------------------------------

def test_one_excluded_file_per_reason_with_a_plain_reason(ranked, market_slug):
    files = exports.export_market(ranked, market_slug, only=["excluded"])
    assert set(files) == set(exports.EXCLUDED_FILES)
    for reason in EXCLUSION_REASONS:
        data = read_csv(files[f"excluded_{reason}.csv"])
        expected = scalar(ranked, "SELECT COUNT(*) FROM scores WHERE exclusion_reasons LIKE ?", (f'%"{reason}"%',))
        assert len(data) == expected, reason
        assert all(r["Reason"].startswith(EXCLUSION_LABELS[reason]) for r in data)
    far = read_csv(files["excluded_drive_time.csv"])
    assert far and "min" in far[0]["Reason"]


def test_excluded_zip(ranked, market_slug):
    path = exports.export_zip(ranked, market_slug)
    assert path.name == "excluded.zip"
    with zipfile.ZipFile(path) as z:
        assert sorted(z.namelist()) == sorted(exports.EXCLUDED_FILES)


# --- CRM handoff, mailing list, worklist ----------------------------------------------------------------

def test_crm_handoff_lists_confirmed_owners_and_stamps_them(ranked, market_slug):
    pid = ranked_ids(ranked)[0]
    engine.level_action(ranked, pid, "confirm", person="Pat Example")
    path = exports.export_market(ranked, market_slug, only=["crm_handoff.csv"])["crm_handoff.csv"]
    data = read_csv(path)
    assert [int(r["property_id"]) for r in data] == [pid]
    assert data[0]["confirmed_person"] == "Pat Example" and data[0]["new_in_this_export"] == "Yes"
    assert one(ranked, "SELECT handed_off_at FROM drill_state WHERE property_id = ?", (pid,))["handed_off_at"]
    again = read_csv(exports.export_market(ranked, market_slug, only=["crm_handoff.csv"])["crm_handoff.csv"])
    assert again[0]["new_in_this_export"] == "No"


def test_crm_handoff_with_nobody_confirmed_has_a_header(ranked, market_slug):
    path = exports.export_market(ranked, market_slug, only=["crm_handoff.csv"])["crm_handoff.csv"]
    assert read_csv(path) == []
    assert path.read_text(encoding="utf-8-sig").startswith("property_id,")


def test_mailing_list_never_includes_do_not_contact_or_confirmed(ranked, market_slug):
    ids = ranked_ids(ranked)
    engine.level_action(ranked, ids[0], "dnc")
    engine.level_action(ranked, ids[1], "confirm", person="Pat Example")
    data = read_csv(exports.export_market(ranked, market_slug, only=["mailing_list.csv"])["mailing_list.csv"])
    listed = {int(r["Property ID"]) for r in data}
    assert listed and ids[0] not in listed and ids[1] not in listed
    row = data[0]
    assert row["Name"] and row["Full mailing address"] and row["Why on this list"]
    record = one(ranked, "SELECT mailing_address, mailing_zip FROM properties WHERE id = ?", (int(row["Property ID"]),))
    assert row["Mailing address"] == record["mailing_address"] and row["Mailing ZIP"] == record["mailing_zip"]


def test_mailing_list_when_letters_are_next(ranked, market_slug, monkeypatch):
    pid = ranked_ids(ranked)[0]
    real = engine.summaries

    def letters_next(conn, slug=None, property_ids=None, **kw):
        out = real(conn, slug, property_ids, **kw)
        out[pid]["next_step"] = {"kind": "stage", "stage_key": "letters", "text": "Letters"}
        return out

    monkeypatch.setattr(engine, "summaries", letters_next)
    ranked.execute("UPDATE contact_points SET status = 'verified' WHERE property_id = ? AND channel = 'mail'", (pid,))
    ranked.commit()
    data = read_csv(exports.export_market(ranked, market_slug, only=["mailing_list.csv"])["mailing_list.csv"])
    mine = [r for r in data if int(r["Property ID"]) == pid]
    assert mine and mine[0]["Why on this list"] == "Letters are the next DRILL step"


def test_drill_worklist(ranked, market_slug):
    ids = ranked_ids(ranked)
    engine.level_action(ranked, ids[0], "dnc")
    data = read_csv(exports.export_market(ranked, market_slug, only=["drill_worklist.csv"])["drill_worklist.csv"])
    listed = [int(r["Property ID"]) for r in data]
    assert ids[0] not in listed and ids[1] in listed
    assert listed.index(ids[1]) == 0
    row = data[0]
    assert row["Next step"] and row["Contact level"].startswith(("1 ", "2 ", "3 "))
    assert row["Mail"] in ("none", "untested", "verified", "wrong", "dnc")


# --- run_summary.md ---------------------------------------------------------------------------------------

def test_run_summary(ranked, market_slug):
    ranked.execute("INSERT INTO api_calls(provider, sku, units, created_at) VALUES('google', 'route_matrix_element', "
                   "10500, ?)", (now_iso(),))
    ranked.execute("INSERT INTO api_calls(provider, sku, units, created_at) VALUES('claude_cli', 'vision_claude_cli', "
                   "7, ?)", (now_iso(),))
    ranked.commit()
    text = exports.export_market(ranked, market_slug, only=["run_summary.md"])["run_summary.md"].read_text()
    for heading in ("## The list", "## Colors in the buy box", "## Steps and layers", "## DRILL", "## Crexi credits",
                    "## API calls this month", "## Files"):
        assert heading in text
    assert "| Google Routes (drive times) | 10,500 | 500 | $2.50 |" in text
    assert "Claude plan (no per-call bill)" in text
    assert "Estimated cost this month: $2.50" in text
    assert "credits left this month cover" in text and "rank #" in text
    assert "Too far to drive" in text and "Penetration" in text


def test_credit_reach(ranked, market_slug):
    n = len(ranked_ids(ranked))
    reach = exports.credit_reach(ranked, market_slug, n)
    assert reach["left"] == 500 and 0 < reach["planned"] <= n
    assert reach["reach_rank"] <= n and "credits have no target yet" in reach["reach_text"]
