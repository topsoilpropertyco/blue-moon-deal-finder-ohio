"""Import a Crexi export (or another provider's property list) into the master list.

Flow: read the file (CSV or XLSX) -> map columns -> parse each row into a PropertyRecord ->
give each row a property key -> collapse duplicate rows inside the file -> merge into the
`properties` table (new rows inserted, known rows updated, nothing ever deleted) -> make sure
every property has a DRILL state.

Property key:
- `apn:{STATE}:{COUNTY}:{APN}` when the row has a parcel number. The APN part keeps only letters
  and digits, so "A01-0001 0001" and "A0100010001" are the same parcel (Crexi often lists one
  parcel twice with different punctuation).
- `apn:...|{street}` when one parcel shows up with two different street addresses in a file:
  both are kept as separate properties and flagged `shared_apn`.
- `addr:{normalized address}` when there is no parcel number.
- `src:{source}:{id}` when there is neither (rare; e.g. a Crexi link only).
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import logging
import math
import re
import sqlite3
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any, Callable, Iterable, Mapping
from urllib.parse import parse_qs, urlsplit

from pydantic import BaseModel, ConfigDict

from ..config import Settings, get_settings
from ..crexi import merge as cmerge
from ..db import dumps, loads, market_by_slug, now_iso, one, rows
from ..util import clean_str, normalize_address, normalize_county, normalize_zip, parse_date, to_float, yes_no
from .columns import (
    DEFAULT_FUZZY_MIN_SCORE,
    KEY_FIELDS,
    OWNER_FIELDS,
    PROVIDER_LABELS,
    check_overrides,
    detect_provider,
    field_label,
    known_header_score,
    looks_like_blue_moon,
    map_columns,
    norm_header,
    other_blue_moon_export,
)
from .edits import EDITABLE_BY_FIELD, edits_count_text, latest_edits, same_value

log = logging.getLogger(__name__)

Progress = Callable[[float, str], None] | None

HEADER_SCAN_ROWS = 15            # look this far down a sheet for the header row
PREFERRED_SHEET = "All properties"   # our own ranked_list.xlsx: every property (Ranked has the same columns)
BLUE_MOON_CREATOR = "Blue Moon Deal Finder"   # workbook metadata on our exports (bluemoon/exports.py)
MIN_PLAUSIBLE_YEAR = 1700        # a "year built" before this is a typo, not a year
EXCEL_EPOCH = date(1899, 12, 30)
CSV_EXTS = {".csv", ".tsv", ".txt"}
XLSX_EXTS = {".xlsx", ".xlsm"}
NULL_TOKENS = {"", "n/a", "na", "null", "none", "nan", "-", "--", "#n/a", "unknown", "not available"}

# Flags this module manages. Flags other modules add to flags_json are kept on re-import.
ROW_NOTES = frozenset({"year_invalid", "units_from_addresses", "shared_apn"})
INGEST_FLAGS = ROW_NOTES | frozenset({"year_unknown", "units_unknown", "no_owner_name", "no_coordinates",
                                      "no_address"})
STICKY_FLAGS = frozenset({"shared_apn"})
FLAG_HELP = {
    "year_unknown": "no year built",
    "year_invalid": "impossible year built (like 0 or 9999), left blank",
    "units_unknown": "no unit count",
    "units_from_addresses": "unit count taken from the number of addresses",
    "shared_apn": "parcel number shared with another row; check it",
    "no_owner_name": "no owner name",
    "no_coordinates": "no map coordinates (will be geocoded)",
    "no_address": "no street address",
}
SIZE_FIELDS = ("units", "building_sqft", "lot_acres", "building_count", "stories")

US_STATES = {
    "ALABAMA": "AL", "ALASKA": "AK", "ARIZONA": "AZ", "ARKANSAS": "AR", "CALIFORNIA": "CA", "COLORADO": "CO",
    "CONNECTICUT": "CT", "DELAWARE": "DE", "DISTRICT OF COLUMBIA": "DC", "FLORIDA": "FL", "GEORGIA": "GA",
    "HAWAII": "HI", "IDAHO": "ID", "ILLINOIS": "IL", "INDIANA": "IN", "IOWA": "IA", "KANSAS": "KS",
    "KENTUCKY": "KY", "LOUISIANA": "LA", "MAINE": "ME", "MARYLAND": "MD", "MASSACHUSETTS": "MA",
    "MICHIGAN": "MI", "MINNESOTA": "MN", "MISSISSIPPI": "MS", "MISSOURI": "MO", "MONTANA": "MT",
    "NEBRASKA": "NE", "NEVADA": "NV", "NEW HAMPSHIRE": "NH", "NEW JERSEY": "NJ", "NEW MEXICO": "NM",
    "NEW YORK": "NY", "NORTH CAROLINA": "NC", "NORTH DAKOTA": "ND", "OHIO": "OH", "OKLAHOMA": "OK",
    "OREGON": "OR", "PENNSYLVANIA": "PA", "RHODE ISLAND": "RI", "SOUTH CAROLINA": "SC", "SOUTH DAKOTA": "SD",
    "TENNESSEE": "TN", "TEXAS": "TX", "UTAH": "UT", "VERMONT": "VT", "VIRGINIA": "VA", "WASHINGTON": "WA",
    "WEST VIRGINIA": "WV", "WISCONSIN": "WI", "WYOMING": "WY", "PUERTO RICO": "PR",
}


class IngestError(ValueError):
    """A problem with the file itself. The message is written for Seth."""


# --- models -----------------------------------------------------------------------------------

class PropertyRecord(BaseModel):
    """One parsed row, in the shape of the `properties` table."""
    model_config = ConfigDict(extra="forbid", validate_assignment=True)

    source: str | None = None
    source_id: str | None = None
    crexi_url: str | None = None
    name: str | None = None
    property_type: str | None = None
    address: str | None = None
    unit: str | None = None
    city: str | None = None
    state: str | None = None
    zip: str | None = None
    county: str | None = None
    apn: str | None = None
    lat: float | None = None
    lng: float | None = None
    geocode_quality: str | None = None
    units: int | None = None
    building_sqft: float | None = None
    lot_acres: float | None = None
    building_count: int | None = None
    stories: int | None = None
    year_built: int | None = None
    sale_date: str | None = None
    sale_price: float | None = None
    lender: str | None = None
    loan_amount: float | None = None
    loan_type: str | None = None
    interest_rate: float | None = None
    loan_maturity: str | None = None
    improvement_value: float | None = None
    land_value: float | None = None
    total_value: float | None = None
    value_type: str | None = None
    tax_year: int | None = None
    annual_tax: float | None = None
    usps_vacant: int | None = None
    opportunity_zone: int | None = None
    preforeclosure: int | None = None
    reo: int | None = None
    owner_name: str | None = None
    mailing_address: str | None = None
    mailing_unit: str | None = None
    mailing_city: str | None = None
    mailing_state: str | None = None
    mailing_zip: str | None = None
    mailing_care_of: str | None = None


RECORD_COLUMNS: tuple[str, ...] = tuple(PropertyRecord.model_fields)


class ImportOptions(BaseModel):
    min_score: float = DEFAULT_FUZZY_MIN_SCORE
    units_from_addresses_min: int = 2
    units_metric: bool = False
    address_units: bool = True     # count addresses as units when there is no unit count (unit-sized classes only)
    year_in_buybox: bool = True    # flag a missing year built (only when the asset class's buy box uses it)
    today_year: int
    source: str = "crexi"


@dataclass
class ParsedRow:
    record: PropertyRecord
    raw: dict[str, Any]
    notes: set[str] = field(default_factory=set)
    units_source: str | None = None      # "units" | "addresses" | None
    apn_base: str | None = None
    street: str | None = None
    addr_key: str | None = None
    key: str | None = None
    dups: list["ParsedRow"] = field(default_factory=list)   # later rows of the same file that were this property
    fills: int = 0                                          # blanks of this row filled from those rows


@dataclass
class Table:
    filename: str
    headers: list[str]
    rows: list[dict[str, Any]]
    sheet: str | None = None
    header_row: int = 1
    skipped_columns: list[str] = field(default_factory=list)
    skipped_sheets: list[str] = field(default_factory=list)
    made_by_blue_moon: bool = False      # the workbook says this app wrote it (our own export)
    created: datetime | None = None      # our workbook's own time stamp (when it was exported; Excel keeps it)


# --- reading files ----------------------------------------------------------------------------

def read_table(path: str | Path) -> list[dict]:
    """Read a CSV or XLSX property list into a list of {original header: value} dicts.

    For XLSX the sheet that looks like a property list wins (e.g. "Properties", not "Key").
    Formula columns like "=IF(...)" are dropped. Dates come back as ISO strings.
    """
    return _read(Path(path)).rows


def _read(path: Path) -> Table:
    if not path.exists():
        raise IngestError(f"I can't find the file {path}.")
    ext = path.suffix.lower()
    if ext in XLSX_EXTS:
        return _read_xlsx(path)
    if ext in CSV_EXTS:
        return _read_csv(path)
    if ext == ".xls":
        raise IngestError(f"{path.name} is an old Excel .xls file. Open it in Excel or Google Sheets and save it "
                          f"as .xlsx or .csv, then import that.")
    raise IngestError(f"{path.name} isn't a CSV or Excel (.xlsx) file, so I can't import it.")


def _read_csv(path: Path, keep_formulas: bool = False) -> Table:
    text = _decode(path.read_bytes())
    sample = text[:65536]
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        delimiter = _count_delimiter(sample, path.suffix.lower())
    grid = [[_csv_cell(c) for c in r] for r in csv.reader(io.StringIO(text), delimiter=delimiter)]
    return _grid_to_table(path.name, grid, sheet=None, keep_formulas=keep_formulas)


def _count_delimiter(sample: str, ext: str) -> str:
    """Fallback when csv.Sniffer can't decide: the candidate seen most in the first lines."""
    lines = [ln for ln in sample.splitlines()[:20] if ln.strip()]
    counts = {d: sum(ln.count(d) for ln in lines) for d in (",", "\t", ";", "|")}
    best = max(counts, key=lambda d: counts[d])
    if counts[best] == 0:
        return "\t" if ext == ".tsv" else ","
    return best


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _csv_cell(v: str) -> Any:
    s = v.strip()
    if not s:
        return None
    m = re.fullmatch(r'="(.*)"', s)  # Excel's "keep my leading zeros" trick: ="00123"
    return m.group(1) if m else s


def _read_xlsx(path: Path) -> Table:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    except Exception as exc:  # corrupt file, password protected, not really xlsx
        raise IngestError(f"I couldn't open {path.name} as an Excel file ({exc.__class__.__name__}). "
                          f"Try exporting it again, or save it as .csv.") from exc
    try:
        best: tuple[int, str] | None = None
        titles = [ws.title for ws in wb.worksheets]
        for ws in wb.worksheets:
            head = [list(r) for r in ws.iter_rows(min_row=1, max_row=HEADER_SCAN_ROWS, values_only=True)]
            if not any(any(c is not None for c in r) for r in head):
                continue
            score = max((known_header_score(r) for r in head), default=0)
            # On a tie the first sheet wins, except our own export's 'All properties' (every row, not just
            # the ranked ones).
            if best is None or score > best[0] or (score == best[0] and ws.title == PREFERRED_SHEET):
                best = (score, ws.title)
        if best is None:
            raise IngestError(f"{path.name} has no rows in any sheet.")
        ws = wb[best[1]]
        grid = [[_xlsx_cell(c) for c in r] for r in ws.iter_rows(values_only=True)]
        ours = _made_by_blue_moon(wb)
    finally:
        wb.close()
    table = _grid_to_table(path.name, grid, sheet=best[1])
    table.skipped_sheets = [t for t in titles if t != best[1]]
    table.made_by_blue_moon = ours
    return table


def _made_by_blue_moon(wb: Any) -> bool:
    try:
        return str(getattr(wb.properties, "creator", "") or "").startswith(BLUE_MOON_CREATOR)
    except Exception:  # unreadable metadata never stops an import
        return False


def _workbook_created(wb: Any) -> datetime | None:
    try:
        created = wb.properties.created
        return created if isinstance(created, datetime) else None
    except Exception:  # unreadable metadata never stops an import
        return None


def read_sheets(path: str | Path) -> list[Table]:
    """Every sheet of an XLSX that looks like a list (a header row with known column names), or the one table
    of a CSV. Used for our own exports, whose property sheets (Ranked, All properties, Excluded) can each
    carry edits.

    There an empty cell clears a value, so a formula is never read as empty (as a Crexi import does): it reads
    as the result Excel saved with it (=36+2 is 38), or, with no saved result, as its text, which the round
    trip reports as a cell it can't read and leaves alone."""
    path = Path(path).expanduser()
    if path.suffix.lower() in CSV_EXTS and path.exists():
        return [_read_csv(path, keep_formulas=True)]
    if path.suffix.lower() not in XLSX_EXTS:
        return [_read(path)]
    import openpyxl

    if not path.exists():
        raise IngestError(f"I can't find the file {path}.")
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    except Exception as exc:
        raise IngestError(f"I couldn't open {path.name} as an Excel file ({exc.__class__.__name__}). "
                          f"Try exporting it again, or save it as .csv.") from exc
    grids: dict[str, list[list[Any]]] = {}
    try:
        ours = _made_by_blue_moon(wb)
        created = _workbook_created(wb) if ours else None
        for ws in wb.worksheets:
            grid = [[_xlsx_cell(c) for c in r] for r in ws.iter_rows(values_only=True)]
            if not any(any(c is not None for c in r) for r in grid[:HEADER_SCAN_ROWS]):
                continue
            if max((known_header_score(r) for r in grid[:HEADER_SCAN_ROWS]), default=0) == 0:
                continue
            grids[ws.title] = grid
    finally:
        wb.close()
    _use_saved_results(path, grids)
    out: list[Table] = []
    for title, grid in grids.items():
        table = _grid_to_table(path.name, grid, sheet=title, keep_formulas=True)
        table.made_by_blue_moon = ours
        table.created = created
        out.append(table)
    if not out:
        raise IngestError(f"{path.name} has no rows in any sheet.")
    return out


EXCEL_ERRORS = frozenset({"#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A", "#SPILL!", "#CALC!",
                          "#GETTING_DATA"})


def _use_saved_results(path: Path, grids: dict[str, list[list[Any]]]) -> None:
    """Put the result Excel saved with each formula in place of the formula's text. A formula with no saved
    result (a file no spreadsheet app has saved), or whose result is an error like #REF!, keeps its text.
    The workbook is opened a second time only when it holds formulas."""
    spots: dict[str, dict[int, list[int]]] = {}
    for title, grid in grids.items():
        for i, r in enumerate(grid):
            for j, c in enumerate(r):
                if _is_formula(c):
                    spots.setdefault(title, {}).setdefault(i, []).append(j)
    if not spots:
        return
    import openpyxl

    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception:  # the formulas keep their text: read as unreadable, never as empty
        return
    try:
        for title, rows_ in spots.items():
            last = max(rows_)
            for i, r in enumerate(wb[title].iter_rows(values_only=True)):
                if i > last:
                    break
                for j in rows_.get(i, ()):
                    v = _xlsx_cell(r[j]) if j < len(r) else None
                    if v is not None and not (isinstance(v, str) and v.upper() in EXCEL_ERRORS):
                        grids[title][i][j] = v
    finally:
        wb.close()


def _xlsx_cell(v: Any) -> Any:
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == time(0, 0) else v.isoformat(timespec="seconds")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, time):
        return v.isoformat()
    if isinstance(v, float) and math.isnan(v):
        return None
    if isinstance(v, str):
        s = v.strip()
        return s or None
    return v


def _is_formula(v: Any) -> bool:
    return isinstance(v, str) and v.startswith("=")


def _grid_to_table(filename: str, grid: list[list[Any]], sheet: str | None, keep_formulas: bool = False) -> Table:
    """The header row and the rows under it. A formula cell reads as empty and a column of only formulas is
    skipped, unless keep_formulas (our own files: read_sheets), where a formula keeps its text."""
    grid = [r for r in grid if any(c is not None for c in r)]
    if not grid:
        raise IngestError(f"{filename} is empty.")
    scan = grid[:HEADER_SCAN_ROWS]
    scores = [known_header_score(r) for r in scan]
    h_idx = scores.index(max(scores)) if max(scores) > 0 else 0
    header_cells, body = grid[h_idx], grid[h_idx + 1:]
    width = max(len(header_cells), max((len(r) for r in body), default=0))
    names = _header_names(header_cells, width)

    keep: list[int] = []
    skipped: list[str] = []
    for i in range(width):
        values = [r[i] for r in body if i < len(r) and r[i] is not None]
        blank_header = i >= len(header_cells) or clean_str(header_cells[i]) is None
        if blank_header and not values:
            continue
        if values and not keep_formulas and all(_is_formula(v) for v in values):
            skipped.append(names[i])
            continue
        keep.append(i)

    headers = [names[i] for i in keep]
    out_rows = []
    for r in body:
        row = {names[i]: (None if i >= len(r) or (_is_formula(r[i]) and not keep_formulas) else r[i]) for i in keep}
        if any(v is not None for v in row.values()):
            out_rows.append(row)
    return Table(filename=filename, headers=headers, rows=out_rows, sheet=sheet, header_row=h_idx + 1,
                 skipped_columns=skipped)


def _header_names(cells: list[Any], width: int) -> list[str]:
    names: list[str] = []
    seen: Counter[str] = Counter()
    for i in range(width):
        base = clean_str(cells[i]) if i < len(cells) else None
        base = re.sub(r"\s+", " ", base) if base else f"Column {i + 1}"
        seen[base] += 1
        names.append(base if seen[base] == 1 else f"{base} ({seen[base]})")
    return names


# --- value parsing ----------------------------------------------------------------------------

def _text(v: Any) -> str | None:
    if v is None or isinstance(v, bool):
        return None if v is None else ("Yes" if v else "No")
    if isinstance(v, float):
        if math.isnan(v):
            return None
        return str(int(v)) if v.is_integer() else str(v)
    s = clean_str(v)
    if s is None or s.lower() in NULL_TOKENS:
        return None
    return re.sub(r"\s+", " ", s)


_NUM_WITH_UNIT = re.compile(r"\$?\s*(-?[\d,]*\.?\d+)\s*[A-Za-z%][A-Za-z%.\s/]*")


def _num(v: Any) -> float | None:
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return None if isinstance(v, float) and (math.isnan(v) or math.isinf(v)) else float(v)
    s = _text(v)
    if s is None:
        return None
    f = to_float(s)
    if f is None:
        m = _NUM_WITH_UNIT.fullmatch(s)
        f = to_float(m.group(1)) if m else None
    return f


def _pos(v: Any) -> float | None:
    """A positive number, else None (0 means "not recorded" in these exports)."""
    f = _num(v)
    return f if f is not None and f > 0 else None


def _pos_int(v: Any) -> int | None:
    f = _pos(v)
    return int(round(f)) if f is not None else None


def _date(v: Any) -> str | None:
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return (EXCEL_EPOCH + timedelta(days=int(v))).isoformat() if 20000 <= v <= 80000 else None
    s = _text(v)
    if s is None:
        return None
    if s.isdigit():
        if len(s) == 8:  # 20210315
            try:
                return datetime.strptime(s, "%Y%m%d").date().isoformat()
            except ValueError:
                return None
        return _date(int(s))  # an Excel day number saved as text
    if re.match(r"^\d{4}-\d{2}-\d{2}", s):
        try:
            return date.fromisoformat(s[:10]).isoformat()
        except ValueError:
            return None
    d = parse_date(s)
    return d.isoformat() if d else None


def _year(v: Any, today_year: int) -> tuple[int | None, bool]:
    """(year, invalid). 0, 9999, future years and typos become (None, True); blank is (None, False)."""
    f = _num(v)
    if f is None:
        return None, _text(v) is not None
    if not float(f).is_integer():
        return None, True
    y = int(f)
    if y < MIN_PLAUSIBLE_YEAR or y > today_year:
        return None, True
    return y, False


def _state(v: Any) -> str | None:
    s = _text(v)
    if s is None:
        return None
    s = s.upper().replace(".", "").strip()
    if len(s) == 2 and s.isalpha():
        return s
    return US_STATES.get(s, s)


def _apn(v: Any) -> str | None:
    s = _text(v)
    if s is None:
        return None
    s = s.upper()
    alnum = re.sub(r"[^0-9A-Z]", "", s)
    if not alnum or set(alnum) == {"0"}:
        return None
    return s


def apn_compact(apn: str | None) -> str | None:
    """Letters and digits only, uppercase: the form used in property keys."""
    if not apn:
        return None
    return re.sub(r"[^0-9A-Z]", "", apn.upper()) or None


def _value_type(v: Any) -> str | None:
    s = _text(v)
    if s is None:
        return None
    if s.startswith("["):
        try:
            items = json.loads(s)
            if isinstance(items, list):
                s = "/".join(str(x).strip() for x in items if str(x).strip())
        except ValueError:
            s = s.strip("[]").replace('"', "").strip()
    return s.upper() or None


def _flag(v: Any) -> int | None:
    s = _text(v)
    return yes_no(s) if s is not None else None


def _filled_not_no(v: Any) -> bool:
    s = _text(v)
    return s is not None and yes_no(s) != 0


def _coords(lat_v: Any, lng_v: Any) -> tuple[float | None, float | None]:
    lat, lng = _num(lat_v), _num(lng_v)
    if lat is None or lng is None:
        return None, None
    if not (-90 <= lat <= 90 and -180 <= lng <= 180) or (abs(lat) < 1e-9 and abs(lng) < 1e-9):
        return None, None
    return lat, lng


def source_id_from_link(url: str | None) -> str | None:
    """The record id inside a Crexi link.

    /property-records/{slug}/{hash} -> hash; /search?selectedId=X -> X; /properties/{id}/{slug} -> id.
    """
    s = _text(url)
    if not s or "/" not in s:
        return None
    parts = urlsplit(s if "://" in s else f"https://{s}")
    query = parse_qs(parts.query)
    if query.get("selectedId"):
        return query["selectedId"][0].strip() or None
    segs = [p for p in parts.path.split("/") if p]
    if not segs:
        return None
    if segs[0] == "properties" and len(segs) > 1 and segs[1].isdigit():
        return segs[1]
    return segs[-1] if len(segs) > 1 else None


# --- rows -> records --------------------------------------------------------------------------

_TEXT_FIELDS = ("name", "property_type", "address", "unit", "city", "lender", "loan_type", "owner_name",
                "mailing_address", "mailing_unit", "mailing_city", "mailing_care_of")
_POS_FLOAT_FIELDS = ("building_sqft", "sale_price", "loan_amount", "interest_rate", "improvement_value",
                     "land_value", "total_value", "annual_tax")
_POS_INT_FIELDS = ("building_count", "stories", "tax_year")
_DATE_FIELDS = ("sale_date", "loan_maturity")


def parse_row(row: Mapping[str, Any], column_map: Mapping[str, str], opts: ImportOptions) -> ParsedRow:
    """Turn one original row into a ParsedRow (record + raw + notes + key parts)."""
    def get(f: str) -> Any:
        h = column_map.get(f)
        return row.get(h) if h else None

    data: dict[str, Any] = {f: _text(get(f)) for f in _TEXT_FIELDS}
    data.update({f: _pos(get(f)) for f in _POS_FLOAT_FIELDS})
    data.update({f: _pos_int(get(f)) for f in _POS_INT_FIELDS})
    data.update({f: _date(get(f)) for f in _DATE_FIELDS})
    data.update(
        state=_state(get("state")), mailing_state=_state(get("mailing_state")),
        zip=normalize_zip(_text(get("zip"))), mailing_zip=normalize_zip(_text(get("mailing_zip"))),
        county=normalize_county(_text(get("county"))), apn=_apn(get("apn")),
        value_type=_value_type(get("value_type")),
        usps_vacant=_flag(get("usps_vacant")), opportunity_zone=_flag(get("opportunity_zone")),
    )
    notes: set[str] = set()
    _fill_owner(data, get)
    _fill_location(data, get)
    _fill_distress(data, get, column_map)
    year, invalid = _year(get("year_built"), opts.today_year)
    data["year_built"] = year
    if invalid:
        notes.add("year_invalid")
    units_source = _fill_units(data, get, opts)
    if units_source == "addresses":
        notes.add("units_from_addresses")
    lot_acres = _pos(get("lot_acres"))
    lot_sqft = _pos(get("lot_sqft"))
    data["lot_acres"] = lot_acres if lot_acres is not None else (round(lot_sqft / 43560.0, 4) if lot_sqft else None)
    _fill_source(data, get, opts)
    if _filled_not_no(get("shared_apn_flag")):
        notes.add("shared_apn")

    record = PropertyRecord(**data)
    parsed = ParsedRow(record=record, raw=dict(row), notes=notes, units_source=units_source)
    _key_parts(parsed)
    return parsed


def _fill_owner(data: dict[str, Any], get: Callable[[str], Any]) -> None:
    if data["owner_name"] is None:
        parts = [_text(get("owner_first_name")), _text(get("owner_last_name"))]
        data["owner_name"] = " ".join(p for p in parts if p) or None


def _fill_location(data: dict[str, Any], get: Callable[[str], Any]) -> None:
    lat, lng = _coords(get("lat"), get("lng"))
    data["lat"], data["lng"] = lat, lng
    data["geocode_quality"] = "source" if lat is not None else None


def _fill_distress(data: dict[str, Any], get: Callable[[str], Any], column_map: Mapping[str, str]) -> None:
    pfc_fields = ("pfc_recording_date", "pfc_indicator", "pfc_document_type")
    if any(f in column_map for f in pfc_fields):
        data["preforeclosure"] = int(any(_filled_not_no(get(f)) for f in pfc_fields))
    else:
        data["preforeclosure"] = None
    if "reo" in column_map:
        flag = _flag(get("reo"))
        # Crexi leaves the REO flag blank for "no".
        data["reo"] = flag if flag is not None else (1 if _text(get("reo")) else 0)
    else:
        data["reo"] = None


def _fill_units(data: dict[str, Any], get: Callable[[str], Any], opts: ImportOptions) -> str | None:
    units = _pos_int(get("units"))
    if units is not None:
        data["units"] = units
        return "units"
    addresses = _pos_int(get("number_of_addresses"))
    if opts.address_units and addresses is not None and addresses >= opts.units_from_addresses_min:
        data["units"] = addresses
        return "addresses"
    data["units"] = None
    return None


def _fill_source(data: dict[str, Any], get: Callable[[str], Any], opts: ImportOptions) -> None:
    url = _text(get("crexi_url"))
    data["crexi_url"] = url
    data["source_id"] = _text(get("source_id")) or source_id_from_link(url)
    is_crexi = bool(url and "crexi.com" in url.lower())
    data["source"] = "crexi" if is_crexi else opts.source


def _key_parts(p: ParsedRow) -> None:
    r = p.record
    compact = apn_compact(r.apn)
    p.apn_base = f"apn:{r.state or ''}:{r.county or ''}:{compact}" if compact else None
    p.street = normalize_address(r.address)
    p.addr_key = address_key(r.address, r.unit, r.city, r.state, r.zip)


def address_key(address: str | None, unit: str | None, city: str | None, state: str | None,
                zip_: str | None) -> str | None:
    """`addr:{normalized address}` (unit included when it isn't already part of the street line)."""
    address, unit = clean_str(address), clean_str(unit)
    if not address:
        return None
    line = address
    if unit and not address.upper().endswith(unit.upper()):
        line = f"{address} {unit}"
    return f"addr:{normalize_address(line, city, state, zip_)}"


def assign_keys(parsed: list[ParsedRow]) -> list[ParsedRow]:
    """Set .key on each row. One parcel with two different street addresses keeps both rows."""
    groups: dict[str, list[ParsedRow]] = defaultdict(list)
    for p in parsed:
        if p.apn_base:
            groups[p.apn_base].append(p)
        elif p.addr_key:
            p.key = p.addr_key
        elif p.record.source_id:
            p.key = f"src:{p.record.source or 'list'}:{p.record.source_id}"
    for base, items in groups.items():
        streets = {p.street for p in items if p.street}
        shared = len(streets) > 1
        for p in items:
            if shared:
                p.notes.add("shared_apn")
                p.key = f"{base}|{p.street}" if p.street else base
            else:
                p.key = base
    return parsed


def collapse_duplicates(parsed: list[ParsedRow]) -> tuple[list[ParsedRow], int]:
    """Merge rows with the same key. The first row is the original: blanks of it are filled from the later rows,
    and any value a later row disagrees about is kept as an alternate once the property is saved
    (crexi/merge.py). Returns (rows, merged)."""
    by_key: dict[str, ParsedRow] = {}
    merged = 0
    for p in parsed:
        if p.key is None:
            continue
        if p.key in by_key:
            _combine(by_key[p.key], p)
            merged += 1
        else:
            by_key[p.key] = p
    return list(by_key.values()), merged


def _combine(a: ParsedRow, b: ParsedRow) -> None:
    ra, rb = a.record, b.record
    if ra.source_id and rb.source_id and ra.source_id != rb.source_id and a.apn_base:
        a.notes.add("shared_apn")  # Crexi listed this parcel on more than one record
    updates: dict[str, Any] = {}
    for f in RECORD_COLUMNS:
        if f == "units":
            continue  # handled by _combine_units
        va, vb = getattr(ra, f), getattr(rb, f)
        if va is None and vb is not None:
            updates[f] = vb
            a.fills += 1
    if updates:
        a.record = ra.model_copy(update=updates)
    a.dups.append(b)
    _combine_units(a, b)
    a.notes |= b.notes - {"units_from_addresses"}
    for k, v in b.raw.items():
        if a.raw.get(k) is None and v is not None:
            a.raw[k] = v


def _combine_units(a: ParsedRow, b: ParsedRow) -> None:
    """Prefer a real unit count over one taken from the address count; otherwise the first row's count stays
    (a different count in a later row becomes an alternate)."""
    rank = {"units": 2, "addresses": 1, None: 0}
    ua, ub = a.record.units, b.record.units
    if rank[b.units_source] > rank[a.units_source] or (
            b.units_source == a.units_source and ub is not None and ua is None):
        a.record = a.record.model_copy(update={"units": ub})
        a.units_source = b.units_source
    a.notes.discard("units_from_addresses")
    if a.units_source == "addresses":
        a.notes.add("units_from_addresses")


def final_flags(rec: PropertyRecord, notes: Iterable[str], units_metric: bool,
                year_in_buybox: bool = True) -> list[str]:
    """Flags for flags_json, worked out from the record's final values. A missing year built is only flagged
    when the asset class's buy box uses the year."""
    flags = set(notes) & ROW_NOTES
    if rec.year_built is not None:
        flags.discard("year_invalid")
    elif "year_invalid" not in flags and year_in_buybox:
        flags.add("year_unknown")
    if rec.units is None:
        flags.discard("units_from_addresses")
        if units_metric:
            flags.add("units_unknown")
    if not rec.owner_name:
        flags.add("no_owner_name")
    if rec.lat is None or rec.lng is None:
        flags.add("no_coordinates")
    if not rec.address:
        flags.add("no_address")
    return sorted(flags)


# --- options ---------------------------------------------------------------------------------

def _options(settings: Settings, asset_class: str | None, provider: str) -> ImportOptions:
    cfg = settings.section("crexi")
    units_metric, address_units, year_in_buybox = False, True, True
    if asset_class:
        try:
            ac = settings.asset_class(asset_class)
            units_metric = ac.size.metric == "units"
            address_units = units_metric           # sq ft and acre classes never get units from addresses
            year_in_buybox = bool(ac.year_built_min)
        except KeyError:
            units_metric = False
    min_score = cfg.get("fuzzy_min_score")
    addr_min = cfg.get("units_from_addresses_min")
    return ImportOptions(
        min_score=float(min_score) if min_score is not None else DEFAULT_FUZZY_MIN_SCORE,
        units_from_addresses_min=int(addr_min) if addr_min is not None else 2,
        units_metric=units_metric,
        address_units=address_units,
        year_in_buybox=year_in_buybox,
        today_year=date.today().year,
        source="crexi" if provider == "crexi" else (provider if provider != "other" else "list"),
    )


def _key_fields(settings: Settings, asset_class: str | None) -> tuple[list[str], str | None]:
    """(the fields the import screen checks for, the size metric) for the market's asset class: its size
    measure (units, sq ft or acres) and the year built only when its buy box uses them."""
    try:
        ac = settings.asset_class(asset_class) if asset_class else None
    except KeyError:
        ac = None
    if ac is None:
        return list(KEY_FIELDS), None
    base = [f for f in KEY_FIELDS if f not in ("units", "building_sqft", "year_built")]
    extra = [ac.size.metric] + ([ac.size.fallback] if ac.size.fallback else [])
    extra += ["year_built"] if ac.year_built_min else []
    at = base.index("lng") + 1 if "lng" in base else len(base)
    return base[:at] + list(dict.fromkeys(extra)) + base[at:], ac.size.metric


TYPE_CHECK_MIN_ROWS = 5   # rows with a Property Type before the import screen checks the market


def _type_mismatch(settings: Settings, asset_class: str | None, table: Table,
                   column_map: Mapping[str, str]) -> dict | None:
    """When fewer than half the rows with a Property Type fit this market's asset class, say which asset class
    (and market) they look like, so a file never lands in the wrong market by accident."""
    col = column_map.get("property_type")
    if not asset_class or not col:
        return None
    try:
        ac = settings.asset_class(asset_class)
    except KeyError:
        return None
    if not ac.crexi_property_types:
        return None
    # Imported here: the scoring package loads more than the importer needs.
    from ..scoring.buybox import property_type_matches
    types = [t for t in (_text(r.get(col)) for r in table.rows) if t]
    if len(types) < TYPE_CHECK_MIN_ROWS:
        return None
    matched = sum(1 for t in types if property_type_matches(t, ac))
    if matched * 2 >= len(types):
        return None
    best, best_n = None, matched
    for key, other in settings.asset_classes.items():
        if key == ac.key or not other.crexi_property_types:
            continue
        n = sum(1 for t in types if property_type_matches(t, other))
        if n > best_n:
            best, best_n = other, n
    best_market = next((m.slug for m in settings.markets if best and m.asset_class == best.key), None)
    msg = f"Only {matched:,} of {len(types):,} rows are {ac.label} (Property Type column)."
    if best:
        msg += f" They look like {best.label} ({best_n:,} rows)."
        if best_market:
            msg += f" Switch to the {best.label} market before importing."
    msg += (" Here they would all be set aside as 'Not this kind of property', and imports can't be undone.")
    return {"label": ac.label, "matched": matched, "typed": len(types), "best_label": best.label if best else None,
            "best_matched": best_n if best else None, "best_market": best_market, "message": msg}


def _merged_overrides(settings: Settings, overrides: Mapping[str, str | None] | None) -> dict[str, str | None]:
    cfg = settings.section("crexi").get("column_overrides") or {}
    if not isinstance(cfg, dict):
        log.warning("config.yaml crexi.column_overrides should be a mapping of field: column name; ignored")
        cfg = {}
    return {**cfg, **dict(overrides or {})}


def _check_identity(table: Table, column_map: Mapping[str, str]) -> None:
    if not any(f in column_map for f in ("apn", "address", "crexi_url", "source_id")):
        shown = ", ".join(table.headers[:25]) + (" ..." if len(table.headers) > 25 else "")
        raise IngestError(
            f"{table.filename} doesn't look like a property list: I couldn't find an address or parcel number "
            f"column. The columns I saw: {shown}. If one of them is the address, add it in config.yaml under "
            f"crexi.column_overrides (for example  address: \"Site Addr\").")


# --- public: preview ------------------------------------------------------------------------

def preview_file(path: str | Path, overrides: Mapping[str, str | None] | None = None,
                 market_slug: str | None = None, conn: sqlite3.Connection | None = None) -> dict:
    """What the import screen shows before importing: sheet, column matches and 5 sample rows.

    The sample never includes owner names or mailing addresses. A Blue Moon export (our own file, edited in
    Excel) gets the edits preview instead: how many properties would change, and what (conn is used to
    compare; without one the app's database is opened).
    """
    settings = get_settings()
    table = _read(Path(path).expanduser())
    provider = _provider(table)
    if provider == "bluemoon":
        from .roundtrip import preview_blue_moon_edits
        return preview_blue_moon_edits(conn, market_slug, Path(path).expanduser())
    asset_class = _asset_class_for(settings, market_slug)
    opts = _options(settings, asset_class, provider)
    all_overrides = _merged_overrides(settings, overrides)
    column_map, unmapped = map_columns(table.headers, all_overrides, opts.min_score)
    sample = []
    for row in table.rows[:5]:
        p = parse_row(row, column_map, opts)
        item = {k: v for k, v in p.record.model_dump().items() if k not in OWNER_FIELDS}
        item["flags"] = final_flags(p.record, p.notes, opts.units_metric, opts.year_in_buybox)
        sample.append(item)
    key_fields, size_metric = _key_fields(settings, asset_class)
    have = set(column_map) | ({"lot_acres"} if "lot_sqft" in column_map else set())   # acres come from lot sq ft
    mismatch = _type_mismatch(settings, asset_class, table, column_map)
    warnings = _read_warnings(table) + check_overrides(table.headers, all_overrides)
    return {
        "filename": table.filename,
        "sheet": table.sheet,
        "provider": provider,
        "provider_label": PROVIDER_LABELS.get(provider, provider),
        "headers": table.headers,
        "column_map": column_map,
        "column_labels": {f: field_label(f) for f in column_map},
        "unmapped_columns": unmapped,
        "missing_fields": [{"field": f, "label": field_label(f)} for f in key_fields if f not in have],
        "key_fields": key_fields,
        "size_metric": size_metric,
        "type_mismatch": mismatch,
        "rows": len(table.rows),
        "sample": sample,
        "skipped_columns": table.skipped_columns,
        "warnings": ([mismatch["message"]] if mismatch else []) + warnings,
    }


def _provider(table: Table) -> str:
    """detect_provider, plus our own workbooks known by their metadata (with a Property ID column).
    IngestError for our own files that can't come back: one of the other exports (DRILL worklist, mailing
    list, CRM handoff), or a property list whose Property ID column was deleted. Through the Crexi merge
    they would add duplicates and write our columns into the original data."""
    provider = detect_provider(table.headers)
    if provider == "bluemoon":
        return provider
    other = other_blue_moon_export(table.headers)
    if other:
        raise IngestError(f"This is Blue Moon's {other}, one of its own exports, not a property list, so it can't "
                          f"be imported. To bring fixes back, import ranked_list or a review file.")
    has_id = any(norm_header(h) == "property id" for h in table.headers)
    if table.made_by_blue_moon and has_id:
        return "bluemoon"
    if not has_id and (table.made_by_blue_moon or looks_like_blue_moon(table.headers)):
        raise IngestError("This is one of Blue Moon's own files, but its Property ID column is missing, so I can't "
                          "tell which property each row is. Download a fresh copy, make your changes there, and "
                          "leave the Property ID and Row as of columns in place.")
    return provider


def _asset_class_for(settings: Settings, market_slug: str | None) -> str | None:
    if not market_slug:
        return None
    try:
        return settings.market(market_slug).asset_class
    except KeyError:
        return None


def _read_warnings(table: Table) -> list[str]:
    out = []
    if table.sheet and table.skipped_sheets:
        names = ", ".join(f"'{s}'" for s in table.skipped_sheets)
        out.append(f"Read the '{table.sheet}' sheet; skipped {names} (not a property list).")
    if table.skipped_columns:
        names = ", ".join(f"'{c}'" for c in table.skipped_columns)
        out.append(f"Ignored formula column(s) {names}; they are worked out from other columns.")
    if table.header_row > 1:
        out.append(f"Column names were on row {table.header_row}; skipped the rows above them.")
    return out


# --- public: import -------------------------------------------------------------------------

def import_file(conn: sqlite3.Connection, market_slug: str, path: str | Path, limit: int | None = None,
                overrides: Mapping[str, str | None] | None = None, progress: Progress = None,
                actor: str | None = None) -> dict:
    """Import one Crexi (or other provider) export into a market. Safe to re-run: rows merge, never delete.

    Returns {import_id, filename, rows_total, rows_new, rows_updated, rows_unchanged, rows_skipped, column_map,
    unmapped_columns, warnings, ...extra counts}. rows_unchanged is the part of rows_updated whose stored values
    did not change (re-importing the same rows). Corrections made by hand (property_edits) are put back after
    the merge (edits_kept counts the ones the file disagreed with).

    A Blue Moon export (one of our own files, edited in Excel) never goes through the merge: it goes to
    roundtrip.import_blue_moon_edits, which changes only the editable columns of known properties.
    actor: who is importing (the dashboard's signed-in person; 'cli' when None).
    """
    _progress(progress, 0.02, "Reading the file")
    settings = get_settings()
    try:
        market = market_by_slug(conn, market_slug)
    except KeyError:
        known = ", ".join(r["slug"] for r in rows(conn, "SELECT slug FROM markets ORDER BY id"))
        raise IngestError(f"There is no market called '{market_slug}'. Markets: {known}.") from None
    path = Path(path).expanduser()
    table = _read(path)
    provider = _provider(table)
    if provider == "bluemoon":
        from .roundtrip import import_blue_moon_edits
        return import_blue_moon_edits(conn, market_slug, path, limit=limit, actor=actor, progress=progress)
    opts = _options(settings, market["asset_class"], provider)
    all_overrides = _merged_overrides(settings, overrides)
    column_map, unmapped = map_columns(table.headers, all_overrides, opts.min_score)
    _check_identity(table, column_map)

    data_rows = table.rows if limit is None else table.rows[:max(0, int(limit))]
    _progress(progress, 0.15, f"Checking {len(data_rows):,} rows")
    parsed = [parse_row(r, column_map, opts) for r in data_rows]
    assign_keys(parsed)
    no_identity = sum(1 for p in parsed if p.key is None)
    records, duplicates = collapse_duplicates(parsed)

    file_hash = _file_hash(path)
    mismatch = _type_mismatch(settings, market["asset_class"], table, column_map)
    warnings = ([mismatch["message"]] if mismatch else []) + _read_warnings(table)
    warnings += check_overrides(table.headers, all_overrides)
    warnings += _previous_import_warning(conn, market["id"], file_hash)

    _progress(progress, 0.3, f"Saving {len(records):,} properties")
    with conn:
        import_id = _start_import(conn, market["id"], table.filename, file_hash, column_map)
        merge = _merge_records(conn, market["id"], records, import_id, opts, progress)
        skipped = no_identity + duplicates
        warnings += _row_warnings(no_identity, duplicates, parsed, merge["flag_counts"], unmapped)
        warnings += _kept_edits_warnings(merge["edits_kept"], merge["edits_kept_fields"], merge["edits_crexi_new"])
        conn.execute(
            "UPDATE imports SET rows_total=?, rows_new=?, rows_updated=?, rows_skipped=?, notes=? WHERE id=?",
            (len(data_rows), merge["new"], merge["updated"], skipped, "\n".join(warnings), import_id))

    _progress(progress, 0.97, "Setting up DRILL levels")
    _ensure_drill_states(conn, market_slug, warnings)
    _progress(progress, 1.0, "Import finished")
    enrichment = merge["report"]
    enrichment.add("rows_read", len(data_rows))
    enrichment.add("matched", merge["updated"])
    enrichment_dict = enrichment.to_dict()
    enrichment_dict["duplicates"] = max(enrichment_dict["duplicates"], duplicates)
    log.info("Imported %s into %s: %d rows, %d new, %d updated (%d of them unchanged), %d skipped", table.filename,
             market_slug, len(data_rows), merge["new"], merge["updated"], merge["unchanged"], skipped)
    return {
        "import_id": import_id,
        "filename": table.filename,
        "market": market_slug,
        "provider": provider,
        "sheet": table.sheet,
        "rows_total": len(data_rows),
        "rows_new": merge["new"],
        "rows_updated": merge["updated"],
        "rows_unchanged": merge["unchanged"],   # part of rows_updated: already stored with these exact values
        "rows_skipped": skipped,
        "duplicates_collapsed": duplicates,
        "rows_without_id": no_identity,
        "column_map": column_map,
        "unmapped_columns": unmapped,
        "skipped_columns": table.skipped_columns,
        "mapped_columns": len(table.headers) - len(unmapped),
        "total_columns": len(table.headers),
        "flag_counts": merge["flag_counts"],
        "edits_kept": merge["edits_kept"],
        "enrichment": enrichment_dict,
        "enrichment_text": cmerge.report_text(enrichment_dict),
        "warnings": warnings,
    }


def _kept_edits_warnings(kept: int, fields: Mapping[str, int], crexi_new: int) -> list[str]:
    """'Kept 3 values you corrected by hand ...': a Crexi file never quietly undoes a correction."""
    if not kept:
        return []
    out = [f"Kept {_plural(kept, 'value', 'values')} you corrected by hand ({edits_count_text(fields)}) instead "
           f"of this file's {'value' if kept == 1 else 'values'}. The file's values are still saved with the "
           f"original columns."]
    if crexi_new:
        out.append(f"For {_plural(crexi_new, 'of them', 'of them')} Crexi now shows a different value than when you "
                   f"corrected it. Your correction still counts; look at those properties if you like.")
    return out


def _progress(progress: Progress, fraction: float, message: str) -> None:
    if progress:
        try:
            progress(fraction, message)
        except Exception:  # a broken progress callback must never break an import
            log.debug("progress callback failed", exc_info=True)


def _file_hash(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _previous_import_warning(conn: sqlite3.Connection, market_id: int, file_hash: str) -> list[str]:
    prev = one(conn, "SELECT id, created_at FROM imports WHERE market_id=? AND file_hash=? ORDER BY id DESC LIMIT 1",
               (market_id, file_hash))
    if not prev:
        return []
    return [f"This exact file was imported before (import #{prev['id']} on {str(prev['created_at'])[:10]}). "
            f"Rows were merged again; nothing was duplicated."]


def _start_import(conn: sqlite3.Connection, market_id: int, filename: str, file_hash: str,
                  column_map: Mapping[str, str]) -> int:
    cur = conn.execute(
        "INSERT INTO imports(market_id, filename, file_hash, column_map_json, created_at) VALUES(?,?,?,?,?)",
        (market_id, filename, file_hash, dumps(dict(column_map)), now_iso()))
    return int(cur.lastrowid)


def _plural(n: int, one_word: str, many_word: str) -> str:
    return f"{n:,} {one_word if n == 1 else many_word}"


def _row_warnings(no_identity: int, duplicates: int, parsed: list[ParsedRow], flag_counts: Mapping[str, int],
                  unmapped: list[str]) -> list[str]:
    out = []
    if no_identity:
        out.append(f"{_plural(no_identity, 'row', 'rows')} had no parcel number, address or Crexi link, "
                   f"so I skipped them.")
    if duplicates:
        out.append(f"{_plural(duplicates, 'duplicate row', 'duplicate rows')} inside this file were combined "
                   f"with their match.")
    # A parcel listed with more than one street address is kept as one property per address (key
    # 'apn:...|street'). Address keys also contain '|', so only parcel keys count here.
    split = {p.key for p in parsed if p.apn_base and p.key and p.key.startswith(p.apn_base + "|")}
    if split:
        parcels = len({k.split("|", 1)[0] for k in split})
        out.append(f"{_plural(parcels, 'parcel number was', 'parcel numbers were')} listed with more than one "
                   f"street address. Each address was kept as its own property ({len(split):,} in all) and "
                   f"flagged so you can check them.")
    if flag_counts.get("year_invalid"):
        out.append(f"{_plural(flag_counts['year_invalid'], 'property', 'properties')} had an impossible year "
                   f"built (like 0 or 9999); left it blank.")
    if flag_counts.get("units_from_addresses"):
        out.append(f"{_plural(flag_counts['units_from_addresses'], 'property', 'properties')} had no unit count, "
                   f"so I used the number of addresses as an estimate.")
    if unmapped:
        out.append(f"{_plural(len(unmapped), 'column', 'columns')} didn't match a field and were kept as-is in the "
                   f"original data: " + ", ".join(unmapped[:10]) + (" ..." if len(unmapped) > 10 else ""))
    return out


def _ensure_drill_states(conn: sqlite3.Connection, market_slug: str, warnings: list[str]) -> None:
    try:
        from ..drill.engine import ensure_states
        ensure_states(conn, market_slug)
    except Exception:
        log.exception("ensure_states failed after import")
        warnings.append("The properties were imported, but DRILL levels weren't set up yet; "
                        "they will be on the next DRILL run.")


# --- merge into the database ---------------------------------------------------------------

@dataclass
class _Existing:
    by_key: dict[str, dict] = field(default_factory=dict)
    by_base: dict[str, list[dict]] = field(default_factory=lambda: defaultdict(list))
    by_addr: dict[str, dict] = field(default_factory=dict)
    by_source: dict[tuple[str, str], dict] = field(default_factory=dict)

    def add(self, row: dict) -> None:
        key = row["property_key"]
        self.by_key[key] = row
        if key.startswith("apn:"):
            self.by_base[key.split("|", 1)[0]].append(row)
        addr = _addr_key_of(row)
        if addr:
            self.by_addr.setdefault(addr, row)
        if row.get("source") and row.get("source_id"):
            self.by_source.setdefault((row["source"], row["source_id"]), row)

    def rekey(self, row: dict, new_key: str) -> None:
        self.by_key.pop(row["property_key"], None)
        row["property_key"] = new_key
        self.by_key[new_key] = row
        if new_key.startswith("apn:"):
            base = self.by_base[new_key.split("|", 1)[0]]
            if row not in base:
                base.append(row)


def _addr_key_of(row: Mapping[str, Any]) -> str | None:
    return address_key(row.get("address"), row.get("unit"), row.get("city"), row.get("state"), row.get("zip"))


def _load_existing(conn: sqlite3.Connection, market_id: int) -> _Existing:
    ex = _Existing()
    for r in rows(conn, "SELECT * FROM properties WHERE market_id=? ORDER BY id", (market_id,)):
        ex.add(r)
    return ex


def _find_match(ex: _Existing, p: ParsedRow, taken: set[int]) -> dict | None:
    """Find the stored property this row is. Exact key first, then same parcel + street, address, source id."""
    hit = ex.by_key.get(p.key or "")
    if hit is not None and hit["id"] not in taken:
        return hit
    compact = apn_compact(p.record.apn)
    if p.apn_base:
        for cand in ex.by_base.get(p.apn_base, []):
            if cand["id"] not in taken and p.street and normalize_address(cand.get("address")) == p.street:
                return cand
    for cand in (ex.by_addr.get(p.addr_key or ""),
                 ex.by_source.get((p.record.source or "", p.record.source_id or ""))):
        if cand is None or cand["id"] in taken:
            continue
        cand_apn = apn_compact(cand.get("apn"))
        if not compact or not cand_apn or cand_apn == compact:
            return cand
    return None


def _stored(row: Mapping[str, Any]) -> dict[str, Any]:
    """What an import can change on a stored property (not the import ids or the time), for comparing."""
    out = {c: row.get(c) for c in RECORD_COLUMNS}
    out["property_key"] = row.get("property_key")
    out["flags"] = sorted(loads(row.get("flags_json"), []) or [])
    out["raw"] = loads(row.get("raw_json"), {}) or {}
    return out


def _merge_records(conn: sqlite3.Connection, market_id: int, records: list[ParsedRow], import_id: int,
                   opts: ImportOptions, progress: Progress) -> dict:
    """Insert or merge every record. Returns {new, updated, unchanged, flag_counts, edits_kept,
    edits_kept_fields, edits_crexi_new}: `unchanged` counts the updated properties whose stored values came out
    exactly as they were (a re-import of the same rows); edits_* count the corrections made by hand that this
    file disagreed with and that were kept."""
    ex = _load_existing(conn, market_id)
    edits = latest_edits(conn, market_id)
    kept: Counter[str] = Counter()
    taken: set[int] = set()
    before: dict[int, tuple[dict, dict]] = {}     # property id -> (stored values before, the live row)
    new = updated = 0
    flag_counts: Counter[str] = Counter()
    report = cmerge.Report()
    now = now_iso()
    step = max(1, len(records) // 20)
    for i, p in enumerate(records):
        match = _find_match(ex, p, taken)
        if match is None and p.key and p.key in ex.by_key:
            # An earlier row of this file already matched this property (for example once without its APN
            # and once with it): merge into it rather than insert a second row with the same key.
            match = ex.by_key[p.key]
        if match is None:
            flags = final_flags(p.record, p.notes, opts.units_metric, opts.year_in_buybox)
            row = _insert(conn, market_id, p, flags, import_id, now)
            ex.add(row)
            new += 1
        else:
            if match["id"] not in taken:
                before[match["id"]] = (_stored(match), match)
            if match["property_key"] != p.key and p.key and p.key not in ex.by_key and _upgrade_key(match, p.key):
                ex.rekey(match, p.key)
            flags = _update(conn, match, p, import_id, now, opts, edits.get(match["id"]), kept, report)
            row = match
            updated += match["id"] not in taken   # one property merged twice is still one update
        _record_duplicates(conn, row, p, import_id, now, edits.get(row["id"]), report)
        taken.add(row["id"])
        flag_counts.update(flags)
        if progress and i % step == 0:
            _progress(progress, 0.3 + 0.65 * i / max(1, len(records)), f"Saved {i:,} of {len(records):,}")
    unchanged = sum(1 for old, row in before.values() if _stored(row) == old)
    fields = {f: n for f, n in kept.items() if f in EDITABLE_BY_FIELD}
    return {"new": new, "updated": updated, "unchanged": unchanged, "report": report,
            "flag_counts": dict(sorted(flag_counts.items())),
            "edits_kept": sum(fields.values()), "edits_kept_fields": fields, "edits_crexi_new": kept["_crexi_new"]}


def _upgrade_key(existing: Mapping[str, Any], new_key: str) -> bool:
    """Only move a stored key to a more specific one (src -> addr/apn, addr -> apn, apn -> apn|street)."""
    old = existing["property_key"]
    if old.startswith("src:") and not new_key.startswith("src:"):
        return True
    if new_key.startswith("apn:") and not old.startswith("apn:"):
        return True
    return new_key.startswith(old + "|")


def _insert(conn: sqlite3.Connection, market_id: int, p: ParsedRow, flags: list[str], import_id: int,
            now: str) -> dict:
    data = p.record.model_dump()
    data.update(market_id=market_id, property_key=p.key, raw_json=dumps(p.raw), flags_json=dumps(flags),
                first_import_id=import_id, last_import_id=import_id, created_at=now, updated_at=now)
    cols = list(data)
    cur = conn.execute(f"INSERT INTO properties({','.join(cols)}) VALUES({','.join('?' for _ in cols)})",
                       tuple(data[c] for c in cols))
    data["id"] = int(cur.lastrowid)
    return data


def plan_update(existing: Mapping[str, Any], p: ParsedRow,
                edits: Mapping[str, dict] | None = None) -> tuple["cmerge.FactsPlan", dict[str, Any], set[str]]:
    """What merging row `p` into the stored property `existing` would do (crexi/merge.py), without writing:
    (the plan, the incoming values as the merge sees them, the row's flag notes). The import and the
    Crexi-reply preview both use it, so a preview counts exactly what the import then does."""
    old_flags = set(loads(existing.get("flags_json"), []) or [])
    new_values = p.record.model_dump()
    notes = set(p.notes)
    if p.units_source == "addresses" and existing.get("units") is not None and "units_from_addresses" not in old_flags:
        new_values["units"] = None  # an address count never replaces a real unit count
        notes.discard("units_from_addresses")
    skip = set(edits or {})
    if p.units_source == "addresses":
        skip.add("units")            # an estimate that differs from an estimate is not a disagreement
    estimate = ("units",) if ("units_from_addresses" in old_flags and p.units_source == "units") else ()
    return cmerge.plan_facts(existing, new_values, RECORD_COLUMNS, skip=skip, replace_blank_like=estimate), \
        new_values, notes


def _update(conn: sqlite3.Connection, existing: dict, p: ParsedRow, import_id: int, now: str,
            opts: ImportOptions, edits: Mapping[str, dict] | None = None,
            kept: Counter[str] | None = None, report: "cmerge.Report | None" = None) -> list[str]:
    """The stored property is the original and keeps its values (crexi/merge.py): a blank is filled from the
    file, a value the file disagrees about stays and the file's value is saved as an alternate (plus a new
    person / mailing address to try for a different owner or mailing address); empty cells never wipe what
    we already know. Corrections made by hand (edits: {field: {value, old_value}}) win over the file; each one
    the file disagreed with is counted in kept (by field, and '_crexi_new' when the file's value is not even
    the one that was corrected). updated_at only moves when a stored value changes (the 'Row as of' guard of
    the round trip uses it)."""
    before = _stored(existing)
    old_flags = set(loads(existing.get("flags_json"), []) or [])
    plan, new_values, notes = plan_update(existing, p, edits)
    merged = plan.merged
    if new_values["units"] is None and "units_from_addresses" in old_flags and merged["units"] is not None:
        notes.add("units_from_addresses")
    for f, edit in (edits or {}).items():
        if f not in merged:
            continue
        incoming = new_values[f]
        if incoming is not None and not same_value(f, incoming, edit["value"]) and kept is not None:
            kept[f] += 1
            if not same_value(f, incoming, edit.get("old_value")):
                kept["_crexi_new"] += 1
        merged[f] = edit["value"]
        if f == "units":
            notes.discard("units_from_addresses")   # a count typed by hand is a real unit count
    rec = PropertyRecord.model_construct(**merged)  # values are already typed (ours or the DB's)
    notes |= old_flags & STICKY_FLAGS
    flags = sorted(set(final_flags(rec, notes, opts.units_metric, opts.year_in_buybox)) | (old_flags - INGEST_FLAGS))

    raw = dict(loads(existing.get("raw_json"), {}) or {})
    for k, v in p.raw.items():
        if k not in raw or (raw[k] is None and v is not None):   # the original columns are kept; blanks are filled
            raw[k] = v

    data = rec.model_dump()
    data.update(property_key=existing["property_key"], raw_json=dumps(raw), flags_json=dumps(flags),
                last_import_id=import_id, updated_at=now)
    if _stored({**existing, **data}) == before:
        data["updated_at"] = existing.get("updated_at") or now   # nothing changed: keep its 'as of' time
    sets = ",".join(f"{c}=?" for c in data)
    conn.execute(f"UPDATE properties SET {sets} WHERE id=?", (*data.values(), existing["id"]))
    existing.update(data)
    _save_enrichment(conn, existing["id"], plan, import_id, now, report, p.fills)
    return flags


def _save_enrichment(conn: sqlite3.Connection, property_id: int, plan: "cmerge.FactsPlan", import_id: int, now: str,
                     report: "cmerge.Report | None", fills: int = 0) -> None:
    """Store what the merge rule found for one property (alternates, new person, new mailing address) and count
    it in the report. A re-import finds nothing new, so it adds nothing."""
    notes = f"Crexi import #{import_id}"
    alternates = cmerge.apply_alternates(conn, property_id, plan, import_id=import_id, now=now)
    gained = cmerge.add_new_knowledge(conn, property_id, plan, notes=notes)
    if report is not None:
        filled = cmerge.filled_count(plan) + fills
        report.add("blanks_filled", filled)
        report.add("alternates", alternates)
        report.add("new_names", gained["new_names"])
        report.add("new_mailing", gained["new_mailing"])
        report.add("properties")
        if not (filled or alternates or sum(gained.values())):
            report.add("gained_nothing")


def _record_duplicates(conn: sqlite3.Connection, row: dict, p: ParsedRow, import_id: int, now: str,
                       edits: Mapping[str, dict] | None, report: "cmerge.Report") -> None:
    """Later rows of the file about the same property: the saved property is the original, so what they
    disagree about becomes alternates (and a different owner or mailing address something to try)."""
    for dup in p.dups:
        report.add("duplicates")
        plan = cmerge.plan_facts(row, dup.record.model_dump(), RECORD_COLUMNS, skip=set(edits or {}) | {"units"}
                                if dup.units_source == "addresses" else set(edits or {}))
        alternates = cmerge.apply_alternates(conn, row["id"], plan, import_id=import_id, now=now)
        gained = cmerge.add_new_knowledge(conn, row["id"], plan, notes=f"Crexi import #{import_id}")
        report.add("alternates", alternates)
        report.add("new_names", gained["new_names"])
        report.add("new_mailing", gained["new_mailing"])
