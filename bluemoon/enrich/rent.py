"""Rent check: HUD Fair Market Rents for each property, compared with the asset class rent scale.

Where the rents come from (`config.yaml: rent.source`):
- `hud_arcgis` (default; free, no key): HUD's public map service. Area FMRs come from the
  Fair_Market_Rents layer; ZIP-level Small Area FMRs (SAFMRs) from HUD_PDR_Small_Area_Fair_Market_Rents.
- `hud_file`: a HUD FMR (county) or SAFMR (ZIP) spreadsheet you downloaded into `data/hud/`.
- `hud_api`: the HUD USER API (needs a free `HUD_API_TOKEN` in `.env`).
When the chosen source has nothing for a property, the free map service is tried next
(`rent.fallback_to_arcgis`).

Small Area FMRs are the FMRs HUD publishes for each ZIP inside a metro area. They follow local rents
more closely than one number for the whole metro, so they are used (`rent.prefer_small_area`) when the
property is in a metro FMR area and HUD has a SAFMR for its ZIP in that area. Otherwise the area FMR
(metro or county) is used. FMRs are gross rents (rent plus basic utilities) near the 40th percentile.

The fiscal year comes from HUD's own metadata (the "Date of Coverage" line of the map item, which HUD
keeps current; the layer's own text can be stale), re-read once a month. Every cache key carries the
layer's data-edit stamp, so new HUD rents are fetched when HUD posts them, and a warning says when the
rents are from an older fiscal year than the one in effect.

Being polite to HUD's free service: one request at a time except point lookups (`rent.concurrency`,
default 3); SAFMRs are fetched 50 ZIPs per request and cached per ZIP; a ZIP that lies in a single FMR
area gives the area without a map lookup; other properties get one point lookup cached per point
rounded to 3 decimals (about 100 m), and since FMR areas are whole counties (outside New England) a
county is reused once one lookup confirms its own non-metro area or two lookups in it agree; area rents
are fetched once per FMR area. After repeated failures it stops asking and keeps earlier results.
"""
from __future__ import annotations

import contextlib
import csv
import logging
import re
import sqlite3
import threading
import time
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from collections.abc import Callable, Iterable, Iterator, Mapping
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any

import httpx
from pydantic import BaseModel, ConfigDict, Field, ValidationError, field_validator

from ..cache import cache_get, cache_put, make_key
from ..config import AssetClass, Settings, get_settings, secret
from ..db import dumps, market_by_slug, now_iso, rows, this_month, upsert
from ..util import clean_str, normalize_county, normalize_zip, to_float, to_int

log = logging.getLogger(__name__)

Progress = Callable[[float, str], None] | None

SOURCES = ("hud_arcgis", "hud_file", "hud_api")
BEDROOMS = ("0", "1", "2", "3", "4")
# FMR areas follow whole counties everywhere except New England, where they follow towns.
NEW_ENGLAND = frozenset({"CT", "MA", "ME", "NH", "RI", "VT"})
USER_AGENT = "BlueMoonDealFinder/1.0 (rent check; HUD Fair Market Rents)"
HUD_FILE_EXTS = (".xlsx", ".xlsm", ".csv")

_ARCGIS = "https://services.arcgis.com/VTyQ9soqVukalItT/arcgis/rest/services"
FMR_FIELDS = "FMR_CODE,FMR_AREANAME,FMR_0BDR,FMR_1BDR,FMR_2BDR,FMR_3BDR,FMR_4BDR"
SAFMR_FIELDS = "HUD_CODE,FMR_NAME,SAFMR_0BR,SAFMR_1BR,SAFMR_2BR,SAFMR_3BR,SAFMR_4BR,ID,ZCTA_ID"
ZCTA_FIELDS = "ZCTA5CE20,GEOID"


class HudUnavailable(RuntimeError):
    """HUD's service did not answer (network trouble, outage, rejected token). Nothing is overwritten."""


# --------------------------------------------------------------------------------------------
# Settings
# --------------------------------------------------------------------------------------------

class RentConfig(BaseModel):
    """`config.yaml: rent`. Every key is optional; these defaults apply when a key is missing."""

    model_config = ConfigDict(extra="ignore")

    source: str = "hud_arcgis"
    hud_arcgis_url: str = f"{_ARCGIS}/Fair_Market_Rents/FeatureServer/0/query"
    hud_safmr_url: str = f"{_ARCGIS}/HUD_PDR_Small_Area_Fair_Market_Rents/FeatureServer"
    arcgis_items_url: str = "https://www.arcgis.com/sharing/rest/content/items"
    hud_api_url: str = "https://www.huduser.gov/hudapi/public/fmr"
    hud_api_year: int | None = None          # None = HUD's latest year
    prefer_small_area: bool = True
    fallback_to_arcgis: bool = True
    concurrency: int = 3
    timeout_seconds: float = 30.0
    max_retries: int = 3
    retry_backoff_seconds: float = 1.0
    max_consecutive_errors: int = 8
    zip_batch_size: int = 50
    point_decimals: int = 3

    @classmethod
    def from_settings(cls, settings: Settings, source: str | None = None) -> RentConfig:
        raw = {k: v for k, v in settings.section("rent").items() if v is not None}
        if source:
            raw["source"] = source
        return cls(**raw)

    @property
    def fmr_layer_url(self) -> str:
        return _strip_query(self.hud_arcgis_url)

    @property
    def safmr_service_url(self) -> str:
        return re.sub(r"/\d+$", "", _strip_query(self.hud_safmr_url))

    @property
    def zcta_layer_url(self) -> str:
        return f"{self.safmr_service_url}/0"

    @property
    def safmr_table_url(self) -> str:
        return f"{self.safmr_service_url}/1"


def _strip_query(url: str) -> str:
    return re.sub(r"/query/?$", "", url.strip().rstrip("/"))


# --------------------------------------------------------------------------------------------
# Data models
# --------------------------------------------------------------------------------------------

Rents = dict[str, int | None]


class AreaRents(BaseModel):
    """Rents for one HUD FMR area (a metro or a non-metro county)."""
    code: str
    name: str
    rents: Rents


class SafmrRow(BaseModel):
    """Small Area FMR for one ZIP inside one FMR area."""
    zip: str
    code: str
    name: str
    rents: Rents


class LayerMeta(BaseModel):
    url: str
    fiscal_year: int | None = None
    version: str = "unknown"      # the layer's data-edit stamp; part of every cache key
    basis: str = "unknown"        # item | layer | edit_date | unknown


class RentLookup(BaseModel):
    """What `fmr_for_point` returns (as a dict)."""
    area_name: str
    fiscal_year: str | None = None   # "FY2026"
    rents: Rents
    source: str                       # hud_arcgis | hud_file | hud_api
    kind: str = "area"                # small_area (ZIP-level SAFMR) | area (metro/county FMR)
    zip: str | None = None
    area_code: str | None = None
    area_rents: Rents | None = None


Result = RentLookup | None | HudUnavailable   # per property: rents found, nothing there, or HUD didn't answer


class RentSpec(BaseModel):
    """`asset_classes.yaml: <class>.rent`."""
    model_config = ConfigDict(extra="ignore")

    enabled: bool = False
    scale: dict[int, float] = Field(default_factory=dict)
    focus_bedrooms: int = 2
    margin: float = 0.15                 # within this = "in"
    near_margin: float | None = None     # within this = "near"; default 2 x margin
    scale_label: str = "Brad's"

    @field_validator("scale", mode="before")
    @classmethod
    def _clean_scale(cls, v: Any) -> dict[int, float]:
        out: dict[int, float] = {}
        for k, val in (v or {}).items():
            br, amount = to_int(k), to_float(val)
            if br is not None and amount and amount > 0:
                out[br] = amount
        return out

    @property
    def near(self) -> float:
        return self.near_margin if self.near_margin is not None else 2 * self.margin


@dataclass(frozen=True)
class Location:
    id: int
    lat: float | None = None
    lng: float | None = None
    zip: str | None = None
    state: str | None = None
    county: str | None = None
    city: str | None = None

    @property
    def has_point(self) -> bool:
        return self.lat is not None and self.lng is not None

    @property
    def county_key(self) -> tuple[str, str] | None:
        """(state, county) when FMR areas there follow whole counties."""
        c = _county_name(self.county)
        if not self.state or not c or self.state in NEW_ENGLAND:
            return None
        return (self.state, c)


# ArcGIS and HUD API response shapes ---------------------------------------------------------

class _ArcgisFeature(BaseModel):
    attributes: dict[str, Any] = Field(default_factory=dict)


class _ArcgisQuery(BaseModel):
    model_config = ConfigDict(extra="ignore")
    features: list[_ArcgisFeature] = Field(default_factory=list)
    exceededTransferLimit: bool = False


class _FmrAttrs(BaseModel):
    model_config = ConfigDict(extra="ignore")
    FMR_CODE: str
    FMR_AREANAME: str | None = None
    FMR_0BDR: Any = None
    FMR_1BDR: Any = None
    FMR_2BDR: Any = None
    FMR_3BDR: Any = None
    FMR_4BDR: Any = None

    def to_area(self) -> AreaRents:
        rents = {b: _money(getattr(self, f"FMR_{b}BDR")) for b in BEDROOMS}
        return AreaRents(code=self.FMR_CODE, name=_tidy(self.FMR_AREANAME) or self.FMR_CODE, rents=rents)


class _SafmrAttrs(BaseModel):
    model_config = ConfigDict(extra="ignore")
    HUD_CODE: str
    FMR_NAME: str | None = None
    SAFMR_0BR: Any = None
    SAFMR_1BR: Any = None
    SAFMR_2BR: Any = None
    SAFMR_3BR: Any = None
    SAFMR_4BR: Any = None
    ID: Any = None
    ZCTA_ID: Any = None

    def to_row(self) -> SafmrRow | None:
        z = normalize_zip(self.ID) or normalize_zip(self.ZCTA_ID)
        if not z:
            return None
        rents = {b: _money(getattr(self, f"SAFMR_{b}BR")) for b in BEDROOMS}
        return SafmrRow(zip=z, code=self.HUD_CODE, name=_tidy(self.FMR_NAME) or self.HUD_CODE, rents=rents)


class _LayerJson(BaseModel):
    model_config = ConfigDict(extra="ignore")
    name: str | None = None
    description: str | None = None
    serviceItemId: str | None = None
    editingInfo: dict[str, Any] | None = None


class _ItemJson(BaseModel):
    model_config = ConfigDict(extra="ignore")
    title: str | None = None
    snippet: str | None = None
    description: str | None = None


class _HudApiCounty(BaseModel):
    model_config = ConfigDict(extra="ignore")
    state_code: str | None = None
    fips_code: str
    county_name: str | None = None
    town_name: str | None = None
    category: str | None = None


class _HudApiData(BaseModel):
    model_config = ConfigDict(extra="ignore")
    county_name: str | None = None
    town_name: str | None = None
    metro_status: str | None = None
    metro_name: str | None = None
    area_name: str | None = None
    smallarea_status: str | None = None
    basicdata: Any = None
    year: str | None = None

    @field_validator("metro_status", "smallarea_status", "year", mode="before")
    @classmethod
    def _as_text(cls, v: Any) -> str | None:
        return None if v is None else str(v)


# --------------------------------------------------------------------------------------------
# Small helpers
# --------------------------------------------------------------------------------------------

def _tidy(v: Any) -> str | None:
    s = clean_str(v)
    return re.sub(r"\s+", " ", s) if s else None


def _money(v: Any) -> int | None:
    f = to_float(v)
    return int(round(f)) if f is not None and f > 0 else None


def _has_rents(rents: Mapping[str, Any] | None) -> bool:
    return bool(rents) and any(v for v in rents.values())


def _county_name(c: str | None) -> str | None:
    """'St. Clair County' -> 'ST CLAIR' so HUD and Crexi spellings compare equal."""
    s = normalize_county(c)
    if not s:
        return None
    s = re.sub(r"[.'`]", "", s)
    s = re.sub(r"\s+(PARISH|BOROUGH|CENSUS AREA|MUNICIPALITY)$", "", s)
    return re.sub(r"\s+", " ", s).strip() or None


def _state(v: Any) -> str | None:
    s = clean_str(v)
    return s.upper() if s and len(s) == 2 and s.isalpha() else None


def fy_label(year: int | None) -> str | None:
    return f"FY{year}" if year else None


def current_federal_fy(today: date | None = None) -> int:
    """HUD's fiscal year starts October 1: Oct 2026 is FY2027."""
    d = today or date.today()
    return d.year + 1 if d.month >= 10 else d.year


_FY_COVERAGE = re.compile(r"Date of Coverage\W*(?:FY|Fiscal\s+Year)\s*(\d{4})", re.I)
_FY_ANY = re.compile(r"(?:\bFY|\bFiscal\s+Year)\s*(\d{4})\b", re.I)


def fiscal_year_from_text(text: str | None) -> int | None:
    """Pull the fiscal year out of HUD's layer text ("Date of Coverage: FY2026 ...")."""
    if not text:
        return None
    t = re.sub(r"<[^>]+>", " ", text)
    t = re.sub(r"(&nbsp;|\xa0|\s)+", " ", t)
    m = _FY_COVERAGE.search(t)
    if m:
        return int(m.group(1))
    years = [int(y) for y in _FY_ANY.findall(t)]
    return max(years) if years else None


def fiscal_year_from_edit(epoch_ms: Any) -> int | None:
    """HUD posts next year's FMRs in August-September, so an edit then means next fiscal year."""
    ms = to_float(epoch_ms)
    if not ms:
        return None
    d = datetime.fromtimestamp(ms / 1000, tz=UTC)
    return d.year + 1 if d.month >= 8 else d.year


def _plural(n: int, word: str, many: str | None = None) -> str:
    if n == 1:
        return f"1 {word}"
    if many is None:
        many = word[:-1] + "ies" if word.endswith("y") and word[-2:-1] not in "aeiou" else word + "s"
    return f"{n:,} {many}"


# --------------------------------------------------------------------------------------------
# Rent fit (pure)
# --------------------------------------------------------------------------------------------

def _bedroom_label(br: int) -> str:
    return "studio" if br == 0 else f"{br}BR"


def _usd(v: float) -> str:
    return f"${v:,.0f}"


def _gap_text(gap: float) -> str:
    pct = round(abs(gap) * 100)
    if pct == 0:
        return "about even"
    return f"{pct}% {'under' if gap < 0 else 'over'}"


def _fit_for_gap(gap: float, spec: RentSpec) -> str:
    g = abs(gap) - 1e-9
    if g <= spec.margin:
        return "in"
    if g <= spec.near:
        return "near"
    return "out"


def _clean_rents(rents: Mapping[Any, Any] | None) -> dict[int, float]:
    out: dict[int, float] = {}
    for k, v in (rents or {}).items():
        br, amount = to_int(k), to_float(v)
        if br is not None and amount and amount > 0:
            out[br] = amount
    return out


def rent_fit(rents: Mapping[Any, Any] | None, rent_spec: Mapping[str, Any] | RentSpec | None,
             label: str | None = None) -> tuple[str, str]:
    """Compare area rents with the asset class rent scale.

    Returns (fit, detail): fit is in | near | out | na. The focus bedroom size decides the fit
    (in = within `margin`, near = within `near_margin` (default 2 x margin), else out); the average
    gap across the scale's bedroom sizes is added as extra information.
    Example detail: "Area 2BR rent $1,150 vs Brad's $1,200 (4% under); studio-3BR average 2% under".
    """
    spec = rent_spec if isinstance(rent_spec, RentSpec) else RentSpec(**dict(rent_spec or {}))
    if not spec.enabled:
        return "na", f"Rent check not used for {label or 'this asset class'}"
    if not spec.scale:
        return "na", "No rent scale set for this asset class (config/asset_classes.yaml)"
    area = _clean_rents(rents)
    if not area:
        return "na", "No HUD rent figures for this location"
    gaps = {br: (area[br] - target) / target for br, target in sorted(spec.scale.items()) if br in area}
    if not gaps:
        return "na", "HUD rents don't cover the bedroom sizes in the rent scale"

    average = sum(gaps.values()) / len(gaps)
    span = f"{_bedroom_label(min(gaps))}-{_bedroom_label(max(gaps))}"
    focus = spec.focus_bedrooms
    if focus in gaps:
        fit = _fit_for_gap(gaps[focus], spec)
        detail = (f"Area {_bedroom_label(focus)} rent {_usd(area[focus])} vs {spec.scale_label} "
                  f"{_usd(spec.scale[focus])} ({_gap_text(gaps[focus])})")
        if len(gaps) > 1:
            detail += f"; {span} average {_gap_text(average)}"
        return fit, detail
    fit = _fit_for_gap(average, spec)
    missing = _bedroom_label(focus)
    return fit, f"Area {span} rents average {_gap_text(average)} vs {spec.scale_label} scale (no {missing} figure)"


# --------------------------------------------------------------------------------------------
# HTTP with retries, counting and a circuit breaker
# --------------------------------------------------------------------------------------------

class _Http:
    RETRY = frozenset({429, 500, 502, 503, 504, 599})

    def __init__(self, client: httpx.Client, cfg: RentConfig, sleep: Callable[[float], None] = time.sleep):
        self.client = client
        self.cfg = cfg
        self.sleep = sleep
        self.requests = 0
        self.cache_hits = 0
        self.failures = 0
        self._consecutive = 0
        self._lock = threading.Lock()

    @property
    def tripped(self) -> bool:
        return self._consecutive >= self.cfg.max_consecutive_errors

    def hit(self) -> None:
        with self._lock:
            self.cache_hits += 1

    def get_json(self, url: str, params: Mapping[str, Any] | None = None,
                 headers: Mapping[str, str] | None = None) -> tuple[int, Any]:
        """GET a JSON document. Returns (status, body); ArcGIS errors sent with HTTP 200 count as errors."""
        if self.tripped:
            return 503, {"error": {"message": "Skipped: the service stopped answering."}}
        status, body = 599, {"error": {"message": "No response."}}
        for attempt in range(self.cfg.max_retries + 1):
            if attempt:
                self.sleep(self.cfg.retry_backoff_seconds * 2 ** (attempt - 1))
            status, body = self._once(url, params, headers)
            if status not in self.RETRY:
                break
        with self._lock:
            if status >= 400 and status != 404:
                self.failures += 1
                self._consecutive += 1
                log.warning("HUD rent lookup failed (%s): %s", status, _error_text(body))
            else:
                self._consecutive = 0
        return status, body

    def _once(self, url: str, params: Mapping[str, Any] | None,
              headers: Mapping[str, str] | None) -> tuple[int, Any]:
        with self._lock:
            self.requests += 1
        try:
            r = self.client.get(url, params=dict(params or {}), headers=dict(headers or {}))
        except httpx.HTTPError as e:
            return 599, {"error": {"message": f"Network problem ({type(e).__name__})."}}
        try:
            body = r.json()
        except ValueError:
            return (r.status_code if r.status_code >= 400 else 502), {
                "error": {"message": "The service answered with something that isn't data (maybe a bot check)."}}
        if r.status_code >= 400:
            return r.status_code, body
        err = body.get("error") if isinstance(body, dict) else None
        if err:
            code = to_int(err.get("code")) if isinstance(err, dict) else None
            return (code if code and code >= 400 else 400), body
        return r.status_code, body


def _error_text(body: Any) -> str:
    if isinstance(body, dict):
        err = body.get("error")
        if isinstance(err, dict):
            return str(err.get("message") or err)
        if err:
            return str(err)
    return str(body)[:200]


@contextlib.contextmanager
def _client_scope(client: httpx.Client | None, cfg: RentConfig) -> Iterator[httpx.Client]:
    if client is not None:
        yield client
        return
    own = httpx.Client(timeout=cfg.timeout_seconds, headers={"User-Agent": USER_AGENT}, follow_redirects=True)
    try:
        yield own
    finally:
        own.close()


# --------------------------------------------------------------------------------------------
# Source 1: HUD's free map service (ArcGIS)
# --------------------------------------------------------------------------------------------

_FAILED = "__failed__"   # area code for a location whose lookup failed (not the same as "HUD has nothing here")


class HudArcgis:
    """Area FMRs and Small Area FMRs from HUD's public ArcGIS feature services."""

    PROVIDER = "hud_arcgis"

    def __init__(self, conn: sqlite3.Connection, http: _Http, cfg: RentConfig):
        self.conn = conn
        self.http = http
        self.cfg = cfg
        self._meta: dict[str, LayerMeta] = {}
        self._areas: dict[str, AreaRents | None] = {}
        self._zip_rows: dict[str, list[SafmrRow]] = {}
        self._points: dict[tuple[str, str], AreaRents | None] = {}
        self._failed_points: set[tuple[str, str]] = set()
        self._zctas: dict[tuple[str, str], str | None] = {}
        self._county_areas: dict[tuple[str, str], str] = {}
        self.small_area_down = False
        self.warnings: list[str] = []

    # --- cache plumbing ---------------------------------------------------------------------
    def _key(self, *parts: Any) -> str:
        return make_key(self.PROVIDER, *parts)

    def _from_cache(self, key: str) -> tuple[bool, Any]:
        hit = cache_get(self.conn, key)
        if hit is not None and (hit["status"] or 0) < 400:
            self.http.hit()
            return True, hit["body"]
        return False, None

    def _get_cached(self, key: str, url: str, params: Mapping[str, Any]) -> Any:
        """One cached GET. Raises HudUnavailable on failure (failures are never cached)."""
        found, body = self._from_cache(key)
        if found:
            return body
        status, body = self.http.get_json(url, params)
        if status >= 400:
            raise HudUnavailable(f"HUD's map service didn't answer ({status}): {_error_text(body)}")
        cache_put(self.conn, key, self.PROVIDER, status, body)
        return body

    # --- metadata ---------------------------------------------------------------------------
    def meta(self, layer_url: str) -> LayerMeta:
        """Fiscal year and data version of a layer (re-read once a month)."""
        if layer_url not in self._meta:
            self._meta[layer_url] = self._read_meta(layer_url)
        return self._meta[layer_url]

    def fmr_meta(self) -> LayerMeta:
        return self.meta(self.cfg.fmr_layer_url)

    def safmr_meta(self) -> LayerMeta:
        return self.meta(self.cfg.safmr_table_url)

    def _read_meta(self, layer_url: str) -> LayerMeta:
        month = this_month()
        try:
            layer = _LayerJson.model_validate(
                self._get_cached(self._key("layer_meta", layer_url, month), layer_url, {"f": "json"}))
        except (HudUnavailable, ValidationError) as e:
            log.warning("Could not read HUD layer details for %s: %s", layer_url, e)
            return LayerMeta(url=layer_url)
        edited = (layer.editingInfo or {}).get("dataLastEditDate")
        version = str(edited) if edited else "unknown"
        item_fy = self._item_fiscal_year(layer.serviceItemId, month)
        if item_fy:
            return LayerMeta(url=layer_url, fiscal_year=item_fy, version=version, basis="item")
        layer_fy, edit_fy = fiscal_year_from_text(layer.description), fiscal_year_from_edit(edited)
        if layer_fy and (not edit_fy or layer_fy >= edit_fy):
            return LayerMeta(url=layer_url, fiscal_year=layer_fy, version=version, basis="layer")
        if edit_fy:
            return LayerMeta(url=layer_url, fiscal_year=edit_fy, version=version, basis="edit_date")
        return LayerMeta(url=layer_url, version=version)

    def _item_fiscal_year(self, item_id: str | None, month: str) -> int | None:
        if not item_id or not re.fullmatch(r"[0-9a-fA-F]{32}", item_id):
            return None
        url = f"{self.cfg.arcgis_items_url.rstrip('/')}/{item_id}"
        try:
            item = _ItemJson.model_validate(self._get_cached(self._key("item_meta", url, month), url, {"f": "json"}))
        except (HudUnavailable, ValidationError) as e:
            log.warning("Could not read HUD item details: %s", e)
            return None
        return fiscal_year_from_text(item.description) or fiscal_year_from_text(item.snippet)

    # --- queries ----------------------------------------------------------------------------
    def _point(self, lat: float, lng: float) -> tuple[str, str]:
        d = self.cfg.point_decimals
        return (f"{lat:.{d}f}", f"{lng:.{d}f}")

    @staticmethod
    def _point_params(pt: tuple[str, str], out_fields: str) -> dict[str, str]:
        return {"geometry": f"{pt[1]},{pt[0]}", "geometryType": "esriGeometryPoint", "inSR": "4326",
                "spatialRel": "esriSpatialRelIntersects", "outFields": out_fields,
                "returnGeometry": "false", "f": "json"}

    def _point_key(self, pt: tuple[str, str]) -> str:
        return self._key("fmr_point", self.cfg.fmr_layer_url, self.fmr_meta().version, *pt)

    def _remember_point(self, pt: tuple[str, str], body: Any) -> None:
        areas = _parse_fmr_features(body)
        for a in areas:
            self._areas.setdefault(a.code, a)
        self._points[pt] = areas[0] if areas else None

    def prefetch_points(self, pts: Iterable[tuple[str, str]]) -> None:
        """Look up FMR areas for many rounded points: cache first, then a few requests at a time."""
        todo: list[tuple[str, str]] = []
        for pt in dict.fromkeys(pts):
            if pt in self._points or pt in self._failed_points:
                continue
            found, body = self._from_cache(self._point_key(pt))
            if found:
                self._remember_point(pt, body)
            else:
                todo.append(pt)
        if not todo:
            return
        url = f"{self.cfg.fmr_layer_url}/query"

        def fetch(pt: tuple[str, str]) -> tuple[int, Any]:
            return self.http.get_json(url, self._point_params(pt, FMR_FIELDS))

        with ThreadPoolExecutor(max_workers=max(1, self.cfg.concurrency)) as pool:
            for pt, (status, body) in zip(todo, pool.map(fetch, todo), strict=True):
                if status >= 400:
                    self._failed_points.add(pt)
                    continue
                cache_put(self.conn, self._point_key(pt), self.PROVIDER, status, body)
                self._remember_point(pt, body)

    def area_at_point(self, lat: float, lng: float) -> AreaRents | None:
        pt = self._point(lat, lng)
        self._failed_points.discard(pt)  # ask again
        self.prefetch_points([pt])
        if pt in self._failed_points:
            raise HudUnavailable("HUD's map service didn't answer the location lookup.")
        return self._points.get(pt)

    def zcta_at_point(self, lat: float, lng: float) -> str | None:
        """ZIP (Census ZCTA) at a point, for properties without a ZIP."""
        pt = self._point(lat, lng)
        if pt not in self._zctas:
            meta = self.meta(self.cfg.zcta_layer_url)
            body = self._get_cached(self._key("zcta_point", self.cfg.zcta_layer_url, meta.version, *pt),
                                    f"{self.cfg.zcta_layer_url}/query", self._point_params(pt, ZCTA_FIELDS))
            zips = [normalize_zip(f.attributes.get("ZCTA5CE20") or f.attributes.get("GEOID"))
                    for f in _ArcgisQuery.model_validate(body).features]
            self._zctas[pt] = next((z for z in zips if z), None)
        return self._zctas[pt]

    def _query_all(self, url: str, params: Mapping[str, Any]) -> list[dict[str, Any]]:
        """Attribute query with paging. Raises HudUnavailable on failure."""
        out: list[dict[str, Any]] = []
        offset = 0
        for _ in range(50):
            page = {**params, "resultOffset": str(offset), "orderByFields": "OBJECTID", "f": "json"}
            status, body = self.http.get_json(url, page)
            if status >= 400:
                raise HudUnavailable(f"HUD's map service didn't answer ({status}): {_error_text(body)}")
            q = _ArcgisQuery.model_validate(body)
            out.extend(f.attributes for f in q.features)
            if not q.exceededTransferLimit or not q.features:
                break
            offset += len(q.features)
        return out

    def safmr_rows(self, zips: Iterable[str | None]) -> dict[str, list[SafmrRow]]:
        """Small Area FMR rows per ZIP: cache first, then one request per `zip_batch_size` ZIPs."""
        wanted = sorted({z for z in zips if z and re.fullmatch(r"\d{5}", z)})
        version = self.safmr_meta().version
        todo: list[str] = []
        for z in wanted:
            if z in self._zip_rows:
                continue
            found, body = self._from_cache(self._key("safmr_zip", self.cfg.safmr_table_url, version, z))
            if found:
                self._zip_rows[z] = [SafmrRow.model_validate(r) for r in body or []]
            else:
                todo.append(z)
        size = max(1, self.cfg.zip_batch_size)
        for i in range(0, len(todo), size):
            batch = todo[i:i + size]
            where = "ID IN ({})".format(",".join(f"'{z}'" for z in batch))
            found_rows: dict[str, list[SafmrRow]] = {z: [] for z in batch}
            for attrs in self._query_all(f"{self.cfg.safmr_table_url}/query",
                                         {"where": where, "outFields": SAFMR_FIELDS, "returnGeometry": "false"}):
                row = _parse_model(_SafmrAttrs, attrs)
                row = row.to_row() if row else None
                if row and row.zip in found_rows:
                    found_rows[row.zip].append(row)
            for z, zrows in found_rows.items():
                cache_put(self.conn, self._key("safmr_zip", self.cfg.safmr_table_url, version, z), self.PROVIDER,
                          200, [r.model_dump() for r in zrows])
                self._zip_rows[z] = zrows
        return {z: self._zip_rows.get(z, []) for z in wanted}

    def areas_by_code(self, codes: Iterable[str]) -> dict[str, AreaRents | None]:
        """Area FMRs by HUD area code (e.g. METRO44220M44220): cache first, then batched requests."""
        wanted = sorted({c for c in codes if c and re.fullmatch(r"[A-Za-z0-9]+", c)})
        version = self.fmr_meta().version
        todo: list[str] = []
        for c in wanted:
            if c in self._areas:
                continue
            found, body = self._from_cache(self._key("fmr_code", self.cfg.fmr_layer_url, version, c))
            if found:
                self._areas[c] = AreaRents.model_validate(body) if body else None
            else:
                todo.append(c)
        size = max(1, self.cfg.zip_batch_size)
        for i in range(0, len(todo), size):
            batch = todo[i:i + size]
            where = "FMR_CODE IN ({})".format(",".join(f"'{c}'" for c in batch))
            got: dict[str, AreaRents] = {}
            for attrs in self._query_all(f"{self.cfg.fmr_layer_url}/query",
                                         {"where": where, "outFields": FMR_FIELDS, "returnGeometry": "false"}):
                parsed = _parse_model(_FmrAttrs, attrs)
                if parsed:
                    got.setdefault(parsed.FMR_CODE, parsed.to_area())
            for c in batch:
                area = got.get(c)
                cache_put(self.conn, self._key("fmr_code", self.cfg.fmr_layer_url, version, c), self.PROVIDER,
                          200, area.model_dump() if area else None)
                self._areas[c] = area
        return {c: self._areas.get(c) for c in wanted}

    # --- one property -----------------------------------------------------------------------
    def _single_code_from_zip(self, loc: Location) -> str | None:
        codes = {r.code for r in self._zip_rows.get(loc.zip or "", [])}
        return codes.pop() if len(codes) == 1 else None

    def _guess_code_from_zip(self, loc: Location) -> str | None:
        """No map point and a ZIP that spans areas: take its metro area (SAFMRs exist only in metros)."""
        zrows = self._zip_rows.get(loc.zip or "", [])
        metro = [r.code for r in zrows if r.code.upper().startswith("METRO")]
        return (metro or [r.code for r in zrows] or [None])[0]

    def _small_area_for(self, loc: Location, code: str) -> SafmrRow | None:
        """The ZIP's Small Area FMR inside this metro area, if HUD has one (SAFMRs exist only in metros)."""
        if not self.cfg.prefer_small_area or self.small_area_down or not code.upper().startswith("METRO"):
            return None
        try:
            zip_code = loc.zip
            if not zip_code and loc.has_point:
                zip_code = self.zcta_at_point(loc.lat, loc.lng)  # type: ignore[arg-type]
            if not zip_code:
                return None
            if zip_code not in self._zip_rows:
                self.safmr_rows([zip_code])
        except HudUnavailable as e:
            self._small_area_failed(e)
            return None
        return next((r for r in self._zip_rows.get(zip_code, []) if r.code == code and _has_rents(r.rents)), None)

    def _small_area_failed(self, e: HudUnavailable) -> None:
        if not self.small_area_down:
            self.small_area_down = True
            self.warnings.append(f"HUD's Small Area (ZIP) rent service didn't answer ({e}); used area rents "
                                 "for the rest of this run.")

    def build(self, loc: Location, code: str | None) -> RentLookup | None:
        if code == _FAILED:
            raise HudUnavailable("HUD's map service didn't answer the location lookup.")
        area = self._areas.get(code) if code else None
        if area is None:
            return None
        row = self._small_area_for(loc, code)
        if row:
            return RentLookup(area_name=f"{area.name}, ZIP {row.zip}",
                              fiscal_year=fy_label(self.safmr_meta().fiscal_year), rents=row.rents,
                              source=self.PROVIDER, kind="small_area", zip=row.zip, area_code=code,
                              area_rents=area.rents)
        if not _has_rents(area.rents):
            return None
        return RentLookup(area_name=area.name, fiscal_year=fy_label(self.fmr_meta().fiscal_year), rents=area.rents,
                          source=self.PROVIDER, kind="area", zip=loc.zip, area_code=code, area_rents=area.rents)

    # --- many properties --------------------------------------------------------------------
    def lookup_many(self, locs: list[Location], progress: Progress = None) -> dict[int, Result]:
        _tell(progress, 0.05, "Reading HUD rent layer details")
        self.fmr_meta()
        codes: dict[int, str | None] = {}
        _tell(progress, 0.15, "Getting Small Area rents by ZIP")
        try:
            self.safmr_rows(loc.zip for loc in locs)
        except HudUnavailable as e:
            self._small_area_failed(e)
        try:
            _tell(progress, 0.3, "Finding the HUD rent area for each property")
            self._resolve_codes(locs, codes)
            _tell(progress, 0.6, "Getting area rents")
            self.areas_by_code(c for c in codes.values() if c and c != _FAILED)
        except HudUnavailable as e:
            return {loc.id: e for loc in locs}
        out: dict[int, Result] = {}
        for loc in locs:
            try:
                out[loc.id] = self.build(loc, codes.get(loc.id))
            except HudUnavailable as e:
                out[loc.id] = e
        return out

    def _resolve_codes(self, locs: list[Location], codes: dict[int, str | None]) -> None:
        """Area code per location: ZIP inside one area > county settled by map lookups > own point > ZIP guess.

        FMR areas are whole counties (outside New England), so a county is settled when one map lookup
        returns that county's own non-metro area, or two different points in it return the same area.
        """
        need_point: list[Location] = []
        for loc in locs:
            code = self._single_code_from_zip(loc)
            if code:
                codes[loc.id] = code
            elif loc.has_point:
                need_point.append(loc)
            else:
                codes[loc.id] = self._guess_code_from_zip(loc)
        groups: dict[tuple[str, str] | None, list[Location]] = defaultdict(list)
        for loc in need_point:
            groups[loc.county_key].append(loc)
        rest: list[Location] = groups.pop(None, [])
        samples = {key: list(dict.fromkeys(self._pt(loc) for loc in group))[:2] for key, group in groups.items()}
        # Round 1: one point per county (plus properties without a county).
        self.prefetch_points([pts[0] for pts in samples.values()] + [self._pt(loc) for loc in rest])
        # Round 1b: a second point for counties the first point didn't settle.
        self.prefetch_points([pts[1] for key, pts in samples.items()
                              if len(pts) > 1 and not self._settled_county(key, pts[:1])])
        for key, group in groups.items():
            code = self._settled_county(key, samples[key])
            if code:
                for loc in group:
                    codes[loc.id] = code
            else:
                rest.extend(group)
        # Round 2: everyone else, one point each (cached per rounded point).
        self.prefetch_points([self._pt(loc) for loc in rest])
        for loc in rest:
            pt = self._pt(loc)
            area = self._points.get(pt)
            codes[loc.id] = _FAILED if pt in self._failed_points else (area.code if area else None)

    def _pt(self, loc: Location) -> tuple[str, str]:
        return self._point(loc.lat, loc.lng)  # type: ignore[arg-type]

    def _settled_county(self, key: tuple[str, str], pts: list[tuple[str, str]]) -> str | None:
        if key in self._county_areas:
            return self._county_areas[key]
        areas = [a for a in (self._points.get(pt) for pt in pts) if a is not None]
        if not areas or len(areas) != len(pts) or len({a.code for a in areas}) != 1:
            return None
        if _confirms_county(areas[0], key) or len(areas) == 2:
            self._county_areas[key] = areas[0].code
            return areas[0].code
        return None


def _confirms_county(area: AreaRents, key: tuple[str, str]) -> bool:
    """True when the area is the property's own non-metro county ("Champaign County, OH")."""
    if not area.code.upper().startswith("NCNTY"):
        return False
    name, _, st = area.name.rpartition(",")
    return _county_name(name) == key[1] and st.strip().upper() == key[0]


def _parse_model(model: type[BaseModel], data: Any) -> Any:
    try:
        return model.model_validate(data)
    except ValidationError:
        return None


def _parse_fmr_features(body: Any) -> list[AreaRents]:
    q = _parse_model(_ArcgisQuery, body)
    if not q:
        return []
    areas = [_parse_model(_FmrAttrs, f.attributes) for f in q.features]
    return [a.to_area() for a in areas if a]


# --------------------------------------------------------------------------------------------
# Source 2: a HUD spreadsheet in data/hud/
# --------------------------------------------------------------------------------------------

_RENT_WORDS = {"efficiency": "0", "studio": "0", "zerobedroom": "0", "onebedroom": "1", "twobedroom": "2",
               "threebedroom": "3", "fourbedroom": "4"}
_RENT_RE = re.compile(r"^(?:(?:safmr|fmr|rent)(\d)(?:br|bdr|bed|bedroom|bedrooms)?|(\d)(?:br|bdr|bedroom|bedrooms))$")
_ZIP_HEADERS = {"zip", "zipcode", "zip5", "zcta", "zcta5", "postalcode"}
_CODE_HEADERS = {"hudareacode", "fmrcode", "hudcode", "areacode", "metrocode", "fmrareacode", "cbsasub"}
_STATE_HEADERS = ("stusps", "statealpha", "stateabbr", "stateabbreviation", "stalpha", "st", "state")
_COUNTY_HEADERS = {"countyname", "cntyname", "county"}
_TOWN_HEADERS = {"countytownname", "townname"}
_HEADER_SCAN_ROWS = 10


def _norm_header(h: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(h or "").lower())


def rent_column(header: Any) -> str | None:
    """Bedroom count for a HUD rent column ('fmr_2', 'SAFMR\\n2BR', 'Two-Bedroom' -> '2'), else None.
    Payment-standard columns ('SAFMR 2BR - 90% Payment Standard') are not rent columns."""
    h = _norm_header(header)
    if h in _RENT_WORDS:
        return _RENT_WORDS[h]
    m = _RENT_RE.match(h)
    return (m.group(1) or m.group(2)) if m else None


def map_hud_columns(headers: list[Any]) -> dict[str, int]:
    """Column index per role: zip, code, name, state, county, town, rent:0 .. rent:4."""
    cols: dict[str, int] = {}
    normed = [_norm_header(h) for h in headers]
    for i, h in enumerate(normed):
        br = rent_column(headers[i])
        if br is not None:
            cols.setdefault(f"rent:{br}", i)
        elif h in _ZIP_HEADERS:
            cols.setdefault("zip", i)
        elif h in _CODE_HEADERS:
            cols.setdefault("code", i)
        elif h in _COUNTY_HEADERS:
            cols.setdefault("county", i)
        elif h in _TOWN_HEADERS:
            cols.setdefault("town", i)
        elif "areaname" in h or h in {"fmrname", "metroname"}:
            cols.setdefault("name", i)
    return cols


def fiscal_year_from_filename(name: str) -> int | None:
    m = re.search(r"fy[\s_-]*(\d{4}|\d{2})(?!\d)", name, re.I)
    if m:
        y = int(m.group(1))
        return y + 2000 if y < 100 else y
    m = re.search(r"(?<!\d)(20\d{2})(?!\d)", name)
    return int(m.group(1)) if m else None


class HudFileRow(BaseModel):
    zip: str | None = None
    code: str | None = None
    name: str | None = None
    state: str | None = None
    county: str | None = None
    town: str | None = None
    rents: Rents


@dataclass
class HudFile:
    path: Path
    kind: str                       # safmr (by ZIP) | fmr (by county)
    fiscal_year: int | None
    rows: list[HudFileRow]
    by_zip: dict[str, list[HudFileRow]] = field(default_factory=dict)
    by_county: dict[tuple[str, str], list[HudFileRow]] = field(default_factory=dict)

    def __post_init__(self) -> None:
        for r in self.rows:
            if r.zip:
                self.by_zip.setdefault(r.zip, []).append(r)
            c = _county_name(r.county)
            if r.state and c:
                self.by_county.setdefault((r.state, c), []).append(r)

    def county_row(self, loc: Location) -> HudFileRow | None:
        c = _county_name(loc.county)
        if not loc.state or not c:
            return None
        found = self.by_county.get((loc.state, c), [])
        if len(found) > 1 and loc.city:  # New England: one row per town
            city = loc.city.strip().upper()
            found = [r for r in found if (r.town or "").upper().startswith(city)] or found
        return found[0] if found else None


_FILE_MEMO: dict[str, tuple[tuple[int, int], HudFile | None]] = {}   # path -> (mtime, size), parsed file


def _grid_xlsx(path: Path) -> Iterator[list[list[Any]]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            yield [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def _grid_csv(path: Path) -> Iterator[list[list[Any]]]:
    data = path.read_bytes()
    text = ""
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    yield [list(r) for r in csv.reader(text.splitlines())]


def _find_header(grid: list[list[Any]]) -> tuple[int, dict[str, int]] | None:
    """Row index and column map of the header row (HUD sometimes puts a title row above it)."""
    for i, row in enumerate(grid[:_HEADER_SCAN_ROWS]):
        cols = map_hud_columns(list(row))
        if sum(1 for k in cols if k.startswith("rent:")) >= 3 and (
                "zip" in cols or "county" in cols or "code" in cols):
            return i, cols
    return None


def _row_from_cells(cells: list[Any], cols: dict[str, int], state_idx: list[int]) -> HudFileRow | None:
    def cell(role: str) -> Any:
        i = cols.get(role)
        return cells[i] if i is not None and i < len(cells) else None

    rents = {b: _money(cell(f"rent:{b}")) for b in BEDROOMS}
    if not _has_rents(rents):
        return None
    state = next((s for s in (_state(cells[i]) for i in state_idx if i < len(cells)) if s), None)
    zip_raw = cell("zip")
    zip_code = normalize_zip(f"{int(zip_raw):05d}" if isinstance(zip_raw, (int, float)) else zip_raw)
    return HudFileRow(zip=zip_code, code=clean_str(cell("code")), name=_tidy(cell("name")), state=state,
                      county=clean_str(cell("county")), town=clean_str(cell("town")), rents=rents)


def read_hud_file(path: str | Path) -> HudFile | None:
    """Parse a HUD FMR (county-level) or SAFMR (ZIP-level) spreadsheet. None when it isn't one."""
    p = Path(path)
    st = p.stat()
    memo_key, stamp = str(p.resolve()), (st.st_mtime_ns, st.st_size)
    if memo_key in _FILE_MEMO and _FILE_MEMO[memo_key][0] == stamp:
        return _FILE_MEMO[memo_key][1]
    grids = _grid_csv(p) if p.suffix.lower() == ".csv" else _grid_xlsx(p)
    result: HudFile | None = None
    for grid in grids:
        found = _find_header(grid)
        if not found:
            continue
        i, cols = found
        normed = [_norm_header(h) for h in grid[i]]
        state_idx = [normed.index(s) for s in _STATE_HEADERS if s in normed]  # 2-letter columns first
        parsed = [r for r in (_row_from_cells(list(c), cols, state_idx) for c in grid[i + 1:]) if r]
        kind = "safmr" if "zip" in cols else "fmr"
        result = HudFile(path=p, kind=kind, fiscal_year=fiscal_year_from_filename(p.name), rows=parsed)
        break
    _FILE_MEMO[memo_key] = (stamp, result)
    return result


def hud_files(folder: Path) -> list[HudFile]:
    """Every readable HUD spreadsheet in data/hud/, newest fiscal year first."""
    if not folder.exists():
        return []
    files: list[HudFile] = []
    for p in sorted(folder.iterdir()):
        if p.suffix.lower() not in HUD_FILE_EXTS or p.name.startswith(("~$", ".")):
            continue
        try:
            f = read_hud_file(p)
        except Exception as e:  # a broken download shouldn't stop the run
            log.warning("Could not read HUD file %s: %s", p.name, e)
            continue
        if f and f.rows:
            files.append(f)
    return sorted(files, key=lambda f: f.fiscal_year or 0, reverse=True)


def hud_file_lookup(files: list[HudFile], loc: Location, prefer_small_area: bool = True) -> RentLookup | None:
    area_row, area_file = None, None
    for f in (f for f in files if f.kind == "fmr"):
        area_row = f.county_row(loc)
        if area_row:
            area_file = f
            break
    area_code = area_row.code if area_row else None
    if prefer_small_area and loc.zip and not (area_code and not area_code.upper().startswith("METRO")):
        for f in (f for f in files if f.kind == "safmr"):
            zrows = [r for r in f.by_zip.get(loc.zip, []) if not area_code or r.code == area_code]
            if zrows:
                row = zrows[0]
                base = row.name or (area_row.name if area_row else None) or "HUD Small Area FMR"
                return RentLookup(area_name=f"{base}, ZIP {loc.zip}", fiscal_year=fy_label(f.fiscal_year),
                                  rents=row.rents, source="hud_file", kind="small_area", zip=loc.zip,
                                  area_code=row.code, area_rents=area_row.rents if area_row else None)
    if area_row and area_file:
        return RentLookup(area_name=area_row.name or f"{area_row.county}, {area_row.state}",
                          fiscal_year=fy_label(area_file.fiscal_year), rents=area_row.rents, source="hud_file",
                          kind="area", zip=loc.zip, area_code=area_code, area_rents=area_row.rents)
    return None


# --------------------------------------------------------------------------------------------
# Source 3: HUD USER API (free token)
# --------------------------------------------------------------------------------------------

class HudApi:
    """https://www.huduser.gov/hudapi/public/fmr — listCounties/{state} and data/{county fips}."""

    PROVIDER = "hud_api"

    def __init__(self, conn: sqlite3.Connection, http: _Http, cfg: RentConfig, token: str):
        self.conn = conn
        self.http = http
        self.cfg = cfg
        self._headers = {"Authorization": f"Bearer {token}"}  # never part of a cache key or a log line
        self._counties: dict[str, list[_HudApiCounty]] = {}
        self._data: dict[str, _HudApiData | None] = {}
        self.rejected: str | None = None   # set when HUD turns the token down; then we stop asking

    def _get(self, key: str, url: str, params: Mapping[str, Any] | None = None) -> Any | None:
        if self.rejected:
            raise HudUnavailable(self.rejected)
        hit = cache_get(self.conn, key)
        if hit is not None and (hit["status"] or 0) < 400:
            self.http.hit()
            return hit["body"]
        status, body = self.http.get_json(url, params, self._headers)
        if status == 404:
            return None
        if status in (401, 403):
            self.rejected = ("HUD turned down the HUD_API_TOKEN in .env (check the token, and that the Fair "
                             "Market Rents dataset is selected for it at huduser.gov).")
            raise HudUnavailable(self.rejected)
        if status >= 400:
            raise HudUnavailable(f"The HUD USER API didn't answer ({status}): {_error_text(body)}")
        cache_put(self.conn, key, self.PROVIDER, status, body)
        return body

    def _year_part(self) -> str:
        return str(self.cfg.hud_api_year) if self.cfg.hud_api_year else f"latest:{this_month()}"

    def counties(self, state: str) -> list[_HudApiCounty]:
        if state not in self._counties:
            url = f"{self.cfg.hud_api_url.rstrip('/')}/listCounties/{state}"
            body = self._get(make_key(self.PROVIDER, "counties", url, this_month()), url)
            items = body.get("data", []) if isinstance(body, dict) else (body or [])
            self._counties[state] = [c for c in (_parse_model(_HudApiCounty, i) for i in items) if c]
        return self._counties[state]

    def fips_for(self, loc: Location) -> str | None:
        c = _county_name(loc.county)
        if not loc.state or not c:
            return None
        found = [x for x in self.counties(loc.state) if _county_name(x.county_name) == c]
        if len(found) > 1 and loc.city:
            city = loc.city.strip().upper()
            found = [x for x in found if (x.town_name or "").upper().startswith(city)] or found
        return found[0].fips_code if found else None

    def data(self, fips: str) -> _HudApiData | None:
        if fips not in self._data:
            url = f"{self.cfg.hud_api_url.rstrip('/')}/data/{fips}"
            params = {"year": str(self.cfg.hud_api_year)} if self.cfg.hud_api_year else None
            body = self._get(make_key(self.PROVIDER, "data", url, self._year_part()), url, params)
            inner = body.get("data") if isinstance(body, dict) else None
            self._data[fips] = _parse_model(_HudApiData, inner) if isinstance(inner, dict) else None
        return self._data[fips]

    def lookup(self, loc: Location) -> RentLookup | None:
        fips = self.fips_for(loc)
        d = self.data(fips) if fips else None
        if not d or d.basicdata is None:
            return None
        name = d.area_name or d.metro_name or d.county_name or fips or "HUD FMR area"
        fy = f"FY{d.year}" if d.year and d.year.isdigit() else fy_label(self.cfg.hud_api_year)
        if isinstance(d.basicdata, dict):
            rents = _rents_from_mapping(d.basicdata)
            return RentLookup(area_name=name, fiscal_year=fy, rents=rents, source=self.PROVIDER, kind="area",
                              zip=loc.zip, area_code=fips, area_rents=rents) if _has_rents(rents) else None
        zrows = [r for r in d.basicdata if isinstance(r, dict)]
        area = next((r for r in zrows if "msa" in str(r.get("zip_code", "")).lower()), None)
        area_rents = _rents_from_mapping(area) if area else None
        mine = next((r for r in zrows if normalize_zip(r.get("zip_code")) == loc.zip), None) if loc.zip else None
        if mine and self.cfg.prefer_small_area and _has_rents(_rents_from_mapping(mine)):
            return RentLookup(area_name=f"{name}, ZIP {loc.zip}", fiscal_year=fy, rents=_rents_from_mapping(mine),
                              source=self.PROVIDER, kind="small_area", zip=loc.zip, area_code=fips,
                              area_rents=area_rents)
        if area_rents and _has_rents(area_rents):
            return RentLookup(area_name=name, fiscal_year=fy, rents=area_rents, source=self.PROVIDER, kind="area",
                              zip=loc.zip, area_code=fips, area_rents=area_rents)
        return None


def _rents_from_mapping(m: Mapping[str, Any]) -> Rents:
    rents: Rents = {b: None for b in BEDROOMS}
    for k, v in m.items():
        br = rent_column(k)
        if br is not None and rents.get(br) is None:
            rents[br] = _money(v)
    return rents


# --------------------------------------------------------------------------------------------
# Choosing a source, with fallback to the free map service
# --------------------------------------------------------------------------------------------

class RentLookups:
    """Runs the configured source and falls back to HUD's free map service."""

    def __init__(self, conn: sqlite3.Connection, settings: Settings, client: httpx.Client,
                 source: str | None = None):
        self.cfg = RentConfig.from_settings(settings, source)
        self.http = _Http(client, self.cfg)
        self.arcgis = HudArcgis(conn, self.http, self.cfg)
        self.warnings: list[str] = []
        self.files: list[HudFile] = []
        self.api: HudApi | None = None
        self.primary = self._pick_primary(conn, settings)

    def _pick_primary(self, conn: sqlite3.Connection, settings: Settings) -> str:
        src = (self.cfg.source or "hud_arcgis").strip().lower()
        if src not in SOURCES:
            self.warnings.append(f"rent.source '{src}' isn't one of {', '.join(SOURCES)}; used HUD's free map service.")
            return "hud_arcgis"
        if src == "hud_file":
            folder = _hud_dir(settings)
            self.files = hud_files(folder)
            if not self.files:
                self.warnings.append(f"rent.source is hud_file but no HUD spreadsheet was found in {folder.name}/; "
                                     "used HUD's free map service instead.")
                return "hud_arcgis"
        if src == "hud_api":
            token = secret("HUD_API_TOKEN")
            if not token:
                self.warnings.append("rent.source is hud_api but HUD_API_TOKEN is blank in .env; "
                                     "used HUD's free map service instead.")
                return "hud_arcgis"
            self.api = HudApi(conn, self.http, self.cfg, token)
        return src

    def lookup_many(self, locs: list[Location], progress: Progress = None) -> dict[int, Result]:
        if self.primary == "hud_arcgis":
            return self.arcgis.lookup_many(locs, progress)
        out: dict[int, Result] = {}
        for i, loc in enumerate(locs):
            if self.primary == "hud_api" and self.http.tripped:
                out[loc.id] = HudUnavailable("The HUD USER API stopped answering.")
                continue
            try:
                out[loc.id] = (self.api.lookup(loc) if self.api is not None
                               else hud_file_lookup(self.files, loc, self.cfg.prefer_small_area))
            except HudUnavailable as e:
                out[loc.id] = e
            if i % 100 == 0:
                _tell(progress, 0.05 + 0.5 * i / max(1, len(locs)), "Looking up HUD rents")
        if self.api is not None and self.api.rejected:
            self.warnings.append(self.api.rejected)
        missing = [loc for loc in locs if not isinstance(out.get(loc.id), RentLookup)]
        if missing and self.cfg.fallback_to_arcgis:
            name = "your HUD spreadsheet" if self.primary == "hud_file" else "the HUD USER API"
            self.warnings.append(f"{_plural(len(missing), 'property')} had no rents in {name}; "
                                 "tried HUD's free map service for those.")
            for pid, res in self.arcgis.lookup_many(missing, None).items():
                if isinstance(res, RentLookup) or not isinstance(out.get(pid), HudUnavailable):
                    out[pid] = res
        return out


def _hud_dir(settings: Settings) -> Path:
    try:
        return settings.path("hud_dir")
    except KeyError:
        return settings.root / "data" / "hud"


def _tell(progress: Progress, fraction: float, message: str) -> None:
    if progress:
        try:
            progress(max(0.0, min(1.0, fraction)), message)
        except Exception:  # a broken progress callback must not stop the work
            log.debug("progress callback failed", exc_info=True)


# --------------------------------------------------------------------------------------------
# Public API
# --------------------------------------------------------------------------------------------

def fmr_for_point(conn: sqlite3.Connection, lat: float, lng: float, zip_code: str | None = None, *,
                  state: str | None = None, county: str | None = None, city: str | None = None,
                  settings: Settings | None = None, client: httpx.Client | None = None,
                  source: str | None = None) -> dict | None:
    """HUD rents at a point: {area_name, fiscal_year, rents:{"0".."4"}, source, kind, zip, area_code,
    area_rents}. Uses the Small Area FMR for the ZIP when there is one. None when nothing is found
    or HUD didn't answer (the reason is logged)."""
    s = settings or get_settings()
    loc = Location(id=0, lat=to_float(lat), lng=to_float(lng), zip=normalize_zip(zip_code), state=_state(state),
                   county=clean_str(county), city=clean_str(city))
    cfg = RentConfig.from_settings(s, source)
    with _client_scope(client, cfg) as c:
        res = RentLookups(conn, s, c, source).lookup_many([loc]).get(0)
    if isinstance(res, HudUnavailable):
        log.warning("HUD rent lookup failed: %s", res)
        return None
    return res.model_dump() if res else None


def _load_locations(conn: sqlite3.Connection, market_id: int, limit: int | None) -> list[Location]:
    sql = "SELECT id, lat, lng, zip, state, county, city FROM properties WHERE market_id=? ORDER BY id"
    params: list[Any] = [market_id]
    if limit:
        sql += " LIMIT ?"
        params.append(int(limit))
    out = []
    for r in rows(conn, sql, params):
        lat, lng = to_float(r["lat"]), to_float(r["lng"])
        if lat is None or lng is None or not (-90 <= lat <= 90 and -180 <= lng <= 180) or (lat == 0 and lng == 0):
            lat = lng = None
        out.append(Location(id=r["id"], lat=lat, lng=lng, zip=normalize_zip(r["zip"]), state=_state(r["state"]),
                            county=clean_str(r["county"]), city=clean_str(r["city"])))
    return out


def _save(conn: sqlite3.Connection, pid: int, fit: str, detail: str, res: RentLookup | None, now: str) -> None:
    upsert(conn, "rent_fits", "property_id", {
        "property_id": pid,
        "source": res.source if res else None,
        "area_name": res.area_name if res else None,
        "fiscal_year": res.fiscal_year if res else None,
        "rents_json": dumps(res.rents) if res else None,
        "fit": fit,
        "detail": detail,
        "computed_at": now,
    })


def _no_rent_reason(loc: Location) -> str:
    if loc.has_point or loc.zip:
        return "No HUD rent figures found for this location"
    if loc.state and loc.county:
        return "Only a county, no map point or ZIP, so HUD rents couldn't be matched"
    return "No map point, ZIP or county to look up HUD rents"


def _not_used(conn: sqlite3.Connection, market_slug: str, ac: AssetClass, locs: list[Location],
              dry_run: bool) -> dict:
    detail = f"Rent check not used for {ac.label}"
    if not dry_run:
        now = now_iso()
        for loc in locs:
            _save(conn, loc.id, "na", detail, None, now)
        conn.commit()
    return {"market": market_slug, "asset_class": ac.key, "enabled": False, "dry_run": dry_run,
            "properties": len(locs), "fits": {"in": 0, "near": 0, "out": 0, "na": len(locs)},
            "cost": {"usd": 0.0}, "warnings": [],
            "message": f"{detail}; {_plural(len(locs), 'property')} marked not applicable."}


def compute_rent_fits(conn: sqlite3.Connection, market_slug: str, progress: Progress = None, *,
                      limit: int | None = None, dry_run: bool = False, settings: Settings | None = None,
                      client: httpx.Client | None = None, source: str | None = None,
                      today: date | None = None) -> dict:
    """Look up HUD rents for every property in the market and grade them against the rent scale.

    Free (HUD data), cached, and safe to re-run. When HUD doesn't answer, earlier results are kept.
    Returns counts, the fiscal year(s) used, request counts, warnings and a plain-English message.
    """
    s = settings or get_settings()
    market = market_by_slug(conn, market_slug)
    ac = s.asset_class(market["asset_class"])
    locs = _load_locations(conn, market["id"], limit)
    if not ac.rent_enabled:
        return _not_used(conn, market_slug, ac, locs, dry_run)

    cfg = RentConfig.from_settings(s, source)
    findable = [loc for loc in locs if loc.has_point or loc.zip or (loc.state and loc.county)]
    if dry_run:
        return {"market": market_slug, "asset_class": ac.key, "enabled": True, "dry_run": True,
                "properties": len(locs), "with_location": len(findable), "source": cfg.source,
                "cost": {"usd": 0.0},
                "message": (f"Rent check would look up HUD rents for {_plural(len(findable), 'property')} "
                            f"({len(locs) - len(findable):,} have no map point, ZIP or county). It's free.")}

    _tell(progress, 0.02, f"Rent check for {_plural(len(locs), 'property')}")
    with _client_scope(client, cfg) as c:
        lookups = RentLookups(conn, s, c, source)
        results = lookups.lookup_many(findable, progress)

    _tell(progress, 0.85, "Saving rent fits")
    fits: Counter[str] = Counter()
    kinds: Counter[str] = Counter()
    used_sources: Counter[str] = Counter()
    years: Counter[str] = Counter()
    errors = no_data = 0
    now = now_iso()
    for loc in locs:
        res = results.get(loc.id)
        if isinstance(res, HudUnavailable):
            errors += 1
            continue  # keep whatever was saved before
        if res is None:
            _save(conn, loc.id, "na", _no_rent_reason(loc), None, now)
            fits["na"] += 1
            no_data += 1
            continue
        fit, detail = rent_fit(res.rents, ac.rent, ac.label)
        _save(conn, loc.id, fit, detail, res, now)
        fits[fit] += 1
        kinds[res.kind] += 1
        used_sources[res.source] += 1
        if res.fiscal_year:
            years[res.fiscal_year] += 1
    conn.commit()

    warnings = lookups.warnings + lookups.arcgis.warnings
    warnings += _stale_warning(years, today, lookups.primary)
    if errors:
        warnings.append(f"HUD didn't answer for {_plural(errors, 'property')}; their earlier rent results "
                        "(if any) were kept. Run the rent step again later.")
    result = {
        "market": market_slug, "asset_class": ac.key, "enabled": True, "dry_run": False,
        "source": lookups.primary, "properties": len(locs), "saved": len(locs) - errors,
        "fits": {k: fits.get(k, 0) for k in ("in", "near", "out", "na")},
        "no_data": no_data, "errors": errors,
        "small_area": kinds.get("small_area", 0), "area": kinds.get("area", 0),
        "sources_used": dict(used_sources), "fiscal_years": sorted(years),
        "requests": lookups.http.requests, "cache_hits": lookups.http.cache_hits,
        "cost": {"usd": 0.0}, "warnings": warnings,
    }
    result["message"] = _summary_message(result)
    _tell(progress, 1.0, result["message"])
    return result


def _stale_warning(years: Counter[str], today: date | None, primary: str) -> list[str]:
    """Say so when some rents are from an older HUD fiscal year than the one in effect today."""
    current = current_federal_fy(today)
    old = {int(y[2:]): n for y, n in years.items() if y[2:].isdigit() and int(y[2:]) < current}
    if not old:
        return []
    count = sum(old.values())
    who = "These are" if count == sum(years.values()) else f"{_plural(count, 'property')} got"
    if primary == "hud_arcgis":
        tip = (f"HUD's free map service hasn't posted FY{current} yet. For the newest numbers, put HUD's "
               f"FY{current} FMR and Small Area FMR spreadsheets in data/hud and set rent.source: hud_file "
               "(or add a free HUD_API_TOKEN and set rent.source: hud_api).")
    else:
        tip = f"Check that the HUD spreadsheets in data/hud (or rent.hud_api_year) are for FY{current}."
    return [f"{who} HUD FY{min(old)} rents, but FY{current} rents took effect October 1. {tip}"]


def _summary_message(r: dict) -> str:
    f = r["fits"]
    msg = (f"Rent check done for {_plural(r['properties'], 'property')}: {f['in']:,} in range, "
           f"{f['near']:,} near, {f['out']:,} out, {f['na']:,} without HUD rents.")
    if r["fiscal_years"]:
        msg += f" HUD {', '.join(r['fiscal_years'])} rents"
        msg += (f" ({r['small_area']:,} by ZIP from Small Area FMRs)." if r["small_area"] else ".")
    if r["errors"]:
        msg += f" HUD didn't answer for {r['errors']:,}; try again later."
    return msg
